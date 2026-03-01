"""
Adapter A: Convert ms-preprocessing-toolkit output → DNP input format.

Transformations:
1. RawIntensity sheet: Ensure first column is 'Mz/RT', remove tolerance column
2. SampleInfo sheet: Copy as-is (format already compatible)

Uses openpyxl to preserve Excel formatting (e.g., red ISTD font markers).
"""

from pathlib import Path

import openpyxl


# Accepted first-column names from ms-preprocessing output (current or legacy)
_ACCEPTED_FEATURE_COLS = {"Mz/RT", "FeatureID"}
_TOLERANCE_COL = "m/z Tolerance( ppm)/RT Tolerance"

# Target column name for DNP
_DNP_FEATURE_COL = "Mz/RT"


def convert_preprocessing_to_dnp(input_path: str, output_path: str) -> str:
    """
    Convert ms-preprocessing output to DNP-compatible format.

    Parameters
    ----------
    input_path : str
        Path to ms-preprocessing output Excel file.
    output_path : str
        Path for the converted output file.

    Returns
    -------
    str
        The output file path.

    Raises
    ------
    FileNotFoundError
        If input file does not exist.
    ValueError
        If required sheets or columns are missing.
    """
    input_path = Path(input_path)
    output_path = Path(output_path)

    if not input_path.exists():
        raise FileNotFoundError(f"Input file not found: {input_path}")

    wb = openpyxl.load_workbook(str(input_path))

    # --- Validate required sheet ---
    if "RawIntensity" not in wb.sheetnames:
        raise ValueError(
            f"'RawIntensity' sheet not found. Available sheets: {wb.sheetnames}"
        )

    ws = wb["RawIntensity"]

    # --- Find column indices (1-based) ---
    header_row = 1
    feature_col_idx = None
    tolerance_col_idx = None

    for col_idx in range(1, ws.max_column + 1):
        cell_value = ws.cell(row=header_row, column=col_idx).value
        if cell_value in _ACCEPTED_FEATURE_COLS:
            feature_col_idx = col_idx
        elif cell_value == _TOLERANCE_COL:
            tolerance_col_idx = col_idx

    if feature_col_idx is None:
        raise ValueError(
            f"Feature ID column not found in RawIntensity header. "
            f"Expected one of {_ACCEPTED_FEATURE_COLS}."
        )

    # --- Ensure column name is 'Mz/RT' ---
    ws.cell(row=header_row, column=feature_col_idx).value = _DNP_FEATURE_COL

    # --- Remove tolerance column (if present) ---
    if tolerance_col_idx is not None:
        ws.delete_cols(tolerance_col_idx)

    # --- Save ---
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(output_path))
    wb.close()

    return str(output_path)
