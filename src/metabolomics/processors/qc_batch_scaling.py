import os
from datetime import datetime

import numpy as np
import pandas as pd
from matplotlib.lines import Line2D
from openpyxl import load_workbook
from sklearn.preprocessing import StandardScaler

from metabolomics.utils.constants import (
    COLORBLIND_COLORS,
    DATETIME_FORMAT_FULL,
    FEATURE_ID_COLUMN,
    SHEET_NAMES,
    resolve_sheet_name,
)
from metabolomics.utils.data_helpers import extract_sample_type_row, insert_sample_type_row
from metabolomics.utils.excel_format import copy_sheet_formatting_only
from metabolomics.utils.file_io import build_output_path, build_plots_dir, resolve_session_dir
from metabolomics.utils.plotting import build_batch_group_indices, setup_matplotlib
from metabolomics.utils.results import ProcessingResult
from metabolomics.utils.sample_classification import (
    identify_sample_columns,
    normalize_sample_name,
    normalize_sample_type,
)
from metabolomics.utils.console import safe_print as print


RESULT_SHEET_NAME = SHEET_NAMES.get("qc_batch_scaling", "QC_Batch_Scaling_result")
SUMMARY_SHEET_NAME = SHEET_NAMES.get("qc_batch_scaling_summary", "QC_Batch_Scaling_summary")

setup_matplotlib()
import matplotlib.pyplot as plt


def log_section(title):
    """Print a compact section header."""
    print(f"\n  [{title}]")


def parse_batch_labels(value):
    """Parse semicolon-separated batch labels and trim whitespace."""
    if pd.isna(value):
        return []
    return [part.strip() for part in str(value).split(";") if part.strip()]


def build_batch_membership(sample_info_df, sample_columns=None):
    """Build batch-to-QC and batch-to-sample mappings from SampleInfo."""
    batch_to_qc = {}
    batch_to_samples = {}
    column_lookup = {}

    if sample_columns:
        column_lookup = {
            normalize_sample_name(column): column
            for column in sample_columns
        }

    for _, row in sample_info_df.iterrows():
        sample_name = str(row["Sample_Name"]).strip()
        if column_lookup:
            sample_name = column_lookup.get(normalize_sample_name(sample_name))
            if not sample_name:
                continue
        sample_type = normalize_sample_type(row.get("Sample_Type", ""))
        batches = parse_batch_labels(row.get("Batch", ""))

        if sample_type != "QC" and len(batches) != 1:
            raise ValueError("Non-QC samples must belong to a single batch")

        for batch in batches:
            batch_to_samples.setdefault(batch, []).append(sample_name)
            if sample_type == "QC":
                batch_to_qc.setdefault(batch, []).append(sample_name)

    return batch_to_qc, batch_to_samples


def scale_feature_by_batch_qc_median(feature_row, batch_to_qc, batch_to_samples):
    """Scale one feature row by each batch's QC median."""
    scaled = dict(feature_row)
    medians = {}
    sample_scaled_values = {}

    for batch, qc_samples in batch_to_qc.items():
        qc_values = pd.to_numeric(
            pd.Series([feature_row.get(sample, np.nan) for sample in qc_samples]),
            errors="coerce",
        )
        qc_values = qc_values[np.isfinite(qc_values) & (qc_values > 0)]
        batch_median = float(np.nanmedian(qc_values)) if not qc_values.empty else np.nan
        medians[batch] = batch_median

        if not np.isfinite(batch_median) or batch_median <= 0:
            continue

        for sample in batch_to_samples.get(batch, []):
            value = pd.to_numeric(pd.Series([feature_row.get(sample, np.nan)]), errors="coerce").iloc[0]
            if np.isfinite(value) and value > 0:
                sample_scaled_values.setdefault(sample, []).append(float(value / batch_median))

    for sample, candidates in sample_scaled_values.items():
        if candidates:
            scaled[sample] = float(np.nanmedian(candidates))

    return scaled, medians


def select_source_sheet(sheet_names):
    """Pick the best upstream sheet for Step 3."""
    for sheet_key in ("qc_lowess", "istd_correction", "raw_intensity"):
        sheet_name = resolve_sheet_name(sheet_names, sheet_key)
        if sheet_name is not None:
            return sheet_name
    raise ValueError("No supported upstream data sheet found for QC batch scaling")


def load_and_process_data(input_file):
    """Load the selected upstream data sheet and SampleInfo."""
    if not os.path.exists(input_file):
        raise FileNotFoundError(input_file)

    excel_file = pd.ExcelFile(input_file)
    if SHEET_NAMES["sample_info"] not in excel_file.sheet_names:
        raise ValueError(f"Missing required sheet: {SHEET_NAMES['sample_info']}")

    source_sheet_name = select_source_sheet(excel_file.sheet_names)
    sample_info_df = pd.read_excel(excel_file, sheet_name=SHEET_NAMES["sample_info"])
    data_df = pd.read_excel(excel_file, sheet_name=source_sheet_name)

    feature_col = data_df.columns[0]
    data_df, sample_type_row = extract_sample_type_row(data_df, feature_col)
    if FEATURE_ID_COLUMN in data_df.columns and FEATURE_ID_COLUMN != "FeatureID":
        data_df = data_df.rename(columns={FEATURE_ID_COLUMN: "FeatureID"})
    elif "FeatureID" not in data_df.columns:
        data_df = data_df.rename(columns={feature_col: "FeatureID"})

    sample_columns, _ = identify_sample_columns(data_df, sample_info_df)
    data_df.attrs["sample_columns"] = sample_columns
    data_df.attrs["source_sheet_name"] = source_sheet_name

    return data_df, sample_info_df, sample_columns, source_sheet_name, sample_type_row


def scale_dataframe_by_qc_medians(data_df, sample_columns, batch_to_qc, batch_to_samples):
    """Apply QC median scaling feature by feature."""
    result_df = data_df.copy()
    invalid_median_counts = {batch: 0 for batch in batch_to_qc}
    total_features = len(data_df)
    progress_checkpoints = {
        1,
        total_features,
        max(1, total_features // 4),
        max(1, total_features // 2),
        max(1, (total_features * 3) // 4),
    }

    for idx, row in data_df.iterrows():
        scaled_row, medians = scale_feature_by_batch_qc_median(row, batch_to_qc, batch_to_samples)
        for sample in sample_columns:
            if sample in scaled_row:
                result_df.at[idx, sample] = scaled_row[sample]
        for batch, median in medians.items():
            if not np.isfinite(median) or median <= 0:
                invalid_median_counts[batch] += 1

        if (idx + 1) in progress_checkpoints:
            print(f"  - feature scaling 進度: {idx + 1}/{total_features}")

    return result_df, invalid_median_counts


def build_summary_df(source_sheet_name, batch_to_qc, batch_to_samples, invalid_median_counts):
    """Build a simple summary sheet for QC batch scaling."""
    rows = [
        {"Section": "run", "Item": "source_sheet", "Value": source_sheet_name},
        {"Section": "run", "Item": "batch_count", "Value": len(batch_to_qc)},
    ]

    for batch in sorted(batch_to_samples):
        rows.append(
            {
                "Section": "batch",
                "Item": batch,
                "Value": f"qc={len(batch_to_qc.get(batch, []))}; samples={len(batch_to_samples.get(batch, []))}; invalid_feature_medians={invalid_median_counts.get(batch, 0)}",
            }
        )

    return pd.DataFrame(rows)


def build_plot_metadata(sample_columns, sample_info_df):
    """Build sample type and batch metadata aligned with sample columns."""
    sample_info_norm = sample_info_df.copy()
    sample_info_norm["_norm_name"] = sample_info_norm["Sample_Name"].map(normalize_sample_name)
    sample_info_norm = sample_info_norm[sample_info_norm["_norm_name"].astype(bool)]
    sample_meta = sample_info_norm.drop_duplicates("_norm_name").set_index("_norm_name")

    sample_types = []
    batch_memberships = []
    qc_indices = []

    for index, sample in enumerate(sample_columns):
        meta_key = normalize_sample_name(sample)
        if meta_key in sample_meta.index:
            raw_type = str(sample_meta.loc[meta_key].get("Sample_Type", "Unknown"))
            raw_batch = sample_meta.loc[meta_key].get("Batch", "Unknown")
        else:
            raw_type = "Unknown"
            raw_batch = "Unknown"

        sample_type = normalize_sample_type(raw_type)
        memberships = tuple(parse_batch_labels(raw_batch) or ["Unknown"])

        sample_types.append(sample_type)
        batch_memberships.append(memberships)

        if sample_type == "QC":
            qc_indices.append(index)

    return sample_types, batch_memberships, qc_indices


def prepare_residual_matrix(df, sample_columns):
    """Build a standardized sample-by-feature matrix for residual analysis."""
    data_matrix = (
        df[sample_columns]
        .apply(pd.to_numeric, errors="coerce")
        .T
        .to_numpy(dtype=float)
    )
    data_matrix = np.nan_to_num(data_matrix, nan=0.0, posinf=0.0, neginf=0.0)
    data_matrix[data_matrix < 0] = 0

    non_zero_features = np.any(data_matrix != 0, axis=0)
    data_matrix = data_matrix[:, non_zero_features]
    if data_matrix.shape[1] < 1:
        return None

    data_matrix = np.log2(data_matrix + 1)
    return StandardScaler().fit_transform(data_matrix)


def calculate_batch_residuals(scaled_matrix, sample_columns, batch_memberships):
    """Calculate per-batch residuals against the global mean for each feature."""
    if scaled_matrix is None:
        return {}
    if scaled_matrix.shape[0] != len(sample_columns):
        raise ValueError("scaled_matrix rows must align with sample_columns")

    group_indices = build_batch_group_indices(batch_memberships)
    if not group_indices:
        return {}

    global_mean = np.mean(scaled_matrix, axis=0)
    residuals = {}

    for batch, indices in group_indices.items():
        if not indices:
            continue
        batch_data = scaled_matrix[indices]
        if batch_data.size == 0:
            continue
        batch_mean = np.mean(batch_data, axis=0)
        residuals[batch] = batch_mean - global_mean

    return residuals


def plot_batch_residual_analysis(
    residuals_before,
    residuals_after,
    *,
    output_path=None,
    dpi=300,
):
    """Plot before/after residual scatter for each batch."""
    unique_batches = sorted(set(residuals_before) | set(residuals_after))
    colors = COLORBLIND_COLORS * ((len(unique_batches) // len(COLORBLIND_COLORS)) + 1)
    batch_color_map = {batch: colors[i] for i, batch in enumerate(unique_batches)}

    fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(16, 6))

    for batch in unique_batches:
        before = residuals_before.get(batch)
        if before is not None and len(before):
            ax1.scatter(
                np.arange(len(before)),
                before,
                color=batch_color_map[batch],
                alpha=0.6,
                s=30,
                label=f"Batch {batch}",
            )

        after = residuals_after.get(batch)
        if after is not None and len(after):
            ax2.scatter(
                np.arange(len(after)),
                after,
                color=batch_color_map[batch],
                alpha=0.6,
                s=30,
                label=f"Batch {batch}",
            )

    for ax, title in (
        (ax1, "Before Correction\n(Systematic batch bias visible)"),
        (ax2, "After Correction\n(Random scatter around zero = good)"),
    ):
        ax.axhline(y=0, color="red", linestyle="--", linewidth=2, label="Zero Line")
        ax.set_xlabel("Feature Index", fontsize=12, fontweight="bold")
        ax.set_ylabel("Residual (Batch Mean - Global Mean)", fontsize=12, fontweight="bold")
        ax.set_title(title, fontsize=14, fontweight="bold", pad=15)
        ax.grid(True, alpha=0.3)

    all_residuals = []
    for values in list(residuals_before.values()) + list(residuals_after.values()):
        if values is not None and len(values):
            all_residuals.extend(np.asarray(values, dtype=float).tolist())

    y_max = max(np.abs(all_residuals)) if all_residuals else 1.0
    ax1.set_ylim(-y_max * 1.1, y_max * 1.1)
    ax2.set_ylim(-y_max * 1.1, y_max * 1.1)

    legend_handles = [
        Line2D(
            [],
            [],
            linestyle="",
            marker="o",
            markersize=7,
            color=batch_color_map[batch],
            alpha=0.8,
            label=f"Batch {batch}",
        )
        for batch in unique_batches
    ]
    legend_handles.append(
        Line2D([], [], color="red", linestyle="--", linewidth=2, label="Zero Line")
    )

    plt.suptitle("QC Batch Scaling Residual Analysis", fontsize=16, fontweight="bold", y=0.98)
    ax1.legend(
        handles=legend_handles,
        loc="upper center",
        bbox_to_anchor=(1.05, 1.22),
        ncol=min(4, len(legend_handles)),
        fontsize=9,
        frameon=True,
    )
    plt.tight_layout(rect=[0, 0, 1, 0.84])

    if output_path is not None:
        plt.savefig(output_path, dpi=dpi)

    return fig


def calculate_batch_qc_feature_medians(df, batch_to_qc):
    """Calculate per-feature QC medians for each batch."""
    medians = {}
    for batch, qc_samples in batch_to_qc.items():
        valid_qc_samples = [sample for sample in qc_samples if sample in df.columns]
        if not valid_qc_samples:
            continue
        qc_matrix = df[valid_qc_samples].apply(pd.to_numeric, errors="coerce").to_numpy(dtype=float)
        qc_matrix = np.where(qc_matrix > 0, qc_matrix, np.nan)
        batch_feature_medians = np.nanmedian(qc_matrix, axis=1)
        batch_feature_medians = batch_feature_medians[np.isfinite(batch_feature_medians) & (batch_feature_medians > 0)]
        if len(batch_feature_medians):
            medians[batch] = batch_feature_medians
    return medians


def plot_batch_qc_median_alignment(
    feature_medians_before,
    feature_medians_after,
    *,
    output_path=None,
    dpi=300,
):
    """Plot batch-wise QC median alignment before and after scaling."""
    batches = sorted(set(feature_medians_before) & set(feature_medians_after))
    if len(batches) < 2:
        return None

    before_summary = []
    after_summary = []
    for batch in batches:
        before_log = np.log2(np.asarray(feature_medians_before[batch], dtype=float))
        after_log = np.log2(np.asarray(feature_medians_after[batch], dtype=float))
        before_summary.append(
            (
                np.nanmedian(before_log),
                np.nanpercentile(before_log, 25),
                np.nanpercentile(before_log, 75),
            )
        )
        after_summary.append(
            (
                np.nanmedian(after_log),
                np.nanpercentile(after_log, 25),
                np.nanpercentile(after_log, 75),
            )
        )

    x = np.arange(len(batches), dtype=float)
    fig, ax = plt.subplots(figsize=(10, 6))

    before_medians = np.array([row[0] for row in before_summary], dtype=float)
    after_medians = np.array([row[0] for row in after_summary], dtype=float)
    before_err = np.vstack([
        before_medians - np.array([row[1] for row in before_summary], dtype=float),
        np.array([row[2] for row in before_summary], dtype=float) - before_medians,
    ])
    after_err = np.vstack([
        after_medians - np.array([row[1] for row in after_summary], dtype=float),
        np.array([row[2] for row in after_summary], dtype=float) - after_medians,
    ])

    ax.errorbar(
        x - 0.08,
        before_medians,
        yerr=before_err,
        color="#c44e52",
        marker="o",
        linewidth=2,
        capsize=4,
        label="Before",
    )
    ax.errorbar(
        x + 0.08,
        after_medians,
        yerr=after_err,
        color="#55a868",
        marker="o",
        linewidth=2,
        capsize=4,
        label="After",
    )

    ax.set_xticks(x)
    ax.set_xticklabels([f"Batch {batch}" for batch in batches], fontsize=10)
    ax.set_ylabel("Median log2(QC feature median intensity)", fontsize=11, fontweight="bold")
    ax.set_title("Batch QC Median Alignment", fontsize=15, fontweight="bold")
    ax.grid(True, axis="y", alpha=0.3)
    ax.legend(fontsize=10)
    plt.tight_layout()

    if output_path is not None:
        plt.savefig(output_path, dpi=dpi, bbox_inches="tight")

    return fig


def calculate_sample_log_medians(df, sample_columns):
    """Summarize each sample by the median log2 intensity across features."""
    if not sample_columns:
        return {}

    data_matrix = df[sample_columns].apply(pd.to_numeric, errors="coerce").to_numpy(dtype=float)
    data_matrix = np.where(data_matrix > 0, np.log2(data_matrix), np.nan)
    sample_medians = np.nanmedian(data_matrix, axis=0)

    result = {}
    for sample, median in zip(sample_columns, sample_medians):
        if np.isfinite(median):
            result[sample] = float(median)
    return result


def plot_batch_boxplot(
    source_df,
    result_df,
    sample_columns,
    batch_memberships,
    *,
    output_path=None,
    dpi=300,
):
    """Plot sample-level intensity distributions grouped by batch."""
    batch_groups = build_batch_group_indices(batch_memberships)
    if len(batch_groups) < 2:
        return None

    sample_medians_before = calculate_sample_log_medians(source_df, sample_columns)
    sample_medians_after = calculate_sample_log_medians(result_df, sample_columns)
    ordered_batches = list(batch_groups.keys())
    before_data = []
    after_data = []

    for batch in ordered_batches:
        indices = batch_groups[batch]
        before_values = [
            sample_medians_before[sample_columns[index]]
            for index in indices
            if sample_columns[index] in sample_medians_before
        ]
        after_values = [
            sample_medians_after[sample_columns[index]]
            for index in indices
            if sample_columns[index] in sample_medians_after
        ]
        if before_values and after_values:
            before_data.append(before_values)
            after_data.append(after_values)
        else:
            before_data.append([])
            after_data.append([])

    if sum(bool(values) for values in before_data) < 2 or sum(bool(values) for values in after_data) < 2:
        return None

    fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(14, 6), sharey=True)
    positions = np.arange(1, len(ordered_batches) + 1)

    for ax, dataset, title, color in (
        (ax1, before_data, "Before Scaling", "#c9d6df"),
        (ax2, after_data, "After Scaling", "#bfe3c0"),
    ):
        safe_dataset = [values if values else [np.nan] for values in dataset]
        box = ax.boxplot(
            safe_dataset,
            positions=positions,
            patch_artist=True,
            widths=0.6,
            medianprops=dict(color="#222222", linewidth=2),
        )
        for patch in box["boxes"]:
            patch.set_facecolor(color)
            patch.set_alpha(0.85)

        for pos, values in zip(positions, dataset):
            if not values:
                continue
            jitter = np.random.uniform(-0.08, 0.08, size=len(values))
            ax.scatter(np.full(len(values), pos) + jitter, values, color="#4c4c4c", alpha=0.55, s=20)

        ax.set_title(title, fontsize=13, fontweight="bold")
        ax.set_xticks(positions)
        ax.set_xticklabels([f"Batch {batch}" for batch in ordered_batches], rotation=0)
        ax.grid(True, axis="y", alpha=0.25)
        ax.set_xlabel("Batch", fontsize=11, fontweight="bold")

    ax1.set_ylabel("Sample median log2(intensity)", fontsize=11, fontweight="bold")
    fig.suptitle("Batch-wise Sample Distribution", fontsize=16, fontweight="bold", y=0.98)
    plt.tight_layout(rect=[0, 0, 1, 0.95])

    if output_path is not None:
        plt.savefig(output_path, dpi=dpi, bbox_inches="tight")

    return fig


def generate_step3_plots(
    source_df,
    result_df,
    sample_columns,
    sample_info_df,
    input_file,
    timestamp,
    plots_dir=None,
):
    """Generate diagnostic plots for QC Batch Scaling."""
    created_figures = []
    if plots_dir is None:
        plots_dir = build_plots_dir(
            "QC_Batch_Scaling_plots",
            input_file=input_file,
            timestamp=timestamp,
            session_prefix="QC_Batch_Scaling",
        )

    if len(sample_columns) < 3:
        return str(plots_dir)

    batch_to_qc, _ = build_batch_membership(sample_info_df, sample_columns=sample_columns)
    _, batch_memberships, _ = build_plot_metadata(sample_columns, sample_info_df)
    unique_batches = sorted({batch for memberships in batch_memberships for batch in memberships if batch})

    if len(unique_batches) < 2:
        print("  - 只有單一 batch，跳過 Step 3 batch diagnostics")
        return str(plots_dir)

    qc_feature_medians_before = calculate_batch_qc_feature_medians(source_df, batch_to_qc)
    qc_feature_medians_after = calculate_batch_qc_feature_medians(result_df, batch_to_qc)
    if len(qc_feature_medians_before) >= 2 and len(qc_feature_medians_after) >= 2:
        print("  - 生成 Fig1: Batch QC median alignment")
        alignment_plot_path = os.path.join(plots_dir, f"Step3_Batch_QC_Median_Alignment_{timestamp}.png")
        fig = plot_batch_qc_median_alignment(
            qc_feature_medians_before,
            qc_feature_medians_after,
            output_path=alignment_plot_path,
            dpi=300,
        )
        if fig is not None:
            created_figures.append(fig)

    print("  - 生成 Fig2: Batch boxplot")
    batch_boxplot_path = os.path.join(plots_dir, f"Step3_Batch_Boxplot_{timestamp}.png")
    fig = plot_batch_boxplot(
        source_df,
        result_df,
        sample_columns,
        batch_memberships,
        output_path=batch_boxplot_path,
        dpi=300,
    )
    if fig is not None:
        created_figures.append(fig)

    residual_source = prepare_residual_matrix(source_df, sample_columns)
    residual_result = prepare_residual_matrix(result_df, sample_columns)
    if residual_source is not None and residual_result is not None:
        residuals_before = calculate_batch_residuals(
            residual_source,
            sample_columns,
            batch_memberships,
        )
        residuals_after = calculate_batch_residuals(
            residual_result,
            sample_columns,
            batch_memberships,
        )
        if residuals_before and residuals_after:
            print("  - 生成 Fig3: Residual analysis")
            residual_plot_path = os.path.join(plots_dir, f"Step3_Residual_Analysis_{timestamp}.png")
            fig = plot_batch_residual_analysis(
                residuals_before,
                residuals_after,
                output_path=residual_plot_path,
                dpi=300,
            )
            if fig is not None:
                created_figures.append(fig)

    for fig in created_figures:
        plt.close(fig)
    return str(plots_dir)


def generate_pca_plots(*args, **kwargs):
    """Backward-compatible alias for the Step 3 diagnostics entry point."""
    return generate_step3_plots(*args, **kwargs)


def save_results_to_excel(
    source_df,
    result_df,
    sample_info_df,
    summary_df,
    output_file,
    input_file,
    source_sheet_name,
    sample_type_row=None,
):
    """Save the new Step 3 workbook, keeping only the selected upstream sheet."""
    source_export = source_df.copy()
    result_export = result_df.copy()

    def _rename_feature_col_for_output(df):
        if "FeatureID" in df.columns and FEATURE_ID_COLUMN != "FeatureID":
            return df.rename(columns={"FeatureID": FEATURE_ID_COLUMN})
        return df

    source_export = _rename_feature_col_for_output(source_export)
    result_export = _rename_feature_col_for_output(result_export)

    if sample_type_row is not None:
        source_export = insert_sample_type_row(source_export, sample_type_row, feature_col=FEATURE_ID_COLUMN)
        result_export = insert_sample_type_row(result_export, sample_type_row, feature_col=FEATURE_ID_COLUMN)

    with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
        source_export.to_excel(writer, sheet_name=source_sheet_name, index=False)
        sample_info_df.to_excel(writer, sheet_name=SHEET_NAMES["sample_info"], index=False)
        result_export.to_excel(writer, sheet_name=RESULT_SHEET_NAME, index=False)
        summary_df.to_excel(writer, sheet_name=SUMMARY_SHEET_NAME, index=False)

    original_wb = load_workbook(input_file)
    new_wb = load_workbook(output_file)
    try:
        if source_sheet_name in original_wb.sheetnames and source_sheet_name in new_wb.sheetnames:
            copy_sheet_formatting_only(original_wb[source_sheet_name], new_wb[source_sheet_name])
        new_wb.save(output_file)
    finally:
        original_wb.close()
        new_wb.close()


def main(input_file=None, session_dir=None):
    """Run QC batch scaling on the selected upstream sheet."""
    if not input_file:
        raise ValueError("input_file is required")

    log_section("開始執行 Step 3: QC Batch Scaling")
    print(f"輸入檔案: {os.path.basename(input_file)}")

    data_df, sample_info_df, sample_columns, source_sheet_name, sample_type_row = load_and_process_data(input_file)
    print(f"來源工作表: {source_sheet_name}")
    print(f"待處理特徵數: {len(data_df)}")
    print(f"樣本數: {len(sample_columns)}")

    log_section("建立 batch membership")
    batch_to_qc, batch_to_samples = build_batch_membership(
        sample_info_df,
        sample_columns=sample_columns,
    )
    print(f"Batch 數: {len(batch_to_samples)}")
    for batch in sorted(batch_to_samples):
        print(
            f"  - Batch {batch}: "
            f"QC={len(batch_to_qc.get(batch, []))}, "
            f"samples={len(batch_to_samples.get(batch, []))}"
        )

    if len(batch_to_samples) <= 1:
        print("⚠ 只有單一 batch，跳過 QC Batch Scaling（無需跨批次校正）")
        return ProcessingResult(
            file_path=input_file,
            output_path=input_file,
            metabolites=len(data_df),
            samples=len(sample_columns),
            extra={"batches": len(batch_to_samples), "skipped": True, "skip_reason": "single_batch"},
        )

    log_section("執行 QC batch median scaling")
    result_df, invalid_median_counts = scale_dataframe_by_qc_medians(
        data_df, sample_columns, batch_to_qc, batch_to_samples
    )
    summary_df = build_summary_df(source_sheet_name, batch_to_qc, batch_to_samples, invalid_median_counts)
    invalid_total = sum(invalid_median_counts.values())
    print(f"無效 QC median 次數: {invalid_total}")
    if invalid_total:
        for batch in sorted(invalid_median_counts):
            print(f"  - Batch {batch}: {invalid_median_counts[batch]}")

    session_dir = resolve_session_dir(input_file=input_file, session_dir=session_dir)
    timestamp = datetime.now().strftime(DATETIME_FORMAT_FULL)
    if session_dir is not None:
        from metabolomics.utils.file_io import session_output_path, session_plots_dir
        output_file = session_output_path(session_dir, step=3, prefix="QC_Batch_Scaling")
        _plots_dir = str(session_plots_dir(session_dir))
    else:
        output_file = build_output_path("QC_Batch_Scaling", input_file=input_file, timestamp=timestamp)
        _plots_dir = None  # let generate_step3_plots create its own

    log_section("生成 residual 分析圖")
    plots_dir = generate_step3_plots(
        data_df,
        result_df,
        sample_columns,
        sample_info_df,
        input_file,
        timestamp,
        plots_dir=_plots_dir,
    )
    log_section("寫出 Step 3 Excel")
    save_results_to_excel(
        data_df,
        result_df,
        sample_info_df,
        summary_df,
        output_file,
        input_file,
        source_sheet_name,
        sample_type_row=sample_type_row,
    )
    print(f"\n  ✓ QC Batch Scaling 完成 → {os.path.basename(str(output_file))}")

    return ProcessingResult(
        file_path=input_file,
        output_path=str(output_file),
        metabolites=len(result_df),
        samples=len(sample_columns),
        plots_dir=plots_dir,
        extra={"batches": len(batch_to_qc)},
    )
