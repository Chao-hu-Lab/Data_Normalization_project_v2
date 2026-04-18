"""Excel formatting helpers shared across pipeline steps."""

from copy import copy
import math

from openpyxl.styles import Font, PatternFill

from .constants import CV_QUALITY_THRESHOLDS, VALIDATION_THRESHOLDS


PASS_FILL = PatternFill(fgColor="C6EFCE", fill_type="solid")
WARN_FILL = PatternFill(fgColor="FFEB9C", fill_type="solid")
FAIL_FILL = PatternFill(fgColor="FFC7CE", fill_type="solid")
SIG_FILL = PatternFill(fgColor="BDD7EE", fill_type="solid")
NO_FILL = PatternFill(fill_type=None)

HEADER_FILL = PatternFill(fgColor="D9E1F2", fill_type="solid")
SECTION_TITLE_FILL = PatternFill(fgColor="B7D7F0", fill_type="solid")
SECTION_DIVIDER_FILL = PatternFill(fgColor="D9E1F2", fill_type="solid")
SECTION_LABEL_FILL = PatternFill(fgColor="F2F2F2", fill_type="solid")

PASS_FONT_COLOR = "276321"
WARN_FONT_COLOR = "9C5700"
FAIL_FONT_COLOR = "9C0006"
STRUCTURE_FONT_COLOR = "1F3864"

def _coerce_numeric(value):
    try:
        numeric = float(value)
    except (TypeError, ValueError):
        return None
    if not math.isfinite(numeric):
        return None
    return numeric


def _copy_font(cell):
    if cell.has_style and cell.font is not None:
        return copy(cell.font)
    return Font()


def _apply_style(cell, fill=None, font_color=None, bold=None, size=None, alignment=None):
    if fill is not None:
        cell.fill = copy(fill)

    if any(option is not None for option in (font_color, bold, size)):
        font = _copy_font(cell)
        if font_color is not None:
            font.color = font_color
        if bold is not None:
            font.bold = bold
        if size is not None:
            font.size = size
        cell.font = font

    if alignment is not None:
        cell.alignment = copy(alignment)


def _iter_column_cells(worksheet, col_idx, min_row=2, max_row=None):
    final_row = max_row or worksheet.max_row
    for row in worksheet.iter_rows(
        min_row=min_row,
        max_row=final_row,
        min_col=col_idx,
        max_col=col_idx,
    ):
        yield row[0]


def _resolve_status_style(style_key):
    style_map = {
        "pass": (PASS_FILL, PASS_FONT_COLOR),
        "warn": (WARN_FILL, WARN_FONT_COLOR),
        "fail": (FAIL_FILL, FAIL_FONT_COLOR),
        "sig": (SIG_FILL, None),
        "none": (NO_FILL, None),
    }
    return style_map.get(style_key, (None, None))


def apply_header_fill(
    worksheet,
    fill_color="D9E1F2",
    row_idx=1,
    min_col=1,
    max_col=None,
    font_size=11,
):
    """Apply a standard header fill + bold to a worksheet row."""
    final_max_col = max_col or worksheet.max_column
    fill = HEADER_FILL if fill_color == "D9E1F2" else PatternFill(fgColor=fill_color, fill_type="solid")
    for cell in worksheet[row_idx][min_col - 1:final_max_col]:
        _apply_style(cell, fill=fill, bold=True, size=font_size)


def apply_band_fill(
    worksheet,
    col_idx,
    excellent,
    acceptable,
    *,
    min_row=2,
    max_row=None,
    higher_is_better=False,
    use_abs=False,
):
    """Apply PASS/WARN/FAIL colors based on numeric thresholds."""
    for cell in _iter_column_cells(worksheet, col_idx, min_row=min_row, max_row=max_row):
        numeric = _coerce_numeric(cell.value)
        if numeric is None:
            continue
        value = abs(numeric) if use_abs else numeric
        if higher_is_better:
            style_key = "pass" if value >= excellent else "warn" if value >= acceptable else "fail"
        else:
            style_key = "pass" if value < excellent else "warn" if value < acceptable else "fail"
        fill, font_color = _resolve_status_style(style_key)
        _apply_style(cell, fill=fill, font_color=font_color)


def apply_cv_quality_fill(worksheet, col_idx, thresholds=None, min_row=2, max_row=None):
    """Apply PASS/WARN/FAIL styling to CV-like columns."""
    thresholds = thresholds or CV_QUALITY_THRESHOLDS
    apply_band_fill(
        worksheet,
        col_idx,
        excellent=thresholds["excellent"],
        acceptable=thresholds["acceptable"],
        min_row=min_row,
        max_row=max_row,
    )


def apply_improvement_fill(worksheet, col_idx, threshold=5.0, min_row=2, max_row=None):
    """Apply IMPROVE/NEUTRAL/DEGRADE styling to delta metrics."""
    for cell in _iter_column_cells(worksheet, col_idx, min_row=min_row, max_row=max_row):
        numeric = _coerce_numeric(cell.value)
        if numeric is None:
            continue
        if numeric > threshold:
            style_key = "pass"
        elif numeric < -threshold:
            style_key = "fail"
        else:
            style_key = "warn"
        fill, font_color = _resolve_status_style(style_key)
        _apply_style(cell, fill=fill, font_color=font_color)


def apply_significance_fill(worksheet, col_idx, alpha=None, min_row=2, max_row=None):
    """Fill significant p/q values with the shared significance color."""
    alpha = VALIDATION_THRESHOLDS["alpha"] if alpha is None else alpha
    for cell in _iter_column_cells(worksheet, col_idx, min_row=min_row, max_row=max_row):
        numeric = _coerce_numeric(cell.value)
        if numeric is None:
            continue
        if numeric < alpha:
            _apply_style(cell, fill=SIG_FILL)
        else:
            _apply_style(cell, fill=NO_FILL)


def apply_status_fill(worksheet, col_idx, status_map, min_row=2, max_row=None):
    """Apply styles according to the cell text value."""
    for cell in _iter_column_cells(worksheet, col_idx, min_row=min_row, max_row=max_row):
        style_key = status_map.get(cell.value)
        if style_key is None:
            continue
        fill, font_color = _resolve_status_style(style_key)
        _apply_style(cell, fill=fill, font_color=font_color)


def apply_number_format(worksheet, col_idx, fmt, min_row=2, max_row=None):
    """Apply a number format to numeric cells only."""
    for cell in _iter_column_cells(worksheet, col_idx, min_row=min_row, max_row=max_row):
        if _coerce_numeric(cell.value) is not None:
            cell.number_format = fmt
