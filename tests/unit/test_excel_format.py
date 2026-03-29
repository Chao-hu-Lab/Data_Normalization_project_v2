from openpyxl import Workbook

from metabolomics.utils.excel_format import (
    apply_band_fill,
    apply_cv_quality_fill,
    apply_header_fill,
    apply_improvement_fill,
    apply_number_format,
    apply_significance_fill,
    apply_status_fill,
)


def _rgb(cell):
    rgb = cell.fill.fgColor.rgb
    return rgb[-6:] if rgb else None


class TestExcelFormatHelpers:
    def test_apply_cv_quality_fill_uses_pass_warn_fail_thresholds(self):
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "Corrected_QC_CV%"
        ws["A2"] = 10.0
        ws["A3"] = 25.0
        ws["A4"] = 35.0

        apply_cv_quality_fill(ws, 1)

        assert _rgb(ws["A2"]) == "C6EFCE"
        assert _rgb(ws["A3"]) == "FFEB9C"
        assert _rgb(ws["A4"]) == "FFC7CE"

    def test_apply_improvement_fill_marks_positive_neutral_and_negative(self):
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "CV_Improvement%"
        ws["A2"] = 12.0
        ws["A3"] = 2.0
        ws["A4"] = -6.0

        apply_improvement_fill(ws, 1)

        assert _rgb(ws["A2"]) == "C6EFCE"
        assert _rgb(ws["A3"]) == "FFEB9C"
        assert _rgb(ws["A4"]) == "FFC7CE"

    def test_apply_significance_fill_marks_only_significant_cells(self):
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "Wilcoxon_qvalue"
        ws["A2"] = 0.01
        ws["A3"] = 0.05
        ws["A4"] = 0.20

        apply_significance_fill(ws, 1)

        assert _rgb(ws["A2"]) == "BDD7EE"
        assert ws["A3"].fill.fill_type is None
        assert ws["A4"].fill.fill_type is None

    def test_apply_band_fill_supports_absolute_and_higher_is_better_modes(self):
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "Kendall_Tau"
        ws["A2"] = -0.2
        ws["A3"] = 0.4
        ws["A4"] = 0.6
        ws["B1"] = "LOESS_R2"
        ws["B2"] = 0.95
        ws["B3"] = 0.80
        ws["B4"] = 0.60

        apply_band_fill(ws, 1, excellent=0.3, acceptable=0.5, use_abs=True)
        apply_band_fill(ws, 2, excellent=0.9, acceptable=0.7, higher_is_better=True)

        assert _rgb(ws["A2"]) == "C6EFCE"
        assert _rgb(ws["A3"]) == "FFEB9C"
        assert _rgb(ws["A4"]) == "FFC7CE"
        assert _rgb(ws["B2"]) == "C6EFCE"
        assert _rgb(ws["B3"]) == "FFEB9C"
        assert _rgb(ws["B4"]) == "FFC7CE"

    def test_apply_header_status_and_number_format_work_together(self):
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "Status"
        ws["B1"] = "p_value"
        ws["A2"] = "Yes"
        ws["A3"] = "Marginal"
        ws["A4"] = "No"
        ws["B2"] = 0.03125
        ws["B3"] = 0.5

        apply_header_fill(ws)
        apply_status_fill(ws, 1, {"Yes": "pass", "Marginal": "warn", "No": "fail"})
        apply_number_format(ws, 2, "0.0000")

        assert _rgb(ws["A1"]) == "D9E1F2"
        assert ws["A1"].font.bold is True
        assert _rgb(ws["A2"]) == "C6EFCE"
        assert _rgb(ws["A3"]) == "FFEB9C"
        assert _rgb(ws["A4"]) == "FFC7CE"
        assert ws["B2"].number_format == "0.0000"
        assert ws["B3"].number_format == "0.0000"
