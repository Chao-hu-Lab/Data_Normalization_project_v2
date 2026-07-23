"""
Tests for QC_LOWESS_v2 module (Step 2).

These tests verify:
1. Input validation (requires Step 1 output format)
2. Output structure (correct sheets, columns)
3. LOWESS correction logic
4. Return value format
"""
import pytest
import pandas as pd
import numpy as np
import os
import shutil
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from metabolomics.utils.constants import SHEET_NAMES
from metabolomics.utils.excel_colors import cell_has_red_font
from metabolomics.utils.results import WorkflowOutcome


@pytest.fixture
def lowess_ready_input_file(sample_input_file, tmp_path):
    """Copy the canonical workbook and make pooled QCs shared across all batches."""
    workbook_path = tmp_path / "lowess_ready_input.xlsx"
    shutil.copy2(sample_input_file, workbook_path)
    workbook = load_workbook(workbook_path)
    try:
        worksheet = workbook[SHEET_NAMES["sample_info"]]
        headers = {cell.value: cell.column for cell in worksheet[1]}
        sample_type_col = headers["Sample_Type"]
        batch_col = headers["Batch"]
        batch_names = sorted({
            str(worksheet.cell(row=row, column=batch_col).value).strip()
            for row in range(2, worksheet.max_row + 1)
            if worksheet.cell(row=row, column=batch_col).value not in (None, "")
        })
        shared_batches = ";".join(batch_names)
        for row in range(2, worksheet.max_row + 1):
            sample_type = str(worksheet.cell(row=row, column=sample_type_col).value).upper()
            if "QC" in sample_type:
                worksheet.cell(row=row, column=batch_col).value = shared_batches
        workbook.save(workbook_path)
    finally:
        workbook.close()
    return str(workbook_path)


class TestQCLOWESSInput:
    """Tests for input validation."""

    def test_module_loads(self, qc_lowess_module):
        """Test that module can be imported."""
        assert qc_lowess_module is not None
        assert hasattr(qc_lowess_module, 'main')

    def test_has_required_functions(self, qc_lowess_module):
        """Test that module has expected functions."""
        expected_functions = ['main', 'load_and_process_data', 'get_valid_values']
        for func_name in expected_functions:
            assert hasattr(qc_lowess_module, func_name), f"Missing function: {func_name}"

    def test_step2_batch_parser_keeps_slash_as_one_label(self, qc_lowess_module):
        assert qc_lowess_module.parse_batch_labels("A/B") == ["A/B"]

    def test_load_and_process_data_reports_actual_source_sheet(
        self,
        qc_lowess_module,
        sample_input_file,
        capsys,
    ):
        """Fallback logging should mention RawIntensity when no ISTD sheet exists."""
        qc_lowess_module.load_and_process_data(sample_input_file)

        captured = capsys.readouterr().out
        assert "成功讀取 'RawIntensity' 工作表" in captured
        assert "成功讀取 'ISTD_Correction' 工作表" not in captured

    def test_load_accepts_pathlike_file(self, qc_lowess_module, sample_input_file):
        _, istd_df, sample_info_df, _ = qc_lowess_module.load_and_process_data(Path(sample_input_file))

        assert istd_df is not None
        assert sample_info_df is not None

    @pytest.mark.parametrize("batch_value", [pytest.param(None, id="missing-column"), np.nan, "", "   "])
    def test_load_fails_when_batch_metadata_is_missing(
        self,
        qc_lowess_module,
        tmp_path,
        batch_value,
    ):
        sample_names = ["QC1", "QC2", "QC3", "QC4", "QC5"]
        sample_info = {
            "Sample_Name": sample_names,
            "Sample_Type": ["QC"] * len(sample_names),
            "Injection_Order": range(1, len(sample_names) + 1),
        }
        if batch_value is not None:
            sample_info["Batch"] = ["A", "A", batch_value, "A", "A"]

        workbook_path = tmp_path / "step2_missing_batch.xlsx"
        with pd.ExcelWriter(workbook_path, engine="openpyxl") as writer:
            pd.DataFrame(
                {
                    "Mz/RT": ["100.1/1.0"],
                    **{name: [100.0] for name in sample_names},
                }
            ).to_excel(writer, sheet_name="RawIntensity", index=False)
            pd.DataFrame(sample_info).to_excel(writer, sheet_name="SampleInfo", index=False)

        with pytest.raises(ValueError, match="請先補齊 Batch"):
            qc_lowess_module.load_and_process_data(workbook_path)

    def test_load_fails_when_sample_info_has_no_qc(
        self,
        qc_lowess_module,
        tmp_path,
    ):
        workbook_path = tmp_path / "step2_no_qc.xlsx"
        with pd.ExcelWriter(workbook_path, engine="openpyxl") as writer:
            pd.DataFrame(
                {
                    "Mz/RT": ["100.1/1.0"],
                    "Sample1": [100.0],
                }
            ).to_excel(writer, sheet_name="RawIntensity", index=False)
            pd.DataFrame(
                {
                    "Sample_Name": ["Sample1"],
                    "Sample_Type": ["Exposure"],
                    "Injection_Order": [1],
                    "Batch": ["A"],
                }
            ).to_excel(writer, sheet_name="SampleInfo", index=False)

        with pytest.raises(ValueError, match="未找到 QC 樣本"):
            qc_lowess_module.load_and_process_data(workbook_path)

    def test_perform_fails_when_batch_column_is_missing(self, qc_lowess_module):
        sample_names = ["QC1", "QC2", "QC3", "QC4", "QC5"]
        source_df = pd.DataFrame(
            {
                "FeatureID": ["100.1/1.0"],
                **{name: [100.0] for name in sample_names},
            }
        )
        source_df.attrs["sample_columns"] = sample_names
        sample_info_df = pd.DataFrame(
            {
                "Sample_Name": sample_names,
                "Sample_Type": ["QC"] * len(sample_names),
                "Injection_Order": range(1, len(sample_names) + 1),
            }
        )

        with pytest.raises(ValueError, match="請先補齊 Batch"):
            qc_lowess_module.perform_lowess_normalization(source_df, sample_info_df)

    def test_perform_keeps_values_when_all_features_have_four_qc(
        self,
        qc_lowess_module,
    ):
        sample_names = ["QC1", "QC2", "QC3", "QC4", "Sample1"]
        source_df = pd.DataFrame(
            {
                "FeatureID": ["100.1/1.0"],
                "QC1": [100.0],
                "QC2": [110.0],
                "QC3": [120.0],
                "QC4": [130.0],
                "Sample1": [125.0],
            }
        )
        source_df.attrs["sample_columns"] = sample_names
        sample_info_df = pd.DataFrame(
            {
                "Sample_Name": sample_names,
                "Sample_Type": ["QC", "QC", "QC", "QC", "Exposure"],
                "Injection_Order": range(1, len(sample_names) + 1),
                "Batch": ["A"] * len(sample_names),
            }
        )

        lowess_df, _, _, trend_stats_df, decision_stats, _ = (
            qc_lowess_module.perform_lowess_normalization(source_df, sample_info_df)
        )

        assert lowess_df.loc[0, sample_names].tolist() == source_df.loc[0, sample_names].tolist()
        assert trend_stats_df.loc[0, "Decision_Status"] == "insufficient_qc"
        assert decision_stats["event_counts"]["insufficient_qc"] == 1

    def test_perform_applies_only_batches_with_enough_valid_qc(
        self,
        qc_lowess_module,
    ):
        a_qc = [f"A_QC{i}" for i in range(1, 5)]
        b_qc = [f"B_QC{i}" for i in range(1, 9)]
        sample_names = [*a_qc, "SampleA", *b_qc, "SampleB"]
        intensities = {
            **{name: [1000.0 + 50.0 * index] for index, name in enumerate(a_qc, 1)},
            "SampleA": [1250.0],
            **{name: [2000.0 + 100.0 * index] for index, name in enumerate(b_qc, 1)},
            "SampleB": [2450.0],
        }
        source_df = pd.DataFrame({"FeatureID": ["100.1/1.0"], **intensities})
        source_df.attrs["sample_columns"] = sample_names
        sample_info_df = pd.DataFrame(
            {
                "Sample_Name": sample_names,
                "Sample_Type": [
                    *(["QC"] * len(a_qc)),
                    "Exposure",
                    *(["QC"] * len(b_qc)),
                    "Control",
                ],
                "Injection_Order": [1, 2, 3, 4, 3.5, 6, 7, 8, 9, 10, 11, 12, 13, 7.5],
                "Batch": [
                    *(["A"] * (len(a_qc) + 1)),
                    *(["B"] * (len(b_qc) + 1)),
                ],
            }
        )

        lowess_df, _, _, trend_stats_df, decision_stats, _ = (
            qc_lowess_module.perform_lowess_normalization(source_df, sample_info_df)
        )

        assert lowess_df.loc[0, "SampleA"] == pytest.approx(source_df.loc[0, "SampleA"])
        assert lowess_df.loc[0, "SampleB"] != pytest.approx(source_df.loc[0, "SampleB"])
        assert trend_stats_df.loc[0, "Decision_Status"] == "partial_success"
        assert trend_stats_df.loc[0, "Fit_Strategy"] == "batch_local_lowess"
        batch_detail = trend_stats_df.loc[0, "Batch_Decision_Detail"]
        assert "A:status=insufficient_qc,fit=unknown,valid_qc=4" in batch_detail
        assert "B:status=success,fit=batch_local_lowess,valid_qc=8" in batch_detail
        assert decision_stats["per_batch"]["A"]["insufficient_qc"] == 1
        assert decision_stats["per_batch"]["B"]["success"] == 1

    def test_perform_reports_log_linear_fallback_per_feature(
        self,
        qc_lowess_module,
    ):
        sample_names = [f"Run{order}" for order in range(1, 14)]
        qc_orders = {1, 3, 5, 7, 9, 11, 13}
        orders = np.arange(1, 14, dtype=float)
        values = 1000.0 * 2.0 ** (0.08 * (orders - 7.0))
        source_df = pd.DataFrame(
            {"FeatureID": ["100.1/1.0"], **{
                name: [value]
                for name, value in zip(sample_names, values, strict=True)
            }}
        )
        source_df.attrs["sample_columns"] = sample_names
        sample_info_df = pd.DataFrame(
            {
                "Sample_Name": sample_names,
                "Sample_Type": [
                    "QC" if order in qc_orders else "Exposure"
                    for order in range(1, 14)
                ],
                "Injection_Order": orders,
                "Batch": ["A"] * len(sample_names),
            }
        )

        corrected, _, _, summary, decision_stats, _ = (
            qc_lowess_module.perform_lowess_normalization(
                source_df,
                sample_info_df,
            )
        )

        assert corrected.loc[0, "Run2"] != pytest.approx(source_df.loc[0, "Run2"])
        assert summary.loc[0, "Decision_Status"] == "success"
        assert summary.loc[0, "Fit_Strategy"] == "log_linear_fallback"
        assert summary.loc[0, "Linear_LOOCV_Gain"] >= 0.10
        assert np.isfinite(summary.loc[0, "Linear_R2"])
        assert np.isfinite(summary.loc[0, "Linear_Residual_RMSE_Log2"])
        assert np.isfinite(summary.loc[0, "Linear_LOOCV_RMSE_Log2"])
        assert summary.loc[0, "Linear_Shrinkage_Factor"] == pytest.approx(0.5)
        assert np.isnan(summary.loc[0, "LOOCV_RMSE"])
        assert np.isnan(summary.loc[0, "LOESS_RMSE"])
        assert decision_stats["event_counts"]["success"] == 1

    def test_perform_does_not_fit_against_temporary_injection_orders(
        self,
        qc_lowess_module,
    ):
        sample_names = [f"Run{order}" for order in range(1, 14)]
        qc_orders = {1, 3, 5, 7, 9, 11, 13}
        orders = np.arange(1, 14, dtype=float)
        values = 1000.0 * 2.0 ** (0.08 * (orders - 7.0))
        source_df = pd.DataFrame(
            {"FeatureID": ["100.1/1.0"], **{
                name: [value]
                for name, value in zip(sample_names, values, strict=True)
            }}
        )
        source_df.attrs["sample_columns"] = sample_names
        observed_orders = orders.copy()
        observed_orders[5] = np.nan
        sample_info_df = pd.DataFrame(
            {
                "Sample_Name": sample_names,
                "Sample_Type": [
                    "QC" if order in qc_orders else "Exposure"
                    for order in range(1, 14)
                ],
                "Injection_Order": observed_orders,
                "Batch": ["A"] * len(sample_names),
            }
        )

        corrected, _, _, summary, decision_stats, _ = (
            qc_lowess_module.perform_lowess_normalization(source_df, sample_info_df)
        )

        assert corrected.loc[0, sample_names].tolist() == pytest.approx(values)
        assert summary.loc[0, "Decision_Status"] == "invalid_injection_order"
        assert decision_stats["event_counts"]["invalid_injection_order"] == 1

    def test_main_skips_when_every_feature_batch_has_insufficient_qc(
        self,
        qc_lowess_module,
        tmp_path,
    ):
        sample_names = ["QC1", "QC2", "QC3", "QC4", "Sample1"]
        workbook_path = tmp_path / "step2_all_insufficient_qc.xlsx"
        with pd.ExcelWriter(workbook_path, engine="openpyxl") as writer:
            pd.DataFrame(
                {
                    "Mz/RT": ["100.1/1.0"],
                    "QC1": [100.0],
                    "QC2": [110.0],
                    "QC3": [120.0],
                    "QC4": [130.0],
                    "Sample1": [125.0],
                }
            ).to_excel(writer, sheet_name="RawIntensity", index=False)
            pd.DataFrame(
                {
                    "Sample_Name": sample_names,
                    "Sample_Type": ["QC", "QC", "QC", "QC", "Exposure"],
                    "Injection_Order": range(1, len(sample_names) + 1),
                    "Batch": ["A"] * len(sample_names),
                }
            ).to_excel(writer, sheet_name="SampleInfo", index=False)

        result = qc_lowess_module.main(input_file=workbook_path)

        assert result.status is WorkflowOutcome.SKIPPED
        assert result.reason == "insufficient_valid_qc_for_correction"
        assert Path(result.output_path) == workbook_path

    def test_main_succeeds_with_gated_seven_qc_log_linear_fallback(
        self,
        qc_lowess_module,
        tmp_path,
    ):
        sample_names = [f"Run{order}" for order in range(1, 14)]
        qc_orders = {1, 3, 5, 7, 9, 11, 13}
        orders = np.arange(1, 14, dtype=float)
        values = 1000.0 * 2.0 ** (0.08 * (orders - 7.0))
        workbook_path = tmp_path / "seven_qc_log_linear.xlsx"
        with pd.ExcelWriter(workbook_path, engine="openpyxl") as writer:
            pd.DataFrame(
                {"Mz/RT": ["100.1/1.0"], **{
                    name: [value]
                    for name, value in zip(sample_names, values, strict=True)
                }}
            ).to_excel(writer, sheet_name="RawIntensity", index=False)
            pd.DataFrame(
                {
                    "Sample_Name": sample_names,
                    "Sample_Type": [
                        "QC" if order in qc_orders else "Exposure"
                        for order in range(1, 14)
                    ],
                    "Injection_Order": orders,
                    "Batch": ["A"] * len(sample_names),
                }
            ).to_excel(writer, sheet_name="SampleInfo", index=False)

        result = qc_lowess_module.main(input_file=workbook_path)

        assert result.status is WorkflowOutcome.SUCCEEDED
        summary = pd.read_excel(result.output_path, sheet_name="LOESS_summary")
        assert summary.loc[0, "Decision_Status"] == "success"
        assert summary.loc[0, "Fit_Strategy"] == "log_linear_fallback"

    def test_main_skips_when_no_feature_correction_is_applied(
        self,
        qc_lowess_module,
        tmp_path,
    ):
        qc_names = [f"QC{index}" for index in range(1, 9)]
        sample_names = [*qc_names, "Sample1"]
        workbook_path = tmp_path / "step2_no_applied_correction.xlsx"
        with pd.ExcelWriter(workbook_path, engine="openpyxl") as writer:
            pd.DataFrame(
                {
                    "Mz/RT": ["100.1/1.0"],
                    **dict(
                        zip(
                            qc_names,
                            [[value] for value in [100, 99, 101, 100, 100, 101, 99, 100]],
                            strict=True,
                        )
                    ),
                    "Sample1": [100.0],
                }
            ).to_excel(writer, sheet_name="RawIntensity", index=False)
            pd.DataFrame(
                {
                    "Sample_Name": sample_names,
                    "Sample_Type": [*(["QC"] * 8), "Exposure"],
                    "Injection_Order": range(1, len(sample_names) + 1),
                    "Batch": ["A"] * len(sample_names),
                }
            ).to_excel(writer, sheet_name="SampleInfo", index=False)

        result = qc_lowess_module.main(input_file=workbook_path)

        assert result.status is WorkflowOutcome.SKIPPED
        assert result.reason == "no_feature_correction_applied"
        assert Path(result.output_path) == workbook_path

    def test_load_validates_sample_info_before_resolving_source(
        self,
        qc_lowess_module,
        tmp_path,
    ):
        workbook_path = tmp_path / "invalid_sample_info_and_missing_source.xlsx"
        with pd.ExcelWriter(workbook_path, engine="openpyxl") as writer:
            pd.DataFrame({"Unrelated": ["value"]}).to_excel(
                writer,
                sheet_name="SampleInfo",
                index=False,
            )

        with pytest.raises(ValueError, match="Step 2 SampleInfo validation failed"):
            qc_lowess_module.load_and_process_data(workbook_path)

    def test_load_preserves_missing_source_error(self, qc_lowess_module, tmp_path):
        workbook_path = tmp_path / "missing_source.xlsx"
        sample_names = [f"QC_{index}" for index in range(1, 6)]
        with pd.ExcelWriter(workbook_path, engine="openpyxl") as writer:
            pd.DataFrame(
                {
                    "Sample_Name": sample_names,
                    "Sample_Type": ["QC"] * 5,
                    "Injection_Order": range(1, 6),
                    "Batch": ["A"] * 5,
                }
            ).to_excel(writer, sheet_name="SampleInfo", index=False)

        with pytest.raises(ValueError, match="Worksheet named 'RawIntensity' not found"):
            qc_lowess_module.load_and_process_data(workbook_path)

    def test_raw_fallback_returns_independent_original_snapshot(
        self,
        qc_lowess_module,
        tmp_path,
    ):
        workbook_path = tmp_path / "raw_fallback_snapshot.xlsx"
        sample_names = [f"QC_{index}" for index in range(1, 6)]
        raw_df = pd.DataFrame(
            {
                "Mz/RT": ["Sample_Type", "100.1/1.0"],
                **{name: ["QC", float(index)] for index, name in enumerate(sample_names, 1)},
            }
        )
        sample_info_df = pd.DataFrame(
            {
                "Sample_Name": sample_names,
                "Sample_Type": ["QC"] * 5,
                "Injection_Order": range(1, 6),
                "Batch": ["A"] * 5,
            }
        )
        with pd.ExcelWriter(workbook_path, engine="openpyxl") as writer:
            raw_df.to_excel(writer, sheet_name="RawIntensity", index=False)
            sample_info_df.to_excel(writer, sheet_name="SampleInfo", index=False)

        loaded_raw, source_df, _, sample_type_row = (
            qc_lowess_module.load_and_process_data(workbook_path)
        )

        pd.testing.assert_frame_equal(loaded_raw, raw_df)
        assert loaded_raw is not source_df
        assert "Mz/RT" in loaded_raw.columns
        assert "FeatureID" in source_df.columns
        assert source_df["FeatureID"].tolist() == ["100.1/1.0"]
        assert sample_type_row is not None

    def test_load_fails_closed_when_no_sample_columns_match_sampleinfo(self, qc_lowess_module, tmp_path):
        raw_df = pd.DataFrame(
            {
                "Mz/RT": ["100.1/1.0"],
                "Unmapped_A": [10.0],
                "Unmapped_B": [20.0],
            }
        )
        sample_info_df = pd.DataFrame(
            {
                "Sample_Name": ["Sample_A", "Sample_B"],
                "Sample_Type": ["QC", "QC"],
                "Injection_Order": [1, 2],
                "Batch": ["A", "A"],
            }
        )
        workbook_path = tmp_path / "unmapped_step2.xlsx"
        with pd.ExcelWriter(workbook_path, engine="openpyxl") as writer:
            raw_df.to_excel(writer, sheet_name="RawIntensity", index=False)
            sample_info_df.to_excel(writer, sheet_name="SampleInfo", index=False)

        with pytest.raises(ValueError, match="找不到任何可與 SampleInfo 對齊的樣本欄位"):
            qc_lowess_module.load_and_process_data(workbook_path)

    def test_collect_red_marked_feature_ids_accepts_argb_red(self, qc_lowess_module, tmp_path):
        """Fallback red-font parsing should accept ARGB red strings from openpyxl."""
        raw_df = pd.DataFrame(
            {
                "Mz/RT": ["Sample_Type", "100.1/1.0", "200.2/2.0"],
                "QC_1": ["QC", 10.0, 20.0],
                "QC_2": ["QC", 11.0, 21.0],
                "QC_3": ["QC", 9.0, 19.0],
            }
        )
        sample_info_df = pd.DataFrame(
            {
                "Sample_Name": ["QC_1", "QC_2", "QC_3"],
                "Sample_Type": ["QC", "QC", "QC"],
            }
        )
        workbook_path = tmp_path / "argb_red_fallback.xlsx"
        with pd.ExcelWriter(workbook_path, engine="openpyxl") as writer:
            raw_df.to_excel(writer, sheet_name="RawIntensity", index=False)
            sample_info_df.to_excel(writer, sheet_name="SampleInfo", index=False)

        workbook = load_workbook(workbook_path)
        try:
            sheet = workbook["RawIntensity"]
            sheet["A3"].font = Font(color="00FF0000")
            workbook.save(workbook_path)
        finally:
            workbook.close()

        detected = qc_lowess_module.collect_red_marked_feature_ids(
            str(workbook_path),
            SHEET_NAMES["raw_intensity"],
        )

        assert detected == {"100.1/1.0"}

    def test_main_raises_when_results_cannot_be_saved(
        self,
        qc_lowess_module,
        tmp_path,
        monkeypatch,
    ):
        input_path = tmp_path / "input.xlsx"
        output_path = tmp_path / "output.xlsx"
        plots_path = tmp_path / "plots"
        raw_df = pd.DataFrame(
            {"FeatureID": ["F1"], "QC_1": [1.0], "Sample_1": [2.0]}
        )
        sample_info_df = pd.DataFrame(
            {
                "Sample_Name": ["QC_1", "Sample_1"],
                "Sample_Type": ["QC", "Exposure"],
                "Batch": ["A", "A"],
                "Injection_Order": [1, 2],
            }
        )

        monkeypatch.setattr(qc_lowess_module, "resolve_session_dir", lambda **_kwargs: None)
        monkeypatch.setattr(qc_lowess_module, "get_output_root", lambda **_kwargs: str(tmp_path))
        monkeypatch.setattr(
            qc_lowess_module,
            "load_and_process_data",
            lambda _path: (raw_df, raw_df, sample_info_df, None),
        )
        monkeypatch.setattr(
            qc_lowess_module,
            "perform_lowess_normalization",
            lambda *_args: (raw_df, ["QC_1"], {}, pd.DataFrame(), {}, {}),
        )
        monkeypatch.setattr(
            qc_lowess_module,
            "build_output_path",
            lambda *_args, **_kwargs: output_path,
        )
        monkeypatch.setattr(
            qc_lowess_module,
            "build_plots_dir",
            lambda *_args, **_kwargs: plots_path,
        )
        monkeypatch.setattr(qc_lowess_module, "save_results_to_excel", lambda *_args, **_kwargs: False)

        with pytest.raises(RuntimeError, match="Failed to save QC-LOESS results"):
            qc_lowess_module.main(input_file=str(input_path))


class TestQCLOWESSOutput:
    """Tests for output validation."""

    @pytest.mark.slow
    @pytest.mark.integration
    def test_canonical_workbook_succeeds_when_any_feature_is_corrected(
        self,
        qc_lowess_module,
        sample_input_file,
    ):
        result = qc_lowess_module.main(input_file=sample_input_file)

        assert result.status is WorkflowOutcome.SUCCEEDED
        assert Path(result.output_path) != Path(sample_input_file)
        summary = pd.read_excel(result.output_path, sheet_name="LOESS_summary")
        assert "Decision_Status" in summary.columns
        assert summary["Decision_Status"].isin(
            {"success", "partial_success"}
        ).any()

    @pytest.mark.slow
    @pytest.mark.integration
    def test_main_falls_back_to_raw_intensity_when_istd_sheet_is_missing(
        self,
        qc_lowess_module,
        lowess_ready_input_file,
        workbook_sheet_names,
    ):
        """Step 2 should accept a workbook that only has RawIntensity and SampleInfo."""
        step2_result = qc_lowess_module.main(input_file=lowess_ready_input_file)

        validation_target = (
            step2_result.output_path
            if hasattr(step2_result, "output_path")
            else step2_result.get("output_path")
        )

        assert set(workbook_sheet_names(validation_target)) == {
            "RawIntensity",
            SHEET_NAMES["qc_lowess"],
            SHEET_NAMES["qc_lowess_advanced"],
            "SampleInfo",
            SHEET_NAMES["design_identifiability"],
        }

    @pytest.mark.slow
    @pytest.mark.integration
    def test_fallback_excludes_red_marked_istds_from_qc_lowess_result(
        self,
        istd_module,
        qc_lowess_module,
        lowess_ready_input_file,
    ):
        """When Step 1 is skipped, red-marked ISTDs should not re-enter downstream result sheets."""
        from metabolomics.utils.data_helpers import extract_sample_type_row

        raw_df, _, _, _ = istd_module.load_and_process_data(lowess_ready_input_file)
        raw_df, _ = extract_sample_type_row(raw_df, "FeatureID")
        expected_rows = len(raw_df[~raw_df["is_ISTD"]])

        step2_result = qc_lowess_module.main(input_file=lowess_ready_input_file)
        step2_output = (
            step2_result.output_path
            if hasattr(step2_result, "output_path")
            else step2_result.get("output_path")
        )

        output_df = pd.read_excel(step2_output, sheet_name=SHEET_NAMES["qc_lowess"])
        output_df, _ = extract_sample_type_row(output_df, output_df.columns[0])

        assert len(output_df) == expected_rows

    @pytest.mark.slow
    @pytest.mark.integration
    def test_main_with_step1_output(self, istd_module, qc_lowess_module,
                                     lowess_ready_input_file, validate_processing_result):
        """Test QC-LOWESS with Step 1 output."""
        # First run Step 1
        step1_result = istd_module.main(input_file=lowess_ready_input_file)
        assert step1_result is not None, "Step 1 should succeed"
        step1_output = step1_result.output_path if hasattr(step1_result, "output_path") else step1_result.get('output_path')
        assert step1_output

        # Then run Step 2
        step2_result = qc_lowess_module.main(input_file=step1_output)

        validation = validate_processing_result(
            step2_result,
            required_keys=['output_path', 'metabolites', 'samples']
        )

        assert validation['is_processing_result'], f"Result should be ProcessingResult: {validation['errors']}"

    @pytest.mark.slow
    @pytest.mark.integration
    def test_output_file_structure(self, istd_module, qc_lowess_module,
                                     lowess_ready_input_file, validate_excel_output):
        """Test output Excel file structure."""
        # Run Step 1
        step1_result = istd_module.main(input_file=lowess_ready_input_file)
        step1_output = step1_result.output_path if hasattr(step1_result, "output_path") else step1_result.get('output_path')

        # Run Step 2
        step2_result = qc_lowess_module.main(input_file=step1_output)

        assert step2_result is not None
        step2_output = step2_result.output_path if hasattr(step2_result, "output_path") else step2_result.get('output_path')
        assert step2_output

        validation = validate_excel_output(
            step2_output,
            required_sheets=[SHEET_NAMES["qc_lowess"]],
            min_rows=1
        )

        assert validation['exists'], f"Output file should exist"
        assert validation['readable'], f"Output file should be readable"

    @pytest.mark.slow
    @pytest.mark.integration
    def test_qc_lowess_preserves_current_step4_handoff_without_imputing(
        self,
        qc_lowess_module,
        lowess_ready_input_file,
        output_dir,
    ):
        input_path = os.path.join(output_dir, "qc_lowess_step4_handoff.xlsx")
        if os.path.exists(input_path):
            os.remove(input_path)
        shutil.copy2(lowess_ready_input_file, input_path)

        workbook = load_workbook(input_path)
        try:
            raw = workbook[SHEET_NAMES["raw_intensity"]]
            sample_info = workbook[SHEET_NAMES["sample_info"]]
            sample_info_headers = {
                cell.value: cell.column for cell in sample_info[1]
            }
            sample_name_col = sample_info_headers["Sample_Name"]
            expected_sample_count = sum(
                sample_info.cell(row=row, column=sample_name_col).value
                not in (None, "")
                for row in range(2, sample_info.max_row + 1)
            )

            study_columns = [
                column
                for column in range(2, raw.max_column + 1)
                if str(raw.cell(row=2, column=column).value).upper() != "QC"
            ]
            analyte_rows = [
                row
                for row in range(3, raw.max_row + 1)
                if not cell_has_red_font(raw.cell(row=row, column=1))
            ]

            metadata_headers = [
                "Cohort_7_ratio",
                "RareSubtype_ratio",
                "QC_ratio",
                "is_Presence_Absence_Marker",
                "Feature_Filter_Keep_Reasons",
                "Imputation_Tag_Reasons",
            ]
            metadata_type_row = [
                "na",
                "na",
                "na",
                "is_Presence_Absence_Marker",
                "Feature_Filter_Keep_Reasons",
                "Imputation_Tag_Reasons",
            ]
            patterns = [
                (0.20, 0.20, 0.125, False, "stable", ""),
                (0.30, 0.15, 0.0, True, "ratio_rescue", "low_overall_detection"),
                (0.20, 0.05, 0.0, True, "mnar", "low_overall_detection"),
            ]
            missing_cells = []
            seen_markers = set()
            for row in analyte_rows:
                marker = patterns[(row - 3) % len(patterns)][3]
                if marker in seen_markers:
                    continue
                study_column = study_columns[len(seen_markers) % len(study_columns)]
                raw.cell(row=row, column=study_column).value = None
                missing_cells.append(
                    (
                        raw.cell(row=row, column=1).value,
                        raw.cell(row=1, column=study_column).value,
                    )
                )
                seen_markers.add(marker)
            assert seen_markers == {False, True}

            expected_metadata = {}
            metadata_start = raw.max_column + 1
            for offset, (header, type_value) in enumerate(
                zip(metadata_headers, metadata_type_row, strict=True),
            ):
                column = metadata_start + offset
                raw.cell(row=1, column=column, value=header)
                raw.cell(row=2, column=column, value=type_value)

            for row in range(3, raw.max_row + 1):
                feature_id = raw.cell(row=row, column=1).value
                pattern = patterns[(row - 3) % len(patterns)]
                if row in analyte_rows:
                    expected_metadata[feature_id] = dict(
                        zip(metadata_headers, pattern, strict=True)
                    )
                for offset, value in enumerate(pattern):
                    raw.cell(row=row, column=metadata_start + offset, value=value)
            workbook.save(input_path)
        finally:
            workbook.close()

        result = qc_lowess_module.main(input_file=input_path)
        assert result.samples == expected_sample_count
        result_wb = load_workbook(result.output_path, read_only=True, data_only=True)
        try:
            output = result_wb[SHEET_NAMES["qc_lowess"]]
            output_headers = {
                cell.value: cell.column for cell in output[1]
            }
            assert "Imputation_Strategy" not in output_headers
            assert set(metadata_headers) <= set(output_headers)

            output_rows = {
                output.cell(row=row, column=1).value: row
                for row in range(3, output.max_row + 1)
            }
            for feature_id, expected in expected_metadata.items():
                output_row = output_rows[feature_id]
                for header, expected_value in expected.items():
                    actual = output.cell(
                        row=output_row,
                        column=output_headers[header],
                    ).value
                    if expected_value == "":
                        actual = actual or ""
                    assert actual == expected_value

            for feature_id, sample_name in missing_cells:
                assert output.cell(
                    row=output_rows[feature_id],
                    column=output_headers[sample_name],
                ).value is None
        finally:
            result_wb.close()

    @pytest.mark.slow
    @pytest.mark.integration
    def test_cv_improvement_tracking(self, istd_module, qc_lowess_module, lowess_ready_input_file):
        """Test that CV improvement is tracked in output."""
        # Run Step 1
        step1_result = istd_module.main(input_file=lowess_ready_input_file)
        step1_output = step1_result.output_path if hasattr(step1_result, "output_path") else step1_result.get('output_path')

        # Run Step 2
        step2_result = qc_lowess_module.main(input_file=step1_output)

        # Check output contains CV statistics
        step2_output = step2_result.output_path if hasattr(step2_result, "output_path") else step2_result.get('output_path')
        output_df = pd.read_excel(step2_output, sheet_name=SHEET_NAMES["qc_lowess"])

        # Should have CV-related columns
        cv_columns = [col for col in output_df.columns if 'CV' in col.upper()]
        assert len(cv_columns) > 0, "Output should have CV-related columns"
        assert {'Original_QC_CV%', 'Corrected_QC_CV%'}.issubset(output_df.columns)

        numeric = output_df[['Original_QC_CV%', 'Corrected_QC_CV%']].apply(
            pd.to_numeric, errors='coerce'
        )
        cv_diff = (numeric['Original_QC_CV%'] - numeric['Corrected_QC_CV%']).dropna()
        assert not cv_diff.empty, "CV comparison should contain numeric values"
        assert (cv_diff.abs() > 1e-9).any(), "QC correction should change at least one feature CV"

    @pytest.mark.slow
    @pytest.mark.integration
    def test_plots_are_generated(self, istd_module, qc_lowess_module, lowess_ready_input_file):
        """Test that QC-LOWESS produces plot artifacts."""
        step1_result = istd_module.main(input_file=lowess_ready_input_file)
        step1_output = step1_result.output_path if hasattr(step1_result, "output_path") else step1_result.get('output_path')

        step2_result = qc_lowess_module.main(input_file=step1_output)

        plots_dir = step2_result.plots_dir if hasattr(step2_result, "plots_dir") else step2_result.get('plots_dir')
        assert plots_dir, "QC-LOWESS should report a plots directory"

        from pathlib import Path

        plot_files = sorted(Path(plots_dir).glob("*.png"))
        assert plot_files, "QC-LOWESS should generate PNG plots"

    @pytest.mark.slow
    @pytest.mark.integration
    def test_output_workbook_only_keeps_required_sheets(
        self,
        istd_module,
        qc_lowess_module,
        lowess_ready_input_file,
        copy_workbook_with_extra_sheet,
        workbook_sheet_names,
    ):
        """Step 2 output should only keep the previous step sheet and required metadata."""
        step1_result = istd_module.main(input_file=lowess_ready_input_file)
        step1_output = step1_result.output_path if hasattr(step1_result, "output_path") else step1_result.get('output_path')
        step1_with_extra_sheet = copy_workbook_with_extra_sheet(step1_output)

        step2_result = qc_lowess_module.main(input_file=step1_with_extra_sheet)
        step2_output = step2_result.output_path if hasattr(step2_result, "output_path") else step2_result.get('output_path')

        expected_source_sheet = 'RawIntensity'
        if 'ISTD_Correction' in workbook_sheet_names(step1_with_extra_sheet):
            expected_source_sheet = 'ISTD_Correction'

        assert set(workbook_sheet_names(step2_output)) == {
            expected_source_sheet,
            SHEET_NAMES["qc_lowess"],
            SHEET_NAMES["qc_lowess_advanced"],
            'SampleInfo',
            SHEET_NAMES["design_identifiability"],
        }

    @pytest.mark.slow
    @pytest.mark.integration
    def test_loess_summary_sheet_keeps_feature_table_and_adds_summary_block(
        self,
        qc_lowess_module,
        lowess_ready_input_file,
    ):
        from openpyxl import load_workbook

        step2_result = qc_lowess_module.main(input_file=lowess_ready_input_file)
        step2_output = (
            step2_result.output_path
            if hasattr(step2_result, "output_path")
            else step2_result.get("output_path")
        )

        workbook = load_workbook(step2_output, read_only=False)
        try:
            worksheet = workbook[SHEET_NAMES["qc_lowess_advanced"]]
            assert worksheet["A1"].value == "Mz/RT"
            summary_col_idx = next(
                cell.column
                for cell in worksheet[1]
                if cell.value == "LOESS Summary"
            )
            summary_col_letter = get_column_letter(summary_col_idx)
            value_col_idx = summary_col_idx + 1
            assert worksheet.cell(row=1, column=summary_col_idx).value == "LOESS Summary"
            summary_labels = [worksheet[f"{summary_col_letter}{row_idx}"].value for row_idx in range(1, worksheet.max_row + 1)]
            assert "Overview" in summary_labels
            assert "Features processed" in summary_labels
            assert "Overall readout" in summary_labels

            features_row = next(
                row_idx
                for row_idx, label in enumerate(summary_labels, start=1)
                if label == "Features processed"
            )
            assert worksheet.cell(row=features_row, column=value_col_idx).value is not None
        finally:
            workbook.close()


    @pytest.mark.slow
    def test_main_writes_to_session_dir(self, qc_lowess_module, lowess_ready_input_file, tmp_path):
        """When session_dir is provided, output goes into that directory."""
        from pathlib import Path
        from metabolomics.utils.file_io import create_session_dir

        session = create_session_dir(output_root=tmp_path)
        result = qc_lowess_module.main(input_file=lowess_ready_input_file, session_dir=session)
        assert Path(result.output_path).is_relative_to(session)
        assert "Step2_" in Path(result.output_path).name


class TestQCLOWESSHelpers:
    """Tests for helper functions."""

    def test_feature_missing_terminal_qc_never_extrapolates(
        self,
        qc_lowess_module,
    ):
        qc_orders = np.array([1, 2, 3, 4, 5, 6, 7, 8, 10], dtype=float)
        qc_values = np.array(
            [100, 200, 300, 400, 500, 600, 700, 800, np.nan],
            dtype=float,
        )
        all_orders = np.arange(1, 11, dtype=float)
        all_values = np.array(
            [100, 200, 300, 400, 500, 600, 700, 800, 900, np.nan],
            dtype=float,
        )

        corrected, info = qc_lowess_module.apply_lowess_correction(
            qc_orders,
            qc_values,
            all_orders,
            all_values,
        )

        assert info["status"] == "success"
        assert info["valid_qc_count"] == 8
        assert info["outside_qc_range_count"] == 1
        assert corrected[8] == pytest.approx(all_values[8])

    def test_perform_lowess_supports_multi_batch_qc_membership(self, qc_lowess_module):
        """QC samples tagged as A;B should contribute to both batches instead of forming a new batch."""
        sample_info_df = pd.DataFrame(
            {
                "Sample_Name": [
                    "QC1",
                    "QC2",
                    "QC3",
                    "QC4",
                    "QC5",
                    "QC6",
                    "QC7",
                    "QC8",
                    "QC9",
                    "QC10",
                    "QC11",
                    "SampleA",
                    "SampleB",
                    "SampleC",
                ],
                "Sample_Type": [
                    "QC",
                    "QC",
                    "QC",
                    "QC",
                    "QC",
                    "QC",
                    "QC",
                    "QC",
                    "QC",
                    "QC",
                    "QC",
                    "Exposure",
                    "Control",
                    "Exposure",
                ],
                "Batch": [
                    "A;B",
                    "A;B;C",
                    "A;B;C",
                    "A;B;C",
                    "A;B;C",
                    "A;B;C",
                    "A;B;C",
                    "A;B;C",
                    "A;B;C",
                    "A;B;C",
                    "A;B;C",
                    "A",
                    "B",
                    "C",
                ],
                "Injection_Order": [1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 3.5, 6.5, 9.5],
            }
        )
        istd_df = pd.DataFrame(
            [
                {
                    "FeatureID": "100.1/5.0",
                    "QC1": 10.0,
                    "QC2": 20.0,
                    "QC3": 30.0,
                    "QC4": 40.0,
                    "QC5": 50.0,
                    "QC6": 60.0,
                    "QC7": 70.0,
                    "QC8": 80.0,
                    "QC9": 90.0,
                    "QC10": 100.0,
                    "QC11": 110.0,
                    "SampleA": 35.0,
                    "SampleB": 65.0,
                    "SampleC": 95.0,
                }
            ]
        )
        istd_df.attrs["sample_columns"] = [
            "QC1",
            "QC2",
            "QC3",
            "QC4",
            "QC5",
            "QC6",
            "QC7",
            "QC8",
            "QC9",
            "QC10",
            "QC11",
            "SampleA",
            "SampleB",
            "SampleC",
        ]

        lowess_df, _, _, _, decision_stats, _ = qc_lowess_module.perform_lowess_normalization(
            istd_df, sample_info_df
        )

        assert decision_stats["event_counts"]["insufficient_qc"] == 0
        assert lowess_df.loc[0, "SampleC"] != pytest.approx(95.0)
        assert decision_stats["event_counts"]["outlier_filtering_left_too_few_points"] == 0
        assert decision_stats["success"] == 1

    def test_trend_stats_schema_keeps_kendall_tau_only(self, qc_lowess_module):
        sample_info_df = pd.DataFrame(
            {
                "Sample_Name": ["QC1", "QC2", "QC3", "QC4", "QC5", "SampleA"],
                "Sample_Type": ["QC", "QC", "QC", "QC", "QC", "Exposure"],
                "Batch": ["A", "A", "A", "A", "A", "A"],
                "Injection_Order": [1, 2, 3, 4, 5, 6],
            }
        )
        istd_df = pd.DataFrame(
            [
                {
                    "FeatureID": "100.1/5.0",
                    "QC1": 10.0,
                    "QC2": 12.0,
                    "QC3": 14.0,
                    "QC4": 16.0,
                    "QC5": 18.0,
                    "SampleA": 20.0,
                }
            ]
        )
        istd_df.attrs["sample_columns"] = ["QC1", "QC2", "QC3", "QC4", "QC5", "SampleA"]

        _, _, _, trend_stats_df, _, _ = qc_lowess_module.perform_lowess_normalization(
            istd_df, sample_info_df
        )

        assert "Kendall_Tau" in trend_stats_df.columns
        assert "MK_Trend_pvalue" not in trend_stats_df.columns

    def test_get_valid_values_consistency(self, qc_lowess_module, istd_module):
        """Test that get_valid_values is consistent with ISTD module."""
        import pandas as pd

        row = pd.Series({'A': 100, 'B': 200, 'C': 0, 'D': -5})
        columns = ['A', 'B', 'C', 'D']

        values_qc = qc_lowess_module.get_valid_values(row, columns)
        values_istd = istd_module.get_valid_values(row, columns)

        # Should produce identical results
        assert values_qc == values_istd, "get_valid_values should be consistent across modules"


class TestFracFloorAndLoocv:
    """Tests for the LOWESS anti-overfitting guards."""

    @pytest.mark.parametrize("n_qc", [6, 7])
    def test_six_or_seven_qc_use_log_linear_fallback_when_predictive(
        self,
        qc_lowess_module,
        n_qc,
    ):
        all_orders = np.arange(1, 14, dtype=float)
        qc_orders = np.linspace(1, 13, n_qc)
        all_intensities = 1000.0 * 2.0 ** (0.08 * (all_orders - 7.0))
        qc_intensities = 1000.0 * 2.0 ** (0.08 * (qc_orders - 7.0))

        corrected, info = qc_lowess_module.apply_lowess_correction(
            qc_orders,
            qc_intensities,
            all_orders,
            all_intensities,
        )

        assert info["status"] == "success"
        assert info["fit_strategy"] == "log_linear_fallback"
        assert info["linear_loocv_gain"] >= 0.10
        assert np.isfinite(info["linear_loocv_rmse_log2"])
        assert np.isnan(info["loocv_rmse"])
        assert info["linear_shrinkage_factor"] == pytest.approx(0.5)
        expected = 1000.0 * 2.0 ** (0.04 * (all_orders - 7.0))
        assert corrected == pytest.approx(expected)
        expected_qc = 1000.0 * 2.0 ** (0.04 * (qc_orders - 7.0))
        expected_qc_cv = (
            np.std(expected_qc, ddof=1) / np.mean(expected_qc) * 100.0
        )
        assert info["cv_after"] == pytest.approx(expected_qc_cv)

    def test_apply_lowess_correction_populates_kendall_tau(self, qc_lowess_module):
        qc_orders = np.arange(1, 9, dtype=float)
        qc_intensities = np.array([100.0, 104.0, 109.0, 115.0, 122.0, 130.0, 139.0, 149.0])
        all_orders = qc_orders.copy()
        all_intensities = qc_intensities.copy()

        _, info = qc_lowess_module.apply_lowess_correction(
            qc_orders.tolist(),
            qc_intensities.tolist(),
            all_orders.tolist(),
            all_intensities.tolist(),
        )

        trend_validation = info["trend_validation"]
        assert np.isfinite(trend_validation["trend_tau"])
        assert -1.0 <= trend_validation["trend_tau"] <= 1.0

    def test_frac_floor_values(self, qc_lowess_module):
        floor = qc_lowess_module._frac_floor
        assert floor(8) == 0.70
        assert floor(10) == 0.70
        assert floor(11) == 0.0
        assert floor(20) == 0.0

    def test_loocv_rmse_returns_positive_float(self, qc_lowess_module):
        rng = np.random.default_rng(42)
        x = np.arange(7, dtype=float)
        y = 1000.0 + 50.0 * x + rng.normal(0, 20, 7)

        rmse = qc_lowess_module._loocv_rmse(x, y, frac=0.8)

        assert isinstance(rmse, float)
        assert rmse > 0

    def test_five_qc_points_are_not_corrected(
        self,
        qc_lowess_module,
        monkeypatch,
    ):
        def fail_if_lowess_is_called(*_args, **_kwargs):
            raise AssertionError("LOWESS must not run with fewer than 8 valid QC points")

        monkeypatch.setattr(
            qc_lowess_module.sm.nonparametric,
            "lowess",
            fail_if_lowess_is_called,
        )
        n_qc = 5
        qc_orders = np.arange(1, n_qc + 1, dtype=float)
        qc_intensities = 1000.0 + 100.0 * qc_orders

        corrected, info = qc_lowess_module.apply_lowess_correction(
            qc_orders.tolist(),
            qc_intensities.tolist(),
            qc_orders.tolist(),
            qc_intensities.tolist(),
        )

        assert info["status"] == "insufficient_qc"
        assert info["fit_strategy"] == "unknown"
        assert info["frac_strategy"] == "insufficient_qc"
        assert np.isnan(info["frac_used"])
        assert np.isnan(info["loocv_rmse"])
        assert corrected == pytest.approx(qc_intensities.tolist())

    def test_log_linear_fallback_keeps_original_when_loocv_does_not_improve(
        self,
        qc_lowess_module,
    ):
        qc_orders = np.array([1, 3, 5, 7, 9, 11, 13], dtype=float)
        qc_intensities = np.array([100, 101, 99, 100, 101, 99, 100], dtype=float)
        all_orders = np.arange(1, 14, dtype=float)
        all_intensities = np.full(13, 100.0)

        corrected, info = qc_lowess_module.apply_lowess_correction(
            qc_orders,
            qc_intensities,
            all_orders,
            all_intensities,
        )

        assert info["status"] == "linear_fallback_no_predictive_gain"
        assert corrected == pytest.approx(all_intensities)

    def test_log_linear_fallback_requires_significant_monotonic_trend(
        self,
        qc_lowess_module,
    ):
        qc_orders = np.array([1, 3, 5, 7, 9, 11, 13], dtype=float)
        qc_intensities = np.array(
            [1043.81, 1094.25, 1087.23, 1013.93, 1170.11, 1164.67, 1126.04],
            dtype=float,
        )
        all_orders = np.arange(1, 14, dtype=float)
        all_intensities = np.full(13, 1100.0)

        corrected, info = qc_lowess_module.apply_lowess_correction(
            qc_orders,
            qc_intensities,
            all_orders,
            all_intensities,
        )

        assert info["linear_loocv_gain"] >= 0.10
        assert info["status"] == "linear_fallback_no_monotonic_trend"
        assert corrected == pytest.approx(all_intensities)

    def test_log_linear_fallback_requires_feature_level_endpoint_qc(
        self,
        qc_lowess_module,
    ):
        qc_orders = np.array([1, 3, 5, 7, 9, 11, 12], dtype=float)
        qc_intensities = 1000.0 * 2.0 ** (0.08 * (qc_orders - 7.0))
        all_orders = np.arange(1, 14, dtype=float)
        all_intensities = 1000.0 * 2.0 ** (0.08 * (all_orders - 7.0))

        corrected, info = qc_lowess_module.apply_lowess_correction(
            qc_orders,
            qc_intensities,
            all_orders,
            all_intensities,
        )

        assert info["status"] == "linear_fallback_endpoint_uncovered"
        assert corrected == pytest.approx(all_intensities)

    def test_log_linear_fallback_rejects_duplicate_qc_injection_orders(
        self,
        qc_lowess_module,
    ):
        qc_orders = np.array([1, 3, 5, 7, 9, 13, 13], dtype=float)
        qc_intensities = 1000.0 * 2.0 ** (0.08 * (qc_orders - 7.0))
        all_orders = np.arange(1, 14, dtype=float)
        all_intensities = 1000.0 * 2.0 ** (0.08 * (all_orders - 7.0))

        corrected, info = qc_lowess_module.apply_lowess_correction(
            qc_orders,
            qc_intensities,
            all_orders,
            all_intensities,
        )

        assert info["status"] == "invalid_injection_order"
        assert corrected == pytest.approx(all_intensities)

    def test_log_linear_fallback_rejects_missing_observed_injection_order(
        self,
        qc_lowess_module,
    ):
        qc_orders = np.array([1, 3, 5, 7, 9, 11, 13], dtype=float)
        qc_intensities = 1000.0 * 2.0 ** (0.08 * (qc_orders - 7.0))
        all_orders = np.arange(1, 14, dtype=float)
        all_intensities = 1000.0 * 2.0 ** (0.08 * (all_orders - 7.0))

        corrected, info = qc_lowess_module.apply_lowess_correction(
            qc_orders,
            qc_intensities,
            all_orders,
            all_intensities,
            order_is_observed=False,
        )

        assert info["status"] == "invalid_injection_order"
        assert corrected == pytest.approx(all_intensities)

    def test_log_linear_fallback_rejects_drift_not_larger_than_noise(
        self,
        qc_lowess_module,
        monkeypatch,
    ):
        qc_orders = np.array([1, 3, 5, 7, 9, 11, 13], dtype=float)
        qc_intensities = np.array([1000, 1300, 900, 1200, 900, 1300, 1000], dtype=float)
        all_orders = np.arange(1, 14, dtype=float)
        all_intensities = np.full(13, 1000.0)
        monkeypatch.setattr(qc_lowess_module, "_log_linear_loocv", lambda *_: (1.0, 0.5))
        monkeypatch.setattr(qc_lowess_module, "kendalltau", lambda *_: (0.8, 0.01))

        corrected, info = qc_lowess_module.apply_lowess_correction(
            qc_orders,
            qc_intensities,
            all_orders,
            all_intensities,
        )

        assert info["status"] == "linear_fallback_drift_below_noise"
        assert corrected == pytest.approx(all_intensities)

    def test_log_linear_fallback_rejects_model_requiring_factor_clamping(
        self,
        qc_lowess_module,
    ):
        qc_orders = np.array([1, 3, 5, 7, 9, 11, 13], dtype=float)
        qc_intensities = 1000.0 * 2.0 ** (0.30 * (qc_orders - 7.0))
        all_orders = np.arange(1, 14, dtype=float)
        all_intensities = 1000.0 * 2.0 ** (0.30 * (all_orders - 7.0))

        corrected, info = qc_lowess_module.apply_lowess_correction(
            qc_orders,
            qc_intensities,
            all_orders,
            all_intensities,
        )

        assert info["status"] == "unstable_correction_factors"
        assert info["clamped_count"] > 0
        assert corrected == pytest.approx(all_intensities)

    def test_eight_qc_points_use_lowess(self, qc_lowess_module):
        qc_orders = np.arange(1, 9, dtype=float)
        qc_intensities = 1000.0 + 20.0 * qc_orders + 3.0 * qc_orders**2

        _, info = qc_lowess_module.apply_lowess_correction(
            qc_orders.tolist(),
            qc_intensities.tolist(),
            qc_orders.tolist(),
            qc_intensities.tolist(),
        )

        assert info["fit_strategy"] == "batch_local_lowess"
        assert info["frac_strategy"] != "linear_fallback"
        assert np.isfinite(info["frac_used"])

    def test_eight_raw_qc_are_not_corrected_when_outlier_filter_leaves_seven(
        self,
        qc_lowess_module,
        monkeypatch,
    ):
        def fail_if_lowess_is_called(*_args, **_kwargs):
            raise AssertionError("post-filter count 7 must not run LOWESS")

        monkeypatch.setattr(
            qc_lowess_module.sm.nonparametric,
            "lowess",
            fail_if_lowess_is_called,
        )
        qc_orders = np.arange(1, 9, dtype=float)
        qc_intensities = np.array([100, 105, 110, 115, 120, 125, 130, 1000], dtype=float)

        _, info = qc_lowess_module.apply_lowess_correction(
            qc_orders.tolist(),
            qc_intensities.tolist(),
            qc_orders.tolist(),
            qc_intensities.tolist(),
        )

        assert info["removed_outlier_count"] == 1
        assert info["valid_qc_count"] == 7
        assert info["status"] == "linear_fallback_outlier_filtering_not_allowed"
        assert info["fit_strategy"] == "unknown"

    def test_large_qc_count_skips_loocv(self, qc_lowess_module):
        """With n > 10 QC points, LOOCV should not be triggered."""
        rng = np.random.default_rng(123)
        n_qc = 20
        n_total = 60
        qc_orders = np.linspace(1, n_total, n_qc)
        qc_intensities = 50000.0 + rng.normal(0, 2000, n_qc)
        all_orders = np.arange(1, n_total + 1, dtype=float)
        all_intensities = 50000.0 + rng.normal(0, 3000, n_total)

        _, info = qc_lowess_module.apply_lowess_correction(
            qc_orders.tolist(),
            qc_intensities.tolist(),
            all_orders.tolist(),
            all_intensities.tolist(),
        )

        assert np.isnan(info.get("loocv_rmse", np.nan))
        assert "loocv" not in info["frac_strategy"]
        assert "floor" not in info["frac_strategy"]

    @pytest.mark.parametrize("n_qc", [1, 4])
    def test_one_to_four_valid_qc_points_are_not_corrected(self, qc_lowess_module, n_qc):
        qc_orders = np.arange(1, n_qc + 1, dtype=float)
        qc_intensities = 1000.0 + 100.0 * qc_orders
        all_orders = list(range(1, 21))
        all_intensities = [1000.0] * 20

        corrected, info = qc_lowess_module.apply_lowess_correction(
            qc_orders.tolist(),
            qc_intensities.tolist(),
            all_orders,
            all_intensities,
        )

        assert info["status"] == "insufficient_qc"
        assert info["valid_qc_count"] == n_qc
        assert corrected == all_intensities


class TestStep2ResponsibilityContract:
    """Canonical Step 2 contract tests for current LOWESS correction responsibilities."""

    def test_apply_lowess_correction_no_longer_accepts_global_qc_median(self, qc_lowess_module):
        qc_orders = [1.0, 2.0, 3.0, 4.0, 5.0]
        qc_intensities = [100.0, 105.0, 110.0, 115.0, 120.0]
        all_orders = [1.0, 2.0, 3.0, 4.0, 5.0, 6.0]
        all_intensities = [100.0, 105.0, 110.0, 115.0, 120.0, 125.0]

        with pytest.raises(TypeError):
            qc_lowess_module.apply_lowess_correction(
                qc_orders,
                qc_intensities,
                all_orders,
                all_intensities,
                global_qc_median=999.0,
            )

    def test_apply_lowess_correction_reports_all_qc_invalid_status(self, qc_lowess_module):
        qc_orders = [1.0, 2.0, 3.0, 4.0, 5.0]
        qc_intensities = [0.0, -1.0, np.nan, 0.0, -5.0]
        all_orders = [1.0, 2.0, 3.0, 4.0, 5.0, 6.0]
        all_intensities = [10.0, 11.0, 12.0, 13.0, 14.0, 15.0]

        _, info = qc_lowess_module.apply_lowess_correction(
            qc_orders,
            qc_intensities,
            all_orders,
            all_intensities,
        )

        assert info["status"] == "all_qc_invalid"

    def test_apply_lowess_correction_skips_stable_feature_as_no_drift_detected(self, qc_lowess_module):
        qc_orders = [1.0, 2.0, 3.0, 4.0, 5.0, 6.0, 7.0, 8.0]
        qc_intensities = [100.0, 100.1, 99.9, 100.1, 99.9, 100.1, 99.9, 100.0]
        all_orders = [1.0, 2.0, 3.0, 4.0, 5.0, 6.0, 7.0, 8.0]
        all_intensities = qc_intensities.copy()

        corrected, info = qc_lowess_module.apply_lowess_correction(
            qc_orders,
            qc_intensities,
            all_orders,
            all_intensities,
        )

        assert info["status"] == "no_drift_detected"
        assert corrected == pytest.approx(all_intensities)

    def test_apply_lowess_correction_reports_clamp_and_outside_range_metadata(self, qc_lowess_module):
        qc_orders = [10.0, 20.0, 30.0, 40.0, 50.0, 60.0, 70.0, 80.0]
        qc_intensities = [100.0, 200.0, 300.0, 400.0, 500.0, 600.0, 700.0, 800.0]
        all_orders = [1.0, 10.0, 20.0, 30.0, 40.0, 50.0, 60.0, 70.0, 80.0, 100.0]
        all_intensities = [90.0, 100.0, 200.0, 300.0, 400.0, 500.0, 600.0, 700.0, 800.0, 900.0]

        corrected, info = qc_lowess_module.apply_lowess_correction(
            qc_orders,
            qc_intensities,
            all_orders,
            all_intensities,
        )

        assert "raw_factor_min" in info
        assert "raw_factor_max" in info
        assert "clamped_factor_min" in info
        assert "clamped_factor_max" in info
        assert "clamped_count" in info
        assert "clamped_ratio" in info
        assert "outside_qc_range_count" in info
        assert info["outside_qc_range_count"] == 2
        for idx in (0, 9):
            assert corrected[idx] == pytest.approx(all_intensities[idx])

    def test_rejected_correction_retains_raw_matrix_and_qc_values(
        self,
        qc_lowess_module,
        monkeypatch,
    ):
        sample_names = ["QC1", "QC2", "QC3", "QC4", "QC5", "SampleB"]
        raw_values = [10.0, 20.0, 30.0, 40.0, 50.0, 60.0]
        sample_info_df = pd.DataFrame(
            {
                "Sample_Name": sample_names,
                "Sample_Type": ["QC", "QC", "QC", "QC", "QC", "Control"],
                "Batch": ["B"] * len(sample_names),
                "Injection_Order": list(range(1, len(sample_names) + 1)),
            }
        )
        istd_df = pd.DataFrame(
            [{"FeatureID": "290.1772/8.32", **dict(zip(sample_names, raw_values))}]
        )
        istd_df.attrs["sample_columns"] = sample_names

        def fake_rejected_correction(
            qc_orders,
            qc_intensities,
            all_orders,
            all_intensities,
            debug_flag=None,
        ):
            return [value * 2 for value in all_intensities], {
                "status": "unstable_correction_factors",
                "trend_validation": {},
                "valid_qc_count": len(qc_orders),
                "removed_outlier_count": 0,
                "outlier_filter_applied": False,
                "outside_qc_range_count": 0,
            }

        monkeypatch.setattr(
            qc_lowess_module,
            "apply_lowess_correction",
            fake_rejected_correction,
        )

        lowess_df, _, qc_corrected_values, trend_stats_df, _, _ = (
            qc_lowess_module.perform_lowess_normalization(istd_df, sample_info_df)
        )

        assert lowess_df.loc[0, sample_names].to_numpy(dtype=float) == pytest.approx(raw_values)
        assert list(qc_corrected_values["290.1772/8.32"].values()) == pytest.approx(raw_values[:5])
        assert trend_stats_df.loc[0, "Decision_Status"] == "unstable_correction_factors"

    def test_qc_cv_uses_jointly_valid_named_qc_samples(self, qc_lowess_module):
        qc_columns = ["QC1", "QC2", "QC3", "QC4", "QC5"]
        istd_df = pd.DataFrame(
            [{
                "FeatureID": "100.1/5.0",
                "QC1": 10.0,
                "QC2": np.nan,
                "QC3": 30.0,
                "QC4": 40.0,
                "QC5": 50.0,
            }]
        )
        lowess_df = pd.DataFrame(
            [{"FeatureID": "100.1/5.0", **{qc: 999.0 for qc in qc_columns}}]
        )
        sample_info_df = pd.DataFrame(
            {"Sample_Name": qc_columns, "Sample_Type": ["QC"] * len(qc_columns)}
        )
        qc_corrected_values = {
            "100.1/5.0": {
                "QC1": 11.0,
                "QC2": 22.0,
                "QC3": 33.0,
                "QC4": np.nan,
                "QC5": 55.0,
            }
        }

        result = qc_lowess_module.calculate_qc_cv_with_statistical_test(
            istd_df,
            lowess_df,
            qc_columns,
            sample_info_df,
            qc_corrected_values,
        )

        expected_cv = np.std([10.0, 30.0, 50.0], ddof=1) / np.mean([10.0, 30.0, 50.0]) * 100
        assert result.loc[0, "Original_QC_CV%"] == pytest.approx(expected_cv)
        assert result.loc[0, "Corrected_QC_CV%"] == pytest.approx(expected_cv)
        assert result.loc[0, "CV_Improvement%"] == pytest.approx(0.0)

    def test_qc_cv_overview_includes_partial_success_as_applied(
        self,
        qc_lowess_module,
        monkeypatch,
        tmp_path,
    ):
        from matplotlib.axes import Axes

        scatter_calls = []
        original_scatter = Axes.scatter

        def capture_scatter(axis, x, y, *args, **kwargs):
            scatter_calls.append((np.asarray(x), np.asarray(y)))
            return original_scatter(axis, x, y, *args, **kwargs)

        monkeypatch.setattr(Axes, "scatter", capture_scatter)
        cv_results_df = pd.DataFrame(
            {
                "FeatureID": ["all_success", "partial", "rejected"],
                "Original_QC_CV%": [10.0, 20.0, 30.0],
                "Corrected_QC_CV%": [8.0, 18.0, 5.0],
                "CV_Improvement%": [2.0, 2.0, 25.0],
                "Decision_Status": [
                    "success",
                    "partial_success",
                    "unstable_correction_factors",
                ],
            }
        )
        decision_stats = {
            "event_counts": {"success": 2, "unstable_correction_factors": 1},
            "total_feature_batch_tasks": 3,
        }

        qc_lowess_module.plot_qc_cv_overview(
            cv_results_df,
            decision_stats,
            tmp_path,
            "test",
        )

        assert len(scatter_calls) == 1
        assert scatter_calls[0][0] == pytest.approx([10.0, 20.0])
        assert scatter_calls[0][1] == pytest.approx([8.0, 18.0])


class TestStep2HardeningContract:
    """Tests for the next Task 3/4 Step 2 hardening contract."""

    def test_filter_qc_outliers_iqr_removes_extreme_value_when_keep_thresholds_are_met(
        self,
        qc_lowess_module,
    ):
        filtered_x, filtered_y, meta = qc_lowess_module.filter_qc_outliers_iqr(
            [1.0, 2.0, 3.0, 4.0, 5.0, 6.0],
            [100.0, 102.0, 101.0, 103.0, 99.0, 1000.0],
        )

        assert filtered_x.tolist() == pytest.approx([1.0, 2.0, 3.0, 4.0, 5.0])
        assert filtered_y.tolist() == pytest.approx([100.0, 102.0, 101.0, 103.0, 99.0])
        assert meta["original_valid_count"] == 6
        assert meta["removed_outlier_count"] == 1
        assert meta["outlier_filter_applied"] is True

    def test_apply_lowess_correction_reports_outlier_filtering_left_too_few_points(
        self,
        qc_lowess_module,
    ):
        qc_orders = [1.0, 2.0, 3.0, 4.0, 5.0]
        qc_intensities = [100.0, 101.0, 102.0, 103.0, 1000.0]
        all_orders = [1.0, 2.0, 3.0, 4.0, 5.0, 6.0]
        all_intensities = [100.0, 101.0, 102.0, 103.0, 1000.0, 105.0]

        corrected, info = qc_lowess_module.apply_lowess_correction(
            qc_orders,
            qc_intensities,
            all_orders,
            all_intensities,
        )

        assert info["status"] == "outlier_filtering_left_too_few_points"
        assert corrected == pytest.approx(all_intensities)

    def test_perform_lowess_normalization_treats_no_drift_batches_as_successful(
        self,
        qc_lowess_module,
        monkeypatch,
    ):
        sample_info_df = pd.DataFrame(
            {
                "Sample_Name": [
                    "A_QC1", "A_QC2", "A_QC3", "A_QC4", "A_QC5", "SampleA",
                    "B_QC1", "B_QC2", "B_QC3", "B_QC4", "B_QC5", "SampleB",
                ],
                "Sample_Type": [
                    "QC", "QC", "QC", "QC", "QC", "Exposure",
                    "QC", "QC", "QC", "QC", "QC", "Control",
                ],
                "Batch": ["A", "A", "A", "A", "A", "A", "B", "B", "B", "B", "B", "B"],
                "Injection_Order": list(range(1, 13)),
            }
        )
        istd_df = pd.DataFrame(
            [
                {
                    "FeatureID": "100.1/5.0",
                    "A_QC1": 100.0,
                    "A_QC2": 101.0,
                    "A_QC3": 102.0,
                    "A_QC4": 103.0,
                    "A_QC5": 104.0,
                    "SampleA": 150.0,
                    "B_QC1": 100.0,
                    "B_QC2": 100.5,
                    "B_QC3": 99.8,
                    "B_QC4": 100.1,
                    "B_QC5": 100.2,
                    "SampleB": 149.0,
                }
            ]
        )
        istd_df.attrs["sample_columns"] = list(istd_df.columns[1:])

        def fake_apply_lowess_correction(qc_orders, qc_intensities, all_orders, all_intensities, debug_flag=None):
            status = "success" if min(qc_orders) < 6 else "no_drift_detected"
            return list(all_intensities), {
                "status": status,
                "trend_validation": {
                    "trend_pvalue": 0.5,
                    "trend_tau": 0.02,
                    "r_squared": 0.02,
                    "rmse": 1.0,
                },
                "valid_qc_count": len(qc_orders),
                "removed_outlier_count": 0,
                "outlier_filter_applied": False,
                "normalized_rmse": 0.02,
                "target_strategy": "batch_local_fit_median",
                "clamped_ratio": 0.0,
                "outside_qc_range_count": 0,
                "frac_used": 0.5,
                "qc_cv_for_frac": 10.0,
                "frac_strategy": "moderate_cv",
            }

        monkeypatch.setattr(
            qc_lowess_module,
            "apply_lowess_correction",
            fake_apply_lowess_correction,
        )

        _, _, _, trend_stats_df, decision_stats, _ = qc_lowess_module.perform_lowess_normalization(
            istd_df,
            sample_info_df,
        )

        assert decision_stats["success"] == 1
        assert decision_stats["partial_success"] == 0
        assert trend_stats_df.loc[0, "Decision_Status"] == "success"

    def test_perform_lowess_normalization_surfaces_all_no_drift_feature_status(
        self,
        qc_lowess_module,
        monkeypatch,
    ):
        sample_info_df = pd.DataFrame(
            {
                "Sample_Name": [
                    "A_QC1", "A_QC2", "A_QC3", "A_QC4", "A_QC5", "SampleA",
                    "B_QC1", "B_QC2", "B_QC3", "B_QC4", "B_QC5", "SampleB",
                ],
                "Sample_Type": [
                    "QC", "QC", "QC", "QC", "QC", "Exposure",
                    "QC", "QC", "QC", "QC", "QC", "Control",
                ],
                "Batch": ["A", "A", "A", "A", "A", "A", "B", "B", "B", "B", "B", "B"],
                "Injection_Order": list(range(1, 13)),
            }
        )
        istd_df = pd.DataFrame(
            [
                {
                    "FeatureID": "100.1/5.0",
                    "A_QC1": 100.0,
                    "A_QC2": 100.5,
                    "A_QC3": 99.8,
                    "A_QC4": 100.1,
                    "A_QC5": 100.2,
                    "SampleA": 101.0,
                    "B_QC1": 99.7,
                    "B_QC2": 100.1,
                    "B_QC3": 100.0,
                    "B_QC4": 100.4,
                    "B_QC5": 99.9,
                    "SampleB": 100.8,
                }
            ]
        )
        istd_df.attrs["sample_columns"] = list(istd_df.columns[1:])

        def fake_apply_lowess_correction(qc_orders, qc_intensities, all_orders, all_intensities, debug_flag=None):
            return list(all_intensities), {
                "status": "no_drift_detected",
                "trend_validation": {
                    "trend_pvalue": 0.5,
                    "trend_tau": 0.02,
                    "r_squared": 0.02,
                    "rmse": 1.0,
                },
                "valid_qc_count": len(qc_orders),
                "removed_outlier_count": 0,
                "outlier_filter_applied": False,
                "normalized_rmse": 0.02,
                "target_strategy": "batch_local_fit_median",
                "clamped_ratio": 0.0,
                "outside_qc_range_count": 0,
                "frac_used": 0.5,
                "qc_cv_for_frac": 10.0,
                "frac_strategy": "moderate_cv",
            }

        monkeypatch.setattr(
            qc_lowess_module,
            "apply_lowess_correction",
            fake_apply_lowess_correction,
        )

        _, _, _, trend_stats_df, decision_stats, _ = qc_lowess_module.perform_lowess_normalization(
            istd_df,
            sample_info_df,
        )

        assert decision_stats["success"] == 0
        assert decision_stats["no_drift_detected"] == 1
        assert decision_stats["partial_success"] == 0
        assert trend_stats_df.loc[0, "Decision_Status"] == "no_drift_detected"

    @pytest.mark.slow
    @pytest.mark.integration
    def test_advanced_stats_sheet_includes_step3_contract_columns(
        self,
        qc_lowess_module,
        lowess_ready_input_file,
    ):
        step2_result = qc_lowess_module.main(input_file=lowess_ready_input_file)
        step2_output = (
            step2_result.output_path
            if hasattr(step2_result, "output_path")
            else step2_result.get("output_path")
        )

        workbook = load_workbook(step2_output, read_only=True, data_only=True)
        try:
            worksheet = workbook[SHEET_NAMES["qc_lowess_advanced"]]
            headers = [cell.value for cell in worksheet[1]]
        finally:
            workbook.close()

        expected_headers = {
            "Valid_QC_Count",
            "Removed_QC_Outliers",
            "Outlier_Filter_Applied",
            "Trend_pvalue",
            "Kendall_Tau",
            "LOESS_R2",
            "LOESS_RMSE",
            "Normalized_RMSE",
            "Target_Strategy",
            "Fit_Strategy",
            "Batch_Decision_Detail",
            "Clamped_Factor_Ratio",
            "Outside_QC_Range_Count",
            "Decision_Status",
        }
        assert expected_headers.issubset(set(headers))
