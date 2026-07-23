import json

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from metabolomics.processors import normalization, qc_lowess
from metabolomics.utils.design_identifiability import (
    DesignStatus,
    build_design_identifiability_receipt,
)
from metabolomics.utils.results import WorkflowOutcome
from scripts.synthetic_matrix_vnext import generate_simulation


def _contrast(receipt, left, right):
    row = receipt.contrasts.loc[
        (receipt.contrasts["Left_Sample_Type"] == left)
        & (receipt.contrasts["Right_Sample_Type"] == right)
    ]
    assert len(row) == 1
    return row.iloc[0]


def test_balanced_design_supports_biological_contrast_but_not_cross_batch_alignment():
    simulation = generate_simulation(
        "routine_recoverable",
        seed=51,
        design="balanced",
    )

    receipt = build_design_identifiability_receipt(simulation.sample_info)

    assert receipt.summary["overall_status"] == DesignStatus.ASSUMPTION_DEPENDENT.value
    assert receipt.summary["interaction_estimable"] is True
    assert receipt.summary["within_batch_qc_drift_status"] == DesignStatus.SUPPORTED.value
    assert (
        receipt.summary["cross_batch_level_alignment_status"]
        == DesignStatus.NON_IDENTIFIABLE.value
    )
    assert receipt.qc_coverage["Endpoint_Covered"].all()
    assert set(receipt.batch_sample_counts["Sample_Type"]) == {
        "Control",
        "Exposure",
        "QC",
    }

    contrast = _contrast(receipt, "Control", "Exposure")
    assert contrast["Status"] == DesignStatus.SUPPORTED.value
    assert contrast["Supporting_Batch_Count"] == 3


def test_order_confounded_design_is_not_promoted_when_model_can_still_execute():
    simulation = generate_simulation(
        "routine_recoverable",
        seed=51,
        design="order_confounded",
    )

    receipt = build_design_identifiability_receipt(simulation.sample_info)

    assert receipt.summary["overall_status"] == DesignStatus.NON_IDENTIFIABLE.value
    assert receipt.summary["interaction_estimable"] is False
    assert receipt.summary["empty_batch_sample_type_cells"] == 2
    assert receipt.summary["sample_type_order_eta_squared"] > 0.7

    contrast = _contrast(receipt, "Control", "Exposure")
    assert contrast["Status"] == DesignStatus.ASSUMPTION_DEPENDENT.value
    assert contrast["Supporting_Batch_Count"] == 1
    assert not bool(contrast["Universal_Cross_Batch_Effect"])


def test_optional_pair_and_bridge_metadata_are_reported_without_becoming_required():
    sample_info = pd.DataFrame(
        {
            "Sample_Name": ["A_C", "A_E", "B_C", "B_E", "QCA", "QCB"],
            "Sample_Type": ["Control", "Exposure", "Control", "Exposure", "QC", "QC"],
            "Batch": ["A", "A", "B", "B", "A", "B"],
            "Injection_Order": [2, 3, 6, 7, 1, 8],
            "Pair_ID": ["P1", "P1", "P2", "P2", pd.NA, pd.NA],
            "Bridge_ID": ["BR1", pd.NA, "BR1", pd.NA, pd.NA, pd.NA],
        }
    )

    receipt = build_design_identifiability_receipt(sample_info)

    assert receipt.summary["pair_metadata_status"] == "available"
    assert receipt.summary["complete_cross_type_pair_count"] == 2
    assert receipt.summary["bridge_metadata_status"] == "available"
    assert receipt.summary["cross_batch_bridge_count"] == 1
    assert (
        receipt.summary["cross_batch_level_alignment_status"]
        == DesignStatus.SUPPORTED.value
    )


def test_non_comparable_qc_pool_ids_are_flagged_but_missing_pool_metadata_is_unknown():
    base = pd.DataFrame(
        {
            "Sample_Name": ["QCA1", "A1", "QCA2", "QCB1", "B1", "QCB2"],
            "Sample_Type": ["QC", "Control", "QC", "QC", "Exposure", "QC"],
            "Batch": ["A", "A", "A", "B", "B", "B"],
            "Injection_Order": [1, 2, 3, 4, 5, 6],
        }
    )

    unknown = build_design_identifiability_receipt(base)
    assert unknown.summary["qc_pool_comparability"] == "unknown"

    declared = base.assign(
        QC_Pool_ID=["exp+nor", pd.NA, "exp+nor", "exp+nor+con", pd.NA, "exp+nor+con"]
    )
    non_comparable = build_design_identifiability_receipt(declared)

    assert non_comparable.summary["qc_pool_comparability"] == "non_comparable"
    assert (
        non_comparable.summary["cross_batch_level_alignment_status"]
        == DesignStatus.NON_IDENTIFIABLE.value
    )


def test_semicolon_qc_batch_membership_is_expanded_without_creating_fake_batch():
    sample_info = pd.DataFrame(
        {
            "Sample_Name": ["QC1", "A_C", "A_E", "B_C", "B_E", "QC2"],
            "Sample_Type": ["QC", "Control", "Exposure", "Control", "Exposure", "QC"],
            "Batch": ["A;B", "A", "A", "B", "B", "A;B"],
            "Injection_Order": [1, 2, 3, 6, 7, 8],
        }
    )

    receipt = build_design_identifiability_receipt(sample_info)

    assert set(receipt.batch_sample_counts["Batch"]) == {"A", "B"}
    qc_counts = (
        receipt.batch_sample_counts.loc[
            receipt.batch_sample_counts["Sample_Type"].eq("QC")
        ]
        .set_index("Batch")["Sample_Count"]
        .to_dict()
    )
    assert qc_counts == {"A": 2, "B": 2}
    assert receipt.qc_coverage["Endpoint_Covered"].all()


def test_receipt_exports_machine_readable_summary_and_detail_sheets():
    simulation = generate_simulation("routine_recoverable", seed=51)

    receipt = build_design_identifiability_receipt(simulation.sample_info)
    sheets = receipt.to_excel_sheets()
    payload = receipt.to_processing_extra()

    assert set(sheets) == {"Design_Identifiability"}
    assert set(sheets["Design_Identifiability"]["Record_Type"]) >= {
        "summary",
        "batch_sample_count",
        "qc_coverage",
        "association",
        "design_diagnostics",
        "contrast",
    }
    assert payload["summary"]["overall_status"] == receipt.summary["overall_status"]
    assert isinstance(payload["reasons"], list)
    json.dumps(payload, allow_nan=False)


def test_receipt_uses_only_matrix_samples_and_preserves_unknown_subtype_labels():
    sample_info = pd.DataFrame(
        {
            "Sample_Name": ["Q1", "Subtype_A", "Subtype_B", "Mixed_Extra", "Blank_1", "Q2"],
            "Sample_Type": ["QC", "Rare_A", "Rare_B", "Mixed", "Blank", "QC"],
            "Batch": ["A"] * 6,
            "Injection_Order": [1, 2, 3, 4, 5, 6],
        }
    )

    receipt = build_design_identifiability_receipt(
        sample_info,
        sample_columns=["Q1", "Subtype_A", "Subtype_B", "Blank_1", "Q2"],
    )

    assert receipt.summary["sample_count"] == 5
    assert receipt.summary["study_sample_count"] == 2
    assert set(receipt.batch_sample_counts["Sample_Type"]) == {
        "QC",
        "Rare_A",
        "Rare_B",
        "Blank",
    }
    contrast = _contrast(receipt, "Rare_A", "Rare_B")
    assert contrast["Status"] == DesignStatus.ASSUMPTION_DEPENDENT.value


def test_receipt_marks_missing_or_non_numeric_injection_order_unavailable():
    sample_info = pd.DataFrame(
        {
            "Sample_Name": ["Q1", "S1", "Q2"],
            "Sample_Type": ["QC", "Control", "QC"],
            "Batch": ["A", "A", "A"],
            "Injection_Order": [1, "unknown", 3],
        }
    )

    receipt = build_design_identifiability_receipt(sample_info)

    assert receipt.summary["injection_order_status"] == "unavailable"
    assert receipt.summary["overall_status"] == DesignStatus.NON_IDENTIFIABLE.value
    assert any("Injection_Order" in reason for reason in receipt.reasons)


def test_injection_order_may_restart_in_each_batch():
    sample_info = pd.DataFrame(
        {
            "Sample_Name": ["A_Q1", "A_S", "A_Q2", "B_Q1", "B_S", "B_Q2"],
            "Sample_Type": ["QC", "Control", "QC", "QC", "Exposure", "QC"],
            "Batch": ["A", "A", "A", "B", "B", "B"],
            "Injection_Order": [1, 2, 3, 1, 2, 3],
        }
    )

    receipt = build_design_identifiability_receipt(sample_info)

    assert receipt.summary["injection_order_status"] == "available"


def test_full_rank_but_order_separated_groups_are_assumption_dependent():
    sample_info = pd.DataFrame(
        {
            "Sample_Name": [
                "Q1",
                "C1",
                "C2",
                "Q2",
                "C3",
                "C4",
                "Q3",
                "E1",
                "E2",
                "Q4",
                "E3",
                "E4",
                "Q5",
                "Q6",
            ],
            "Sample_Type": [
                "QC",
                "Control",
                "Control",
                "QC",
                "Control",
                "Control",
                "QC",
                "Exposure",
                "Exposure",
                "QC",
                "Exposure",
                "Exposure",
                "QC",
                "QC",
            ],
            "Batch": ["A"] * 14,
            "Injection_Order": range(1, 15),
        }
    )

    receipt = build_design_identifiability_receipt(sample_info)

    assert receipt.summary["interaction_estimable"] is True
    assert receipt.summary["sample_type_order_eta_squared"] > 0.7
    assert receipt.summary["overall_status"] == DesignStatus.ASSUMPTION_DEPENDENT.value
    contrast = _contrast(receipt, "Control", "Exposure")
    assert contrast["Status"] == DesignStatus.ASSUMPTION_DEPENDENT.value


def test_partial_bridge_graph_does_not_support_uncovered_batch_alignment():
    sample_info = pd.DataFrame(
        {
            "Sample_Name": [
                "A_Q1",
                "A_C",
                "A_E",
                "A_Q2",
                "B_Q1",
                "B_C",
                "B_E",
                "B_Q2",
                "C_Q1",
                "C_C",
                "C_E",
                "C_Q2",
            ],
            "Sample_Type": ["QC", "Control", "Exposure", "QC"] * 3,
            "Batch": ["A"] * 4 + ["B"] * 4 + ["C"] * 4,
            "Injection_Order": [1, 2, 3, 4] * 3,
            "Bridge_ID": [
                pd.NA,
                "AB",
                pd.NA,
                pd.NA,
                pd.NA,
                "AB",
                pd.NA,
                pd.NA,
                pd.NA,
                pd.NA,
                pd.NA,
                pd.NA,
            ],
        }
    )

    receipt = build_design_identifiability_receipt(sample_info)

    assert receipt.summary["cross_batch_bridge_count"] == 1
    assert receipt.summary["bridged_batch_count"] == 2
    assert receipt.summary["bridge_coverage_complete"] is False
    assert (
        receipt.summary["cross_batch_level_alignment_status"]
        == DesignStatus.NON_IDENTIFIABLE.value
    )


def test_order_warning_is_specific_to_each_pairwise_contrast():
    sample_info = pd.DataFrame(
        {
            "Sample_Name": [
                "Q1",
                "C1",
                "E1",
                "Q2",
                "C2",
                "E2",
                "Q3",
                "R1",
                "R2",
                "Q4",
                "Q5",
                "Q6",
            ],
            "Sample_Type": [
                "QC",
                "Control",
                "Exposure",
                "QC",
                "Control",
                "Exposure",
                "QC",
                "Rare",
                "Rare",
                "QC",
                "QC",
                "QC",
            ],
            "Batch": ["A"] * 12,
            "Injection_Order": range(1, 13),
        }
    )

    receipt = build_design_identifiability_receipt(sample_info)

    assert receipt.summary["sample_type_order_eta_squared"] > 0.7
    control_exposure = _contrast(receipt, "Control", "Exposure")
    control_rare = _contrast(receipt, "Control", "Rare")
    assert control_exposure["Status"] == DesignStatus.SUPPORTED.value
    assert control_exposure["Order_Eta_Squared"] < 0.7
    assert control_rare["Status"] == DesignStatus.ASSUMPTION_DEPENDENT.value


def test_step2_skip_still_returns_design_receipt_before_downstream_handoff(
    monkeypatch,
    tmp_path,
):
    input_path = tmp_path / "input.xlsx"
    input_path.touch()
    sample_names = ["QC1", "S1", "QC2", "QC3", "QC4", "QC5"]
    sample_info = pd.DataFrame(
        {
            "Sample_Name": sample_names,
            "Sample_Type": ["QC", "Exposure", "QC", "QC", "QC", "QC"],
            "Batch": ["A"] * len(sample_names),
            "Injection_Order": range(1, len(sample_names) + 1),
        }
    )
    source = pd.DataFrame(
        {
            "FeatureID": ["f1"],
            **{sample: [float(index + 1)] for index, sample in enumerate(sample_names)},
        }
    )
    trend = pd.DataFrame(
        {
            "FeatureID": ["f1"],
            "Decision_Status": ["insufficient_qc"],
        }
    )

    monkeypatch.setattr(
        qc_lowess,
        "load_and_process_data",
        lambda _path: (source.copy(), source.copy(), sample_info.copy(), None),
    )
    monkeypatch.setattr(
        qc_lowess,
        "perform_lowess_normalization",
        lambda _source, _sample_info: (
            source.copy(),
            sample_names,
            {},
            trend,
            {
                "event_counts": {"insufficient_qc": 1},
                "total_feature_batch_tasks": 1,
            },
            {},
        ),
    )

    result = qc_lowess.main(input_file=input_path)

    assert result.status is WorkflowOutcome.SKIPPED
    receipt = result.extra["design_identifiability"]
    assert receipt["summary"]["within_batch_qc_drift_status"] == (
        DesignStatus.ASSUMPTION_DEPENDENT.value
    )
    assert receipt["summary"]["sample_count"] == len(sample_names)


def test_step2_success_passes_receipt_to_writer_and_processing_result(
    monkeypatch,
    tmp_path,
):
    input_path = tmp_path / "input.xlsx"
    input_path.touch()
    sample_names = ["QC1", "S1", "QC2", "QC3", "QC4", "QC5", "QC6"]
    sample_info = pd.DataFrame(
        {
            "Sample_Name": sample_names,
            "Sample_Type": ["QC", "Exposure", "QC", "QC", "QC", "QC", "QC"],
            "Batch": ["A"] * len(sample_names),
            "Injection_Order": range(1, len(sample_names) + 1),
        }
    )
    source = pd.DataFrame(
        {
            "FeatureID": ["f1"],
            **{sample: [100.0 + index] for index, sample in enumerate(sample_names)},
        }
    )
    trend = pd.DataFrame(
        {
            "FeatureID": ["f1"],
            "Decision_Status": ["success"],
        }
    )
    captured = {}

    monkeypatch.setattr(
        qc_lowess,
        "load_and_process_data",
        lambda _path: (source.copy(), source.copy(), sample_info.copy(), None),
    )
    monkeypatch.setattr(
        qc_lowess,
        "perform_lowess_normalization",
        lambda _source, _sample_info: (
            source.copy(),
            sample_names,
            {"f1": np.asarray([100.0] * len(sample_names))},
            trend,
            {
                "event_counts": {"success": 1},
                "total_feature_batch_tasks": 1,
            },
            {},
        ),
    )

    def capture_writer(*_args, **kwargs):
        captured.update(kwargs)
        return True

    monkeypatch.setattr(qc_lowess, "save_results_to_excel", capture_writer)

    result = qc_lowess.main(input_file=input_path)

    assert result.status is WorkflowOutcome.SUCCEEDED
    assert "design_receipt" in captured
    assert (
        result.extra["design_identifiability"]["summary"]["overall_status"]
        == captured["design_receipt"].summary["overall_status"]
    )


def test_step3_writer_persists_canonical_design_receipt_sheet(tmp_path):
    sample_info = pd.DataFrame(
        {
            "Sample_Name": ["QC1", "S1", "QC2"],
            "Sample_Type": ["QC", "Exposure", "QC"],
            "Batch": ["A", "A", "A"],
            "Injection_Order": [1, 2, 3],
        }
    )
    data = pd.DataFrame(
        {
            "Mz/RT": ["100.0/1.0"],
            "QC1": [100.0],
            "S1": [110.0],
            "QC2": [101.0],
        }
    )
    receipt = build_design_identifiability_receipt(
        sample_info,
        sample_columns=["QC1", "S1", "QC2"],
    )
    output_path = tmp_path / "step3.xlsx"

    saved = normalization.save_normalization_results(
        data,
        "summary",
        str(tmp_path / "input.xlsx"),
        "PQN",
        "QC LOESS result",
        "SampleInfo",
        tmp_path,
        preserved_data_df=data,
        sample_info_df=sample_info,
        output_path=output_path,
        design_receipt=receipt,
    )

    assert saved == str(output_path)
    workbook = load_workbook(output_path, read_only=True)
    assert "Design_Identifiability" in workbook.sheetnames
    worksheet = workbook["Design_Identifiability"]
    headers = [cell.value for cell in next(worksheet.iter_rows(min_row=1, max_row=1))]
    assert "Record_Type" in headers
    workbook.close()
