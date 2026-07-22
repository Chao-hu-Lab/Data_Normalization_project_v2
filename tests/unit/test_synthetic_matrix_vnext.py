from pathlib import Path

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from scripts.synthetic_matrix_vnext import (
    evaluate_step2_candidate,
    generate_simulation,
    semantic_workbook_digest,
    write_simulation_workbook,
)


def test_routine_recoverable_exposes_current_input_and_truth_contract():
    result = generate_simulation("routine_recoverable", seed=51)

    assert result.observed.shape == (
        len(result.sample_info),
        len(result.feature_info),
    )
    assert result.truth.step2_counterfactual.shape == result.observed.shape
    assert np.isnan(result.observed).any()

    for _, batch_info in result.sample_info.groupby("Batch", sort=False):
        assert batch_info.iloc[0]["Sample_Type"] == "QC"
        assert batch_info.iloc[-1]["Sample_Type"] == "QC"
        assert int(batch_info["Sample_Type"].eq("QC").sum()) == 8

    analyte_strata = set(
        result.feature_info.loc[~result.feature_info["is_istd"], "biology_stratum"]
    )
    assert analyte_strata == {"null", "positive", "negative"}
    assert set(result.truth.component_matrices) >= {
        "latent_injection_abundance",
        "batch_factor",
        "ionization_factor",
        "drift_factor",
        "measurement_noise",
        "carryover_factor",
        "matrix_effect_factor",
        "saturation_factor",
    }
    assert set(result.truth.rng_stream_ids) == {
        "schedule",
        "biology",
        "measurement",
        "detection",
        "outliers",
        "carryover",
        "feature_assignments",
    }
    assert len(result.truth.pooled_qc_metadata["contributor_sample_names"]) == 48
    assert result.truth.pooled_qc_metadata["equal_volume_weight"] == pytest.approx(
        1.0 / 48.0
    )

    carryover = result.truth.component_matrices["carryover_factor"]
    assert (carryover > 1.0).any()
    for _, batch_info in result.sample_info.groupby("Batch", sort=False):
        np.testing.assert_array_equal(
            carryover[batch_info.index[0]],
            np.ones(len(result.feature_info)),
        )

    latent = result.truth.component_matrices["latent_injection_abundance"]
    study_mask = result.sample_info["Sample_Type"].ne("QC").to_numpy()
    qc_mask = ~study_mask
    assert np.isfinite(result.observed[qc_mask]).all()
    pooled_from_studies = np.mean(latent[study_mask], axis=0)
    np.testing.assert_allclose(
        latent[qc_mask],
        np.broadcast_to(pooled_from_studies, latent[qc_mask].shape),
    )


def test_workbook_round_trip_preserves_nan_schema_and_istd_markers(tmp_path):
    result = generate_simulation("routine_recoverable", seed=51)
    first_path = tmp_path / "first.xlsx"
    second_path = tmp_path / "second.xlsx"

    write_simulation_workbook(result, first_path)
    write_simulation_workbook(result, second_path)

    assert semantic_workbook_digest(first_path) == semantic_workbook_digest(second_path)
    sample_info = pd.read_excel(first_path, sheet_name="SampleInfo")
    raw = pd.read_excel(first_path, sheet_name="RawIntensity")
    assert sample_info.columns.tolist() == [
        "Sample_Name",
        "Sample_Type",
        "Injection_Order",
        "Batch",
        "Injection_Volume",
        "Creatinine_mg_dL",
    ]
    numeric = raw.iloc[1:, 1:].apply(pd.to_numeric, errors="coerce")
    assert int(numeric.isna().sum().sum()) == int(np.isnan(result.observed).sum())

    workbook = load_workbook(first_path, read_only=False, data_only=True)
    try:
        sheet = workbook["RawIntensity"]
        red_rows = {
            row
            for row in range(3, sheet.max_row + 1)
            if str(sheet.cell(row=row, column=1).font.color.rgb).upper() == "FFFF0000"
        }
        assert len(red_rows) == int(result.feature_info["is_istd"].sum())
    finally:
        workbook.close()


def test_tracked_canonical_workbook_matches_generator(tmp_path):
    result = generate_simulation("routine_recoverable", seed=51)
    regenerated = write_simulation_workbook(result, tmp_path / "regenerated.xlsx")
    tracked = (
        Path(__file__).resolve().parents[2]
        / "data"
        / "synthetic_correction_input.xlsx"
    )

    assert tracked.exists()
    assert semantic_workbook_digest(tracked) == semantic_workbook_digest(regenerated)


def test_paired_drift_variants_share_unperturbed_truth_and_random_uniforms():
    low = generate_simulation("routine_recoverable", seed=137, variant="drift_low")
    high = generate_simulation("routine_recoverable", seed=137, variant="drift_high")

    for component in [
        "latent_injection_abundance",
        "batch_factor",
        "ionization_factor",
        "measurement_noise",
        "outlier_factor",
        "carryover_factor",
        "detection_uniform",
    ]:
        np.testing.assert_array_equal(
            low.truth.component_matrices[component],
            high.truth.component_matrices[component],
        )
    assert not np.array_equal(
        low.truth.component_matrices["drift_factor"],
        high.truth.component_matrices["drift_factor"],
    )


def test_step2_metrics_recognize_perfect_counterfactual_recovery():
    result = generate_simulation("routine_recoverable", seed=51)
    perfect = result.truth.step2_counterfactual.copy()
    perfect[np.isnan(result.observed)] = np.nan

    metrics = evaluate_step2_candidate(result, perfect)

    assert metrics["median_technical_recovery_gain"] == 1.0
    assert metrics["false_correction_rate"] == 0.0
    assert metrics["biology_sign_flip_rate"] == 0.0
    assert metrics["median_abs_log2_fold_change_error"] == 0.0


def test_qc_limited_routing_has_exact_observable_counts_and_endpoint_truth():
    result = generate_simulation("qc_limited_routing", seed=51)
    routing = result.truth.routing_truth

    assert routing is not None
    assert set(routing["scheduled_qc_count"]) == {8}
    assert set(routing["effective_qc_count"]) == {4, 5, 6, 7, 8}
    assert set(routing["evidence_class"]) >= {
        "observable_insufficient",
        "characterize_linear",
        "characterize_no_drift",
        "lowess_eligible",
    }
    assert (~routing["endpoint_valid"]).any()
    assert int(routing["qc_outlier_count"].sum()) > 0
    assert result.truth.component_matrices["qc_outlier_factor"].max() > 1.0

    for record in routing.itertuples(index=False):
        feature_index = result.feature_info.index[
            result.feature_info["FeatureID"].eq(record.FeatureID)
        ][0]
        batch_rows = result.sample_info.index[
            result.sample_info["Batch"].eq(record.Batch)
            & result.sample_info["Sample_Type"].eq("QC")
        ]
        observed_count = np.isfinite(
            result.observed[batch_rows, feature_index]
        ).sum()
        assert int(observed_count) == record.effective_qc_count


@pytest.mark.parametrize("variant", ["drift_low", "drift_high"])
def test_qc_limited_routing_supports_holdout_drift_variants(variant):
    baseline = generate_simulation("qc_limited_routing", seed=51, variant="baseline")
    candidate = generate_simulation("qc_limited_routing", seed=51, variant=variant)

    pd.testing.assert_frame_equal(
        candidate.truth.routing_truth,
        baseline.truth.routing_truth,
    )
    np.testing.assert_array_equal(
        candidate.truth.component_matrices["qc_routing_observed_mask"],
        baseline.truth.component_matrices["qc_routing_observed_mask"],
    )
    assert not np.array_equal(
        candidate.truth.component_matrices["drift_factor"],
        baseline.truth.component_matrices["drift_factor"],
    )


def test_qc_limited_observable_insufficiency_is_not_corrected(
    qc_lowess_module,
):
    result = generate_simulation("qc_limited_routing", seed=51)
    sample_names = result.sample_info["Sample_Name"].tolist()
    source = pd.DataFrame(result.observed.T, columns=sample_names)
    source.insert(0, "FeatureID", result.feature_info["FeatureID"])
    source.attrs["sample_columns"] = sample_names

    corrected, _, _, summary, _, _ = (
        qc_lowess_module.perform_lowess_normalization(
            source,
            result.sample_info,
        )
    )
    details = summary.set_index("FeatureID")["Batch_Decision_Detail"]
    routing = result.truth.routing_truth
    assert routing is not None

    observed_decisions: dict[tuple[str, str], str] = {}
    for record in routing.itertuples(index=False):
        observed_decisions[(record.FeatureID, record.Batch)] = next(
            item
            for item in details[record.FeatureID].split("; ")
            if item.startswith(f"{record.Batch}:")
        )
    assert len(observed_decisions) == len(routing)

    for record in routing.loc[
        routing["evidence_class"].eq("observable_insufficient")
    ].itertuples(index=False):
        batch_samples = result.sample_info.loc[
            result.sample_info["Batch"].eq(record.Batch),
            "Sample_Name",
        ].tolist()
        feature_row = result.feature_info.index[
            result.feature_info["FeatureID"].eq(record.FeatureID)
        ][0]
        np.testing.assert_allclose(
            corrected.loc[feature_row, batch_samples].to_numpy(dtype=float),
            source.loc[feature_row, batch_samples].to_numpy(dtype=float),
            equal_nan=True,
        )
        batch_detail = observed_decisions[(record.FeatureID, record.Batch)]
        assert "status=success" not in batch_detail
        assert "valid_qc=" in batch_detail
