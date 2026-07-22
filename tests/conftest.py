"""
Pytest configuration and fixtures for Data Normalization Workflow v2 tests.
"""
import pytest
import os
import sys
import shutil
import tempfile
import re
from pathlib import Path
from datetime import datetime

from openpyxl import load_workbook

# Add src directory to path for imports
sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "src"))


SAFE_TMP_SESSION_KEEP = 3


def _sanitize_tmp_name(name: str) -> str:
    """Return a filesystem-safe, compact temp directory stem."""
    sanitized = re.sub(r"[^\w]+", "_", name).strip("_")
    return sanitized[:80] or "tmp"


def _prune_safe_tmp_sessions(root: Path, keep: int = SAFE_TMP_SESSION_KEEP) -> None:
    """Keep the most recent repo-local tmp fixture sessions and delete older ones."""
    sessions = sorted(
        (path for path in root.glob("session_*") if path.is_dir()),
        key=lambda path: path.stat().st_mtime,
        reverse=True,
    )
    for stale in sessions[keep:]:
        shutil.rmtree(stale, ignore_errors=True)


class RepoTempPathFactory:
    """Minimal tmp-path factory that avoids pytest's Windows-problematic temp flow."""

    def __init__(self, base_dir: Path) -> None:
        self._base_dir = base_dir
        self._counters: dict[str, int] = {}

    def getbasetemp(self) -> Path:
        """Return the session-scoped base temp directory."""
        return self._base_dir

    def mktemp(self, basename: str, numbered: bool = True) -> Path:
        """Create a child temp directory using normal repo-local mkdir behavior."""
        stem = _sanitize_tmp_name(basename)
        if not numbered:
            path = self._base_dir / stem
            path.mkdir(parents=True, exist_ok=True)
            return path

        counter = self._counters.get(stem, 0)
        while True:
            path = self._base_dir / f"{stem}_{counter:03d}"
            if not path.exists():
                path.mkdir(parents=True, exist_ok=False)
                self._counters[stem] = counter + 1
                return path
            counter += 1


@pytest.fixture(scope="session")
def tmp_path_factory(project_root) -> RepoTempPathFactory:
    """Provide a repo-local tmp path factory that stays under build/pytest."""
    safe_root = Path(project_root) / "build" / "pytest" / "tmp-fixtures"
    safe_root.mkdir(parents=True, exist_ok=True)
    _prune_safe_tmp_sessions(safe_root)

    session_name = datetime.now().strftime("session_%Y%m%d_%H%M%S")
    session_root = safe_root / f"{session_name}_{os.getpid()}"
    session_root.mkdir(parents=True, exist_ok=False)
    return RepoTempPathFactory(session_root)


@pytest.fixture
def tmp_path(request, tmp_path_factory: RepoTempPathFactory) -> Path:
    """Return a per-test repo-local temp directory without using pytest's tmpdir plugin."""
    return tmp_path_factory.mktemp(request.node.nodeid, numbered=True)


# ============================================================
# Path Fixtures
# ============================================================

@pytest.fixture(scope="session")
def project_root():
    """Return the project root directory."""
    return os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


@pytest.fixture(scope="session")
def test_data_dir(project_root):
    """Return the test data directory path."""
    return os.path.join(project_root, "data")


@pytest.fixture
def output_dir(tmp_path):
    """Return an isolated output directory path for a single test."""
    path = tmp_path / "output"
    path.mkdir(parents=True, exist_ok=True)
    return str(path)


# ============================================================
# Test Data File Fixtures
# ============================================================

@pytest.fixture(scope="session")
def sample_input_file(test_data_dir):
    """
    Return path to sample input file for testing.
    Uses feature_matrix_with_qc_non_group_AfterVBA.xlsx as it contains QC samples.
    """
    file_path = os.path.join(test_data_dir, "feature_matrix_with_qc_non_group_AfterVBA.xlsx")
    if not os.path.exists(file_path):
        pytest.skip(f"Test data file not found: {file_path}")
    return file_path


@pytest.fixture(scope="session")
def sample_input_file_no_qc(test_data_dir):
    """
    Return path to sample input file without QC samples.
    Uses feature_matrix_control_exposed_AfterVBA.xlsx.
    """
    file_path = os.path.join(test_data_dir, "feature_matrix_control_exposed_AfterVBA.xlsx")
    if not os.path.exists(file_path):
        pytest.skip(f"Test data file not found: {file_path}")
    return file_path


# ============================================================
# Temporary Output Fixtures
# ============================================================

@pytest.fixture
def temp_output_dir():
    """
    Create a temporary directory for test outputs.
    Automatically cleaned up after test.
    """
    temp_dir = tempfile.mkdtemp(prefix="test_normalization_")
    yield temp_dir
    # Cleanup after test
    shutil.rmtree(temp_dir, ignore_errors=True)


@pytest.fixture
def clean_output_dir(output_dir):
    """
    Return output directory and track files created during test.
    Does NOT delete existing files, only tracks new ones for verification.
    """
    # Record existing files before test
    existing_files = set()
    if os.path.exists(output_dir):
        for root, dirs, files in os.walk(output_dir):
            for f in files:
                existing_files.add(os.path.join(root, f))

    yield output_dir, existing_files


@pytest.fixture
def workbook_sheet_names():
    """Return a function that lists worksheet names for a workbook."""

    def _sheet_names(file_path):
        workbook = load_workbook(file_path, read_only=True, data_only=True)
        try:
            return list(workbook.sheetnames)
        finally:
            workbook.close()

    return _sheet_names


@pytest.fixture
def copy_workbook_with_extra_sheet(tmp_path):
    """Copy a workbook and inject one extra sheet for retention tests."""

    def _copy(src_path, extra_sheet_name="UnexpectedHistory"):
        src = Path(src_path)
        dest = tmp_path / f"{src.stem}_with_extra{src.suffix}"
        shutil.copy2(src, dest)
        workbook = load_workbook(dest)
        try:
            if extra_sheet_name in workbook.sheetnames:
                del workbook[extra_sheet_name]
            worksheet = workbook.create_sheet(extra_sheet_name)
            worksheet["A1"] = "should not be copied downstream"
            workbook.save(dest)
        finally:
            workbook.close()
        return str(dest)

    return _copy


# ============================================================
# Module Import Fixtures
# ============================================================

@pytest.fixture(scope="session")
def istd_module():
    """Import and return ISTD processor module."""
    from metabolomics.processors import istd
    return istd


@pytest.fixture(scope="session")
def qc_lowess_module():
    """Import and return QC-LOWESS processor module."""
    from metabolomics.processors import qc_lowess
    return qc_lowess


@pytest.fixture(scope="session")
def qc_batch_scaling_module():
    """Import and return QC batch scaling processor module."""
    from metabolomics.processors import qc_batch_scaling
    return qc_batch_scaling


@pytest.fixture(scope="session")
def conc_norm_module():
    """Import and return Concentration Normalization processor module."""
    from metabolomics.processors import normalization
    return normalization


# ============================================================
# Validation Helpers
# ============================================================

@pytest.fixture
def validate_excel_output():
    """Return a function to validate Excel output files."""
    import pandas as pd

    def _validate(file_path, required_sheets=None, min_rows=1):
        """
        Validate an Excel output file.

        Args:
            file_path: Path to Excel file
            required_sheets: List of required sheet names (optional)
            min_rows: Minimum number of data rows expected

        Returns:
            dict with validation results
        """
        result = {
            'exists': os.path.exists(file_path),
            'readable': False,
            'sheets': [],
            'row_counts': {},
            'errors': []
        }

        if not result['exists']:
            result['errors'].append(f"File not found: {file_path}")
            return result

        try:
            excel_file = pd.ExcelFile(file_path)
            result['readable'] = True
            result['sheets'] = excel_file.sheet_names

            for sheet in excel_file.sheet_names:
                df = pd.read_excel(excel_file, sheet_name=sheet)
                result['row_counts'][sheet] = len(df)

            if required_sheets:
                missing = set(required_sheets) - set(result['sheets'])
                if missing:
                    result['errors'].append(f"Missing sheets: {missing}")

            for sheet, count in result['row_counts'].items():
                if count < min_rows:
                    result['errors'].append(f"Sheet '{sheet}' has {count} rows, expected >= {min_rows}")

        except Exception as e:
            result['errors'].append(f"Error reading file: {str(e)}")

        return result

    return _validate


@pytest.fixture
def validate_processing_result():
    """Return a function that validates the processor result contract."""

    def _validate(result, required_keys=None):
        """
        Validate a result from module main().

        Args:
            result: The ProcessingResult returned by a processor
            required_keys: List of required keys

        Returns:
            dict with validation results
        """
        from metabolomics.utils.results import ProcessingResult

        is_processing_result = isinstance(result, ProcessingResult)
        data = result.to_dict() if is_processing_result else None

        validation = {
            'is_processing_result': is_processing_result,
            'keys': list(data.keys()) if isinstance(data, dict) else [],
            'errors': []
        }

        if data is None:
            validation['errors'].append(f"Result is not a ProcessingResult: {type(result)}")
            return validation

        if required_keys:
            missing = set(required_keys) - set(data.keys())
            if missing:
                validation['errors'].append(f"Missing keys: {missing}")

        # Check output_path if present
        if 'output_path' in data:
            if not os.path.exists(data['output_path']):
                validation['errors'].append(f"output_path does not exist: {data['output_path']}")

        return validation

    return _validate


# ============================================================
# Pipeline Fixture (for integration tests)
# ============================================================

@pytest.fixture(scope="session")
def run_full_pipeline(sample_input_file, istd_module, qc_lowess_module,
                      qc_batch_scaling_module, conc_norm_module):
    """
    Run the full 4-step pipeline once and cache results.
    Used for integration tests.
    """
    from metabolomics.utils.file_io import create_session_dir, get_output_root

    results = {}
    session_dir = create_session_dir(output_root=get_output_root(input_file=sample_input_file))

    # Step 1: ISTD Correction
    result1 = istd_module.main(input_file=sample_input_file, session_dir=session_dir)
    results['step1'] = result1

    if result1 and hasattr(result1, "output_path"):
        # Step 2: QC-LOWESS
        result2 = qc_lowess_module.main(input_file=result1.output_path, session_dir=session_dir)
        results['step2'] = result2

        if result2 and hasattr(result2, "output_path"):
            # Step 3: Concentration Normalization
            result3 = conc_norm_module.main(input_file=result2.output_path, session_dir=session_dir)
            results['step3'] = result3

            if result3 and hasattr(result3, "output_path"):
                # Step 4: QC Batch Scaling
                result4 = qc_batch_scaling_module.main(
                    input_file=result3.output_path,
                    session_dir=session_dir,
                    diagnostics_only=True,
                )
                results['step4'] = result4

    return results
