"""Truth-bearing synthetic correction inputs for scientific regression tests."""

from __future__ import annotations

import argparse
from dataclasses import dataclass
from hashlib import sha256
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font


QC_POSITIONS = (1, 4, 7, 11, 14, 17, 20, 24)
STREAM_NAMES = (
    "schedule",
    "biology",
    "measurement",
    "detection",
    "outliers",
    "carryover",
    "feature_assignments",
)


@dataclass(frozen=True)
class SimulationTruth:
    step2_counterfactual: np.ndarray
    component_matrices: dict[str, np.ndarray]
    missing_reason: np.ndarray
    outlier_mask: np.ndarray
    component_fingerprints: dict[str, str]
    rng_stream_ids: dict[str, str]
    pooled_qc_metadata: dict[str, object]
    routing_truth: pd.DataFrame | None = None


@dataclass(frozen=True)
class SimulationResult:
    recipe: str
    seed: int
    observed: np.ndarray
    sample_info: pd.DataFrame
    feature_info: pd.DataFrame
    truth: SimulationTruth


def _named_rngs(seed: int) -> dict[str, np.random.Generator]:
    children = np.random.SeedSequence(seed).spawn(len(STREAM_NAMES))
    return {
        name: np.random.default_rng(child)
        for name, child in zip(STREAM_NAMES, children, strict=True)
    }


def _build_sample_info() -> pd.DataFrame:
    rows: list[dict[str, object]] = []
    qc_index = exposure_index = control_index = 1
    injection_order = 1
    for batch in ("A", "B", "C"):
        study_types = ["Exposure", "Control"] * 8
        study_cursor = 0
        for local_position in range(1, 25):
            if local_position in QC_POSITIONS:
                sample_name = f"QC{qc_index}"
                sample_type = "QC"
                qc_index += 1
            else:
                sample_type = study_types[study_cursor]
                study_cursor += 1
                if sample_type == "Exposure":
                    sample_name = f"Exposure_{exposure_index}"
                    exposure_index += 1
                else:
                    sample_name = f"Control_{control_index}"
                    control_index += 1
            rows.append(
                {
                    "Sample_Name": sample_name,
                    "Sample_Type": sample_type,
                    "Injection_Order": injection_order,
                    "Batch": batch,
                    "Injection_Volume": 5.0,
                    "Creatinine_mg_dL": (
                        np.nan if sample_type == "QC" else 100.0
                    ),
                    "Local_Position": local_position,
                }
            )
            injection_order += 1
    return pd.DataFrame(rows)


def _build_feature_info(rng: np.random.Generator) -> pd.DataFrame:
    n_analytes = 36
    n_istd = 8
    n_features = n_analytes + n_istd
    mz = np.sort(rng.uniform(350.0, 980.0, size=n_features))
    rt = rng.uniform(5.0, 40.0, size=n_features)
    assignments = np.array(
        ["null"] * 24
        + ["positive"] * 6
        + ["negative"] * 6
        + ["istd"] * n_istd,
        dtype=object,
    )
    analyte_assignments = assignments[:n_analytes].copy()
    rng.shuffle(analyte_assignments)
    assignments[:n_analytes] = analyte_assignments
    drift_kind = np.resize(
        np.array(["none", "linear", "nonlinear"], dtype=object),
        n_features,
    )
    return pd.DataFrame(
        {
            "FeatureID": [
                f"{mz_value:.4f}/{rt_value:.2f}"
                for mz_value, rt_value in zip(mz, rt, strict=True)
            ],
            "mz": mz,
            "rt": rt,
            "is_istd": np.arange(n_features) >= n_analytes,
            "biology_stratum": assignments,
            "drift_kind": drift_kind,
        }
    )


def _fingerprint(values: np.ndarray) -> str:
    finite = np.nan_to_num(values, nan=-1.0, posinf=1e308, neginf=-1e308)
    digest = sha256()
    digest.update(str(finite.shape).encode("ascii"))
    digest.update(finite.tobytes())
    return digest.hexdigest()


def _apply_qc_limited_routing(result: SimulationResult) -> SimulationResult:
    observed = result.observed.copy()
    step2_counterfactual = result.truth.step2_counterfactual.copy()
    missing_reason = result.truth.missing_reason.copy()
    routing_mask = np.ones_like(observed, dtype=bool)
    qc_outlier_factor = np.ones_like(observed)
    outlier_mask = result.truth.outlier_mask.copy()
    records: list[dict[str, object]] = []

    for batch_index, (batch, batch_info) in enumerate(
        result.sample_info.groupby("Batch", sort=False)
    ):
        qc_rows = batch_info.index[
            batch_info["Sample_Type"].eq("QC")
        ].to_numpy()
        for feature_index, feature in result.feature_info.iterrows():
            task_index = feature_index * 3 + batch_index
            effective_count = 4 + task_index % 5
            if effective_count == len(qc_rows):
                retained_rows = qc_rows
            elif task_index % 2 == 0:
                retained_rows = np.concatenate(
                    ([qc_rows[0]], qc_rows[1:-1][: effective_count - 2], [qc_rows[-1]])
                )
            else:
                retained_rows = np.concatenate(
                    (qc_rows[1:-1][: effective_count - 1], [qc_rows[-1]])
                )

            qc_outlier_count = 0
            if effective_count == len(qc_rows) and task_index % 10 == 4:
                qc_outlier_row = qc_rows[len(qc_rows) // 2]
                qc_outlier_factor[qc_outlier_row, feature_index] = 4.0
                step2_counterfactual[qc_outlier_row, feature_index] *= 4.0
                outlier_mask[qc_outlier_row, feature_index] = True
                qc_outlier_count = 1

            observed[qc_rows, feature_index] = (
                step2_counterfactual[qc_rows, feature_index]
                * result.truth.component_matrices["drift_factor"][
                    qc_rows, feature_index
                ]
            )
            missing_reason[qc_rows, feature_index] = ""
            removed_rows = np.setdiff1d(qc_rows, retained_rows)
            observed[removed_rows, feature_index] = np.nan
            missing_reason[removed_rows, feature_index] = "qc_controlled_missingness"
            routing_mask[removed_rows, feature_index] = False

            endpoint_valid = bool(
                qc_rows[0] in retained_rows and qc_rows[-1] in retained_rows
            )
            drift_kind = str(feature["drift_kind"])
            if effective_count < 6 or not endpoint_valid:
                evidence_class = "observable_insufficient"
            elif drift_kind == "none":
                evidence_class = "characterize_no_drift"
            elif drift_kind == "linear":
                evidence_class = "characterize_linear"
            elif effective_count == 8:
                evidence_class = "lowess_eligible"
            else:
                evidence_class = "characterize_nonlinear"
            records.append(
                {
                    "FeatureID": feature["FeatureID"],
                    "Batch": batch,
                    "scheduled_qc_count": len(qc_rows),
                    "effective_qc_count": effective_count,
                    "endpoint_valid": endpoint_valid,
                    "drift_kind": drift_kind,
                    "evidence_class": evidence_class,
                    "qc_outlier_count": qc_outlier_count,
                }
            )

    components = dict(result.truth.component_matrices)
    components["qc_routing_observed_mask"] = routing_mask
    components["qc_outlier_factor"] = qc_outlier_factor
    truth = SimulationTruth(
        step2_counterfactual=step2_counterfactual,
        component_matrices=components,
        missing_reason=missing_reason,
        outlier_mask=outlier_mask,
        component_fingerprints={
            name: _fingerprint(values) for name, values in components.items()
        },
        rng_stream_ids=result.truth.rng_stream_ids,
        pooled_qc_metadata=result.truth.pooled_qc_metadata,
        routing_truth=pd.DataFrame(records),
    )
    return SimulationResult(
        recipe="qc_limited_routing",
        seed=result.seed,
        observed=observed,
        sample_info=result.sample_info,
        feature_info=result.feature_info,
        truth=truth,
    )


def generate_simulation(
    recipe: str,
    *,
    seed: int = 51,
    variant: str | None = None,
) -> SimulationResult:
    """Generate one reviewed synthetic scenario and its reconstructable truth."""
    if recipe not in ("routine_recoverable", "qc_limited_routing"):
        raise ValueError(f"Unknown vNext recipe: {recipe}")
    if recipe == "qc_limited_routing" and variant not in (None, "baseline"):
        raise ValueError(f"Unsupported qc_limited_routing variant: {variant}")
    if variant not in (None, "baseline", "drift_low", "drift_high"):
        raise ValueError(f"Unsupported routine_recoverable variant: {variant}")
    drift_scale = {
        None: 1.0,
        "baseline": 1.0,
        "drift_low": 0.45,
        "drift_high": 1.8,
    }[variant]

    rngs = _named_rngs(seed)
    sample_info = _build_sample_info()
    feature_info = _build_feature_info(rngs["feature_assignments"])
    n_samples = len(sample_info)
    n_features = len(feature_info)
    study_mask = sample_info["Sample_Type"].ne("QC").to_numpy()
    qc_mask = ~study_mask

    biology_rng = rngs["biology"]
    base_levels = np.exp(
        biology_rng.uniform(np.log(2.0e4), np.log(8.0e5), size=n_features)
    )
    study_abundance = np.tile(base_levels, (int(study_mask.sum()), 1))
    study_abundance *= biology_rng.lognormal(
        mean=0.0,
        sigma=0.10,
        size=study_abundance.shape,
    )
    study_types = sample_info.loc[study_mask, "Sample_Type"].to_numpy()
    exposure_rows = study_types == "Exposure"
    positive = feature_info["biology_stratum"].eq("positive").to_numpy()
    negative = feature_info["biology_stratum"].eq("negative").to_numpy()
    study_abundance[np.ix_(exposure_rows, positive)] *= 2.0
    study_abundance[np.ix_(exposure_rows, negative)] *= 0.5

    pooled_qc = np.mean(study_abundance, axis=0)
    latent_injection = np.empty((n_samples, n_features), dtype=float)
    latent_injection[study_mask] = study_abundance
    latent_injection[qc_mask] = pooled_qc

    batch_lookup = {"A": 1.0, "B": 1.12, "C": 0.90}
    batch_factor = np.array(
        [batch_lookup[batch] for batch in sample_info["Batch"]],
        dtype=float,
    )[:, None]
    batch_factor = np.broadcast_to(batch_factor, latent_injection.shape).copy()

    measurement_rng = rngs["measurement"]
    ionization = measurement_rng.lognormal(mean=0.0, sigma=0.06, size=n_samples)
    ionization_factor = np.broadcast_to(
        ionization[:, None], latent_injection.shape
    ).copy()
    noise_sigma = np.where(qc_mask, 0.035, 0.10)
    measurement_noise = measurement_rng.lognormal(
        mean=0.0,
        sigma=noise_sigma[:, None],
        size=latent_injection.shape,
    )

    drift_factor = np.ones_like(latent_injection)
    for batch, batch_info in sample_info.groupby("Batch", sort=False):
        row_indices = batch_info.index.to_numpy()
        position = np.linspace(-1.0, 1.0, len(row_indices))
        for feature_index, drift_kind in enumerate(feature_info["drift_kind"]):
            if drift_kind == "linear":
                drift_factor[row_indices, feature_index] = (
                    1.0 + drift_scale * 0.14 * position
                )
            elif drift_kind == "nonlinear":
                drift_factor[row_indices, feature_index] = (
                    1.0
                    + drift_scale * 0.16 * position
                    + drift_scale * 0.07 * (position**2 - 1.0 / 3.0)
                )

    carryover_factor = np.ones_like(latent_injection)
    carryover_rng = rngs["carryover"]
    analyte_indices = np.where(~feature_info["is_istd"].to_numpy())[0]
    for _, batch_info in sample_info.groupby("Batch", sort=False):
        batch_rows = batch_info.index.to_numpy()
        source_offsets = carryover_rng.choice(
            np.arange(len(batch_rows) - 2),
            size=2,
            replace=False,
        )
        for source_offset in source_offsets:
            affected_features = carryover_rng.choice(
                analyte_indices,
                size=4,
                replace=False,
            )
            carryover_factor[
                batch_rows[source_offset + 1], affected_features
            ] *= 1.12
            carryover_factor[
                batch_rows[source_offset + 2], affected_features
            ] *= 1.05

    step2_counterfactual = (
        latent_injection
        * batch_factor
        * ionization_factor
        * measurement_noise
        * carryover_factor
    )

    outlier_factor = np.ones_like(latent_injection)
    outlier_mask = np.zeros_like(latent_injection, dtype=bool)
    eligible_rows = np.where(study_mask)[0]
    eligible_features = np.where(~feature_info["is_istd"].to_numpy())[0]
    for _ in range(8):
        row = int(rngs["outliers"].choice(eligible_rows))
        column = int(rngs["outliers"].choice(eligible_features))
        outlier_mask[row, column] = True
        outlier_factor[row, column] = float(
            rngs["outliers"].choice([0.25, 4.0])
        )
    step2_counterfactual *= outlier_factor
    pre_detection = step2_counterfactual * drift_factor

    log_signal = np.log2(pre_detection)
    lod_center = np.quantile(log_signal, 0.08, axis=0)
    detection_probability = 0.78 + 0.21 / (
        1.0 + np.exp(-(log_signal - lod_center[None, :]) * 2.0)
    )
    detection_uniform = rngs["detection"].random(pre_detection.shape)
    detected = detection_uniform <= detection_probability
    detected[qc_mask] = True
    observed = pre_detection.copy()
    observed[~detected] = np.nan
    missing_reason = np.full(observed.shape, "", dtype=object)
    missing_reason[~detected] = "below_detection_probability"

    component_matrices = {
        "latent_injection_abundance": latent_injection,
        "batch_factor": batch_factor,
        "ionization_factor": ionization_factor,
        "drift_factor": drift_factor,
        "measurement_noise": measurement_noise,
        "carryover_factor": carryover_factor,
        "matrix_effect_factor": np.ones_like(latent_injection),
        "saturation_factor": np.ones_like(latent_injection),
        "outlier_factor": outlier_factor,
        "detection_uniform": detection_uniform,
    }
    fingerprints = {
        name: _fingerprint(values)
        for name, values in component_matrices.items()
    }
    truth = SimulationTruth(
        step2_counterfactual=step2_counterfactual,
        component_matrices=component_matrices,
        missing_reason=missing_reason,
        outlier_mask=outlier_mask,
        component_fingerprints=fingerprints,
        rng_stream_ids={
            name: f"root={seed}:child={index}"
            for index, name in enumerate(STREAM_NAMES)
        },
        pooled_qc_metadata={
            "contributor_sample_names": tuple(
                sample_info.loc[study_mask, "Sample_Name"]
            ),
            "equal_volume_weight": 1.0 / int(study_mask.sum()),
            "dilution_factor": 1.0,
        },
    )
    result = SimulationResult(
        recipe="routine_recoverable",
        seed=seed,
        observed=observed,
        sample_info=sample_info,
        feature_info=feature_info,
        truth=truth,
    )
    if recipe == "qc_limited_routing":
        return _apply_qc_limited_routing(result)
    return result


def write_simulation_workbook(
    result: SimulationResult,
    output_path: str | Path,
) -> Path:
    """Write the workflow-facing workbook without imputing missing values."""
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    sample_names = result.sample_info["Sample_Name"].tolist()
    feature_rows = pd.DataFrame(result.observed.T, columns=sample_names)
    feature_rows.insert(0, "FeatureID", result.feature_info["FeatureID"].tolist())
    sample_type_row = {"FeatureID": "Sample_Type"}
    sample_type_row.update(
        dict(
            zip(
                result.sample_info["Sample_Name"],
                result.sample_info["Sample_Type"],
                strict=True,
            )
        )
    )
    raw = pd.concat([pd.DataFrame([sample_type_row]), feature_rows], ignore_index=True)
    sample_columns = [
        "Sample_Name",
        "Sample_Type",
        "Injection_Order",
        "Batch",
        "Injection_Volume",
        "Creatinine_mg_dL",
    ]
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        raw.to_excel(writer, sheet_name="RawIntensity", index=False)
        result.sample_info[sample_columns].to_excel(
            writer,
            sheet_name="SampleInfo",
            index=False,
        )

    workbook = load_workbook(output_path)
    try:
        sheet = workbook["RawIntensity"]
        istd_ids = set(
            result.feature_info.loc[result.feature_info["is_istd"], "FeatureID"]
        )
        for row in range(3, sheet.max_row + 1):
            if sheet.cell(row=row, column=1).value in istd_ids:
                sheet.cell(row=row, column=1).font = Font(color="FFFF0000")
        workbook.save(output_path)
    finally:
        workbook.close()
    return output_path


def semantic_workbook_digest(path: str | Path) -> str:
    """Hash workbook meaning, excluding ZIP metadata and irrelevant styling."""
    workbook = load_workbook(path, read_only=False, data_only=True)
    digest = sha256()
    try:
        for sheet_name in workbook.sheetnames:
            digest.update(f"sheet:{sheet_name}\n".encode())
            sheet = workbook[sheet_name]
            for row in sheet.iter_rows():
                values: list[str] = []
                for cell in row:
                    value = cell.value
                    if value is None or (isinstance(value, float) and np.isnan(value)):
                        values.append("<NA>")
                    elif isinstance(value, (int, float)):
                        values.append(f"{float(value):.12g}")
                    else:
                        values.append(str(value))
                digest.update(("\t".join(values) + "\n").encode("utf-8"))
            if sheet_name == "RawIntensity":
                for row in range(3, sheet.max_row + 1):
                    cell = sheet.cell(row=row, column=1)
                    color = cell.font.color
                    rgb = color.rgb if color is not None and color.type == "rgb" else None
                    if rgb and str(rgb).upper().endswith("FF0000"):
                        digest.update(f"istd:{cell.value}\n".encode("utf-8"))
    finally:
        workbook.close()
    return digest.hexdigest()


def _log_rmse(observed: np.ndarray, target: np.ndarray) -> float:
    valid = (
        np.isfinite(observed)
        & np.isfinite(target)
        & (observed > 0)
        & (target > 0)
    )
    if not valid.any():
        return np.nan
    error = np.log2(observed[valid]) - np.log2(target[valid])
    return float(np.sqrt(np.mean(error**2)))


def _group_log2_fold_change(
    values: np.ndarray,
    sample_info: pd.DataFrame,
) -> np.ndarray:
    exposure = sample_info["Sample_Type"].eq("Exposure").to_numpy()
    control = sample_info["Sample_Type"].eq("Control").to_numpy()
    exposure_median = np.nanmedian(values[exposure], axis=0)
    control_median = np.nanmedian(values[control], axis=0)
    with np.errstate(divide="ignore", invalid="ignore"):
        return np.log2(exposure_median / control_median)


def evaluate_step2_candidate(
    result: SimulationResult,
    corrected: np.ndarray,
    *,
    false_correction_tolerance_log2_rmse: float = 0.05,
) -> dict[str, float]:
    """Score a Step 2 candidate against its immediate counterfactual truth."""
    corrected = np.asarray(corrected, dtype=float)
    if corrected.shape != result.observed.shape:
        raise ValueError(
            f"Corrected matrix shape {corrected.shape} does not match "
            f"observed shape {result.observed.shape}"
        )
    truth = result.truth.step2_counterfactual
    recovery_gains: list[float] = []
    false_corrections: list[bool] = []
    for batch, batch_info in result.sample_info.groupby("Batch", sort=False):
        del batch
        rows = batch_info.index.to_numpy()
        for feature_index, drift_kind in enumerate(result.feature_info["drift_kind"]):
            raw_values = result.observed[rows, feature_index]
            corrected_values = corrected[rows, feature_index]
            truth_values = truth[rows, feature_index]
            common_support = (
                np.isfinite(raw_values)
                & np.isfinite(corrected_values)
                & np.isfinite(truth_values)
                & (raw_values > 0)
                & (corrected_values > 0)
                & (truth_values > 0)
            )
            raw_rmse = _log_rmse(
                raw_values[common_support],
                truth_values[common_support],
            )
            corrected_rmse = _log_rmse(
                corrected_values[common_support],
                truth_values[common_support],
            )
            if not np.isfinite(raw_rmse) or not np.isfinite(corrected_rmse):
                continue
            if drift_kind == "none":
                false_corrections.append(
                    corrected_rmse - raw_rmse
                    > false_correction_tolerance_log2_rmse
                )
            elif raw_rmse > 1e-12:
                recovery_gains.append(1.0 - corrected_rmse / raw_rmse)

    truth_on_candidate_support = truth.copy()
    truth_on_candidate_support[
        ~np.isfinite(result.observed) | ~np.isfinite(corrected)
    ] = np.nan
    corrected_on_common_support = corrected.copy()
    corrected_on_common_support[~np.isfinite(result.observed)] = np.nan
    truth_fc = _group_log2_fold_change(
        truth_on_candidate_support,
        result.sample_info,
    )
    corrected_fc = _group_log2_fold_change(
        corrected_on_common_support,
        result.sample_info,
    )
    analyte = ~result.feature_info["is_istd"].to_numpy()
    eligible = analyte & np.isfinite(truth_fc) & np.isfinite(corrected_fc)
    fc_errors = np.abs(corrected_fc[eligible] - truth_fc[eligible])
    non_null = (
        eligible
        & result.feature_info["biology_stratum"].isin(["positive", "negative"]).to_numpy()
    )
    sign_flips = np.sign(corrected_fc[non_null]) != np.sign(truth_fc[non_null])
    return {
        "median_technical_recovery_gain": float(np.median(recovery_gains)),
        "false_correction_rate": float(np.mean(false_corrections)),
        "biology_sign_flip_rate": float(np.mean(sign_flips)),
        "median_abs_log2_fold_change_error": float(np.median(fc_errors)),
    }


def _build_argument_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Generate a truth-backed DNP correction workbook.",
    )
    parser.add_argument(
        "--recipe",
        choices=("routine_recoverable", "qc_limited_routing"),
        default="routine_recoverable",
    )
    parser.add_argument("--seed", type=int, default=51)
    parser.add_argument(
        "--variant",
        choices=("baseline", "drift_low", "drift_high"),
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("data/synthetic_correction_input.xlsx"),
    )
    return parser


def main() -> None:
    args = _build_argument_parser().parse_args()
    result = generate_simulation(
        args.recipe,
        seed=args.seed,
        variant=args.variant,
    )
    output_path = write_simulation_workbook(result, args.output)
    print(f"Saved {args.recipe} seed={args.seed}: {output_path}")
    print(f"Semantic digest: {semantic_workbook_digest(output_path)}")


if __name__ == "__main__":
    main()
