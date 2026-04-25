import argparse
from dataclasses import dataclass, fields, replace
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font


SEED = 51
INJECTION_VOLUME = 20


@dataclass(frozen=True)
class ScenarioConfig:
    name: str
    description: str
    recommended_steps: tuple[str, ...] = ()
    primary_checks: tuple[str, ...] = ()
    n_features: int = 64
    n_istd: int = 12
    batches: tuple[str, ...] = ("A", "B", "C")
    injections_per_batch: int = 24
    qc_positions_within_batch: tuple[int, ...] = (1, 5, 9, 13, 17, 21)
    batch_factors: tuple[float, ...] = (1.00, 1.20, 0.88)
    residual_batch_factors: tuple[float, ...] = (1.00, 1.12, 0.92)
    drift_strength: float = 0.18
    drift_profile: str = "linear"
    ionization_sigma: float = 0.08
    analyte_noise_sigma: float = 0.12
    qc_noise_sigma: float = 0.05
    istd_noise_sigma: float = 0.018
    istd_tracking_strength: float = 1.0
    istd_order_decay_strength: float = 0.0
    differential_feature_count: int = 14
    residual_batch_feature_count: int = 12
    outlier_count: int = 18
    high_jump_range: tuple[float, float] = (3.5, 14.0)
    low_jump_range: tuple[float, float] = (0.08, 0.4)
    carryover_event_count: int = 0
    carryover_feature_fraction: float = 0.0
    carryover_decay_factors: tuple[float, ...] = ()
    analyte_missing_range: tuple[float, float] = (0.02, 0.16)
    istd_missing_range: tuple[float, float] = (0.0, 0.015)
    structured_missing_rate: float = 0.0
    structured_missing_batch: str = ""
    structured_missing_sample_type: str = ""
    matrix_effect_sample_type: str = ""
    matrix_effect_feature_fraction: float = 0.0
    matrix_effect_strength: float = 1.0
    istd_interference_sample_type: str = ""
    istd_interference_feature_fraction: float = 0.0
    istd_interference_strength: float = 1.0
    batch_drift_multipliers: tuple[float, ...] = (1.0, 1.0, 1.0)
    saturation_feature_fraction: float = 0.0
    saturation_threshold_quantile: float = 0.82
    saturation_power: float = 1.0
    order_pattern: str = "alternating"
    creatinine_mode: str = "random"
    creatinine_noise_sigma: float = 0.08


BASE_CONFIG = ScenarioConfig(
    name="balanced_pipeline",
    description="Balanced multi-batch matrix with stable ISTDs, moderate drift, and usable PQN behavior.",
    recommended_steps=("Step 1", "Step 2", "Step 3 PQN"),
    primary_checks=(
        "End-to-end smoke test across the whole pipeline.",
        "Step 1 should not skip and should generate all diagnostics.",
        "Step 3 PQN should run on a realistic but not pathological matrix.",
    ),
)


SCENARIO_LIBRARY = {
    "balanced_pipeline": BASE_CONFIG,
    "strong_drift": replace(
        BASE_CONFIG,
        name="strong_drift",
        description="Long-run instrument drift dominates QC trajectories and LOWESS should have obvious work to do.",
        recommended_steps=("Step 2",),
        primary_checks=(
            "QC trend plots should show clear drift over injection order.",
            "LOWESS should reduce QC CV and flatten drift trajectories.",
        ),
        drift_strength=0.40,
        analyte_noise_sigma=0.14,
    ),
    "nonlinear_drift": replace(
        BASE_CONFIG,
        name="nonlinear_drift",
        description="Piecewise drift simulates warm-up, plateau, and late-run recovery instead of a simple linear trend.",
        recommended_steps=("Step 2",),
        primary_checks=(
            "LOWESS should handle non-linear QC trajectories without crashing.",
            "Trend plots should show a visibly non-linear drift shape over injection order.",
        ),
        drift_strength=0.34,
        drift_profile="piecewise",
        analyte_noise_sigma=0.15,
    ),
    "random_jump": replace(
        BASE_CONFIG,
        name="random_jump",
        description="Frequent abrupt spikes and drops simulate intermittent spray instability and random jump noise.",
        recommended_steps=("Step 2", "Step 3 PQN"),
        primary_checks=(
            "Stress robustness to sudden spikes and dips.",
            "Check whether QC-focused methods avoid overfitting isolated jumps.",
        ),
        analyte_noise_sigma=0.20,
        outlier_count=40,
        high_jump_range=(5.0, 18.0),
        low_jump_range=(0.03, 0.35),
    ),
    "carryover_memory": replace(
        BASE_CONFIG,
        name="carryover_memory",
        description="Localized memory effects trail after strong injections, creating short carryover tails rather than smooth drift.",
        recommended_steps=("Step 2",),
        primary_checks=(
            "Diagnostics should show local contamination clusters after strong injections.",
            "Carryover should not be mistaken for a smooth whole-run trend.",
        ),
        analyte_noise_sigma=0.15,
        carryover_event_count=8,
        carryover_feature_fraction=0.18,
        carryover_decay_factors=(0.18, 0.07),
    ),
    "strong_batch": replace(
        BASE_CONFIG,
        name="strong_batch",
        description="Large batch offsets persist after Step 1/2 and should be very visible in manual Step 4 diagnostics.",
        recommended_steps=("Step 3 PQN", "Step 4 diagnostics (manual)"),
        primary_checks=(
            "Manual Step 4 batch diagnostics should show clear pre/post separation when explicitly run.",
            "Residual analysis should visibly tighten after manual QC batch scaling diagnostics.",
        ),
        batch_factors=(1.00, 1.45, 0.70),
        residual_batch_factors=(1.00, 1.22, 0.78),
        residual_batch_feature_count=22,
    ),
    "mixed_direction_batch_drift": replace(
        BASE_CONFIG,
        name="mixed_direction_batch_drift",
        description="QC trends drift in different directions per batch, stressing Step 2/3 separation of within-batch drift and cross-batch offsets.",
        recommended_steps=("Step 2", "Step 3"),
        primary_checks=(
            "QC trajectories should slope upward in one batch and downward in another.",
            "Step 3 should still produce batch-alignment diagnostics after Step 2.",
        ),
        drift_strength=0.24,
        batch_drift_multipliers=(1.10, -0.95, 0.15),
        residual_batch_factors=(1.00, 1.18, 0.86),
        residual_batch_feature_count=16,
    ),
    "unstable_istd": replace(
        BASE_CONFIG,
        name="unstable_istd",
        description="Internal standards are noisy, partly decoupled from shared sample factors, and contain more missingness.",
        recommended_steps=("Step 1",),
        primary_checks=(
            "Step 1 ISTD tracking should reveal unstable correction compounds.",
            "Use to verify Step 1 gate / quality diagnostics under weak ISTD conditions.",
        ),
        n_istd=10,
        istd_noise_sigma=0.12,
        istd_tracking_strength=0.40,
        istd_missing_range=(0.03, 0.09),
    ),
    "istd_degradation": replace(
        BASE_CONFIG,
        name="istd_degradation",
        description="ISTDs decay with injection order, mimicking gradual standard degradation across the run.",
        recommended_steps=("Step 1",),
        primary_checks=(
            "ISTD tracking should show a clear downward trend over injection order.",
            "Step 1 gating should not treat these ISTDs as fully healthy.",
        ),
        istd_noise_sigma=0.03,
        istd_tracking_strength=0.75,
        istd_order_decay_strength=0.40,
    ),
    "istd_sample_interference": replace(
        BASE_CONFIG,
        name="istd_sample_interference",
        description="ISTDs look acceptable in QC but are selectively suppressed in one real sample class, mimicking matrix-specific interference.",
        recommended_steps=("Step 1",),
        primary_checks=(
            "QC-derived ISTD quality should look better than real-sample behavior.",
            "Use to inspect whether Step 1 looks overconfident when only one sample class is affected.",
        ),
        istd_noise_sigma=0.022,
        istd_interference_sample_type="Exposure",
        istd_interference_feature_fraction=0.55,
        istd_interference_strength=0.52,
    ),
    "order_confounding": replace(
        BASE_CONFIG,
        name="order_confounding",
        description="Injection order is highly confounded with sample type, useful for stress-testing order-sensitive methods.",
        recommended_steps=("Step 2", "Step 3 PQN"),
        primary_checks=(
            "Use when checking for over-correction risk when drift and biology align in time order.",
            "Review whether biological separation is unintentionally flattened after correction.",
        ),
        order_pattern="confounded",
    ),
    "specnorm_friendly": replace(
        BASE_CONFIG,
        name="specnorm_friendly",
        description="Creatinine values are linked to the dominant sample-wise dilution factor so SpecNorm+PQN normalization is meaningful.",
        recommended_steps=("Step 3 SpecNorm+PQN", "Step 3 PQN"),
        primary_checks=(
            "Preferred matrix for validating SpecNorm+PQN behavior.",
            "Compare SpecNorm+PQN against PQN on a matrix with meaningful reference values.",
        ),
        creatinine_mode="linked",
        creatinine_noise_sigma=0.05,
    ),
    "matrix_effect_suppression": replace(
        BASE_CONFIG,
        name="matrix_effect_suppression",
        description="One real sample type experiences broad ion suppression across many analytes while QC remains comparatively well-behaved.",
        recommended_steps=("Step 1", "Step 3 PQN"),
        primary_checks=(
            "Target sample type should show depressed global analyte intensity.",
            "Useful for checking whether normalization appears over-optimistic under matrix-effect bias.",
        ),
        matrix_effect_sample_type="Exposure",
        matrix_effect_feature_fraction=0.60,
        matrix_effect_strength=0.50,
        analyte_noise_sigma=0.14,
    ),
    "signal_saturation": replace(
        BASE_CONFIG,
        name="signal_saturation",
        description="Highest-intensity analytes enter a compressed response regime, simulating detector saturation and nonlinear response.",
        recommended_steps=("Step 3 PQN", "Step 3 SpecNorm+PQN"),
        primary_checks=(
            "Upper-tail intensities should be visibly compressed relative to balanced_pipeline.",
            "Normalization quality metrics should not look unrealistically perfect under nonlinear response.",
        ),
        analyte_noise_sigma=0.11,
        saturation_feature_fraction=0.22,
        saturation_threshold_quantile=0.80,
        saturation_power=0.68,
    ),
    "structured_missingness": replace(
        BASE_CONFIG,
        name="structured_missingness",
        description="Missing values cluster within one batch and sample type, stressing NaN handling in downstream summaries.",
        recommended_steps=("Step 3 PQN",),
        primary_checks=(
            "Structured NaNs should not crash batch statistics or normalization summaries.",
            "Use when validating graceful degradation under non-random missingness.",
        ),
        analyte_missing_range=(0.01, 0.08),
        structured_missing_rate=0.55,
        structured_missing_batch="B",
        structured_missing_sample_type="Exposure",
    ),
    "combined_stress": replace(
        BASE_CONFIG,
        name="combined_stress",
        description="Stacked stress profile: stronger drift, stronger batch effects, random jumps, unstable ISTDs, and order confounding.",
        recommended_steps=("Step 1", "Step 2", "Step 3 PQN"),
        primary_checks=(
            "Use as a broad regression stress test after major refactors.",
            "Expect some steps to degrade gracefully rather than look ideal.",
        ),
        drift_strength=0.38,
        analyte_noise_sigma=0.22,
        batch_factors=(1.00, 1.50, 0.68),
        residual_batch_factors=(1.00, 1.26, 0.76),
        residual_batch_feature_count=24,
        outlier_count=48,
        high_jump_range=(5.5, 20.0),
        low_jump_range=(0.02, 0.30),
        carryover_event_count=10,
        carryover_feature_fraction=0.18,
        carryover_decay_factors=(0.16, 0.06),
        istd_noise_sigma=0.14,
        istd_tracking_strength=0.35,
        istd_order_decay_strength=0.30,
        istd_missing_range=(0.05, 0.10),
        structured_missing_rate=0.35,
        structured_missing_batch="B",
        structured_missing_sample_type="Exposure",
        matrix_effect_sample_type="Exposure",
        matrix_effect_feature_fraction=0.38,
        matrix_effect_strength=0.68,
        batch_drift_multipliers=(1.0, -0.75, 0.20),
        saturation_feature_fraction=0.15,
        saturation_threshold_quantile=0.82,
        saturation_power=0.74,
        order_pattern="confounded",
    ),
}


def compose_scenario_config(scenario_expression: str) -> ScenarioConfig:
    scenario_expression = (scenario_expression or "balanced_pipeline").strip()
    if scenario_expression in SCENARIO_LIBRARY:
        return SCENARIO_LIBRARY[scenario_expression]

    tokens = [token.strip() for token in scenario_expression.replace(",", "+").split("+") if token.strip()]
    if not tokens:
        return BASE_CONFIG

    unknown = [token for token in tokens if token not in SCENARIO_LIBRARY]
    if unknown:
        raise ValueError(f"Unknown scenario token(s): {', '.join(sorted(unknown))}")

    config = BASE_CONFIG
    for token in tokens:
        profile = SCENARIO_LIBRARY[token]
        overrides = {}
        for field in fields(ScenarioConfig):
            if field.name in {"name", "description"}:
                continue
            base_value = getattr(BASE_CONFIG, field.name)
            profile_value = getattr(profile, field.name)
            if profile_value != base_value:
                overrides[field.name] = profile_value
        if overrides:
            config = replace(config, **overrides)

    return replace(
        config,
        name="_".join(tokens),
        description=f"Composed scenario from: {', '.join(tokens)}",
    )


def build_sample_info(rng: np.random.Generator, config: ScenarioConfig) -> pd.DataFrame:
    rows = []
    exposure_counter = 1
    control_counter = 1
    qc_counter = 1
    total_injection_order = 1
    real_samples_per_batch = config.injections_per_batch - len(config.qc_positions_within_batch)

    for batch in config.batches:
        if config.order_pattern == "confounded":
            per_batch_plan = ["Exposure"] * (real_samples_per_batch // 2) + ["Control"] * (real_samples_per_batch // 2)
        else:
            per_batch_plan = ["Exposure" if index % 2 == 0 else "Control" for index in range(real_samples_per_batch)]

        plan_cursor = 0
        for position in range(1, config.injections_per_batch + 1):
            if position in config.qc_positions_within_batch:
                sample_name = f"QC{qc_counter}"
                sample_type = "QC"
                qc_counter += 1
            else:
                sample_type = per_batch_plan[plan_cursor]
                if sample_type == "Exposure":
                    sample_name = f"Exposure_{exposure_counter}"
                    exposure_counter += 1
                else:
                    sample_name = f"Control_{control_counter}"
                    control_counter += 1
                plan_cursor += 1

            rows.append(
                {
                    "Sample_Name": sample_name,
                    "Sample_Type": sample_type,
                    "Injection_Order": total_injection_order,
                    "Batch": batch,
                    "Injection_Volume": INJECTION_VOLUME,
                    "Creatinine_mg_dL": np.nan,
                }
            )
            total_injection_order += 1

    return pd.DataFrame(rows)


def build_feature_metadata(rng: np.random.Generator, config: ScenarioConfig) -> pd.DataFrame:
    mz_values = np.sort(rng.uniform(350.0, 980.0, size=config.n_features))
    rt_values = np.sort(rng.uniform(5.0, 40.0, size=config.n_features))
    istd_indices = np.sort(rng.choice(np.arange(config.n_features), size=config.n_istd, replace=False))

    rows = []
    for index in range(config.n_features):
        rows.append(
            {
                "FeatureID": f"{mz_values[index]:.4f}/{rt_values[index]:.2f}",
                "mz": float(mz_values[index]),
                "rt": float(rt_values[index]),
                "is_istd": bool(index in istd_indices),
            }
        )

    return pd.DataFrame(rows)


def build_creatinine_values(
    sample_info_df: pd.DataFrame,
    shared_sample_factor: np.ndarray,
    rng: np.random.Generator,
    config: ScenarioConfig,
) -> np.ndarray:
    creatinine = np.full(len(sample_info_df), np.nan, dtype=float)
    for index, row in enumerate(sample_info_df.itertuples()):
        if row.Sample_Type == "QC":
            continue
        if config.creatinine_mode == "linked":
            value = 100.0 * shared_sample_factor[index] * rng.lognormal(mean=0.0, sigma=config.creatinine_noise_sigma)
            creatinine[index] = float(np.clip(value, 25.0, 240.0))
        else:
            creatinine[index] = round(float(rng.uniform(40.0, 180.0)), 1)
    return np.round(creatinine, 1)


def build_drift_vector(scaled_orders: np.ndarray, config: ScenarioConfig) -> np.ndarray:
    if config.drift_profile == "piecewise":
        warmup = np.clip(scaled_orders / 0.25, 0.0, 1.0)
        plateau = np.clip((scaled_orders - 0.25) / 0.40, 0.0, 1.0)
        recovery = np.clip((scaled_orders - 0.65) / 0.35, 0.0, 1.0)
        return (
            1.0
            + config.drift_strength * 1.15 * warmup
            - config.drift_strength * 0.15 * plateau
            - config.drift_strength * 0.40 * recovery
        )

    return 1.0 + config.drift_strength * scaled_orders


def build_batch_local_positions(sample_info_df: pd.DataFrame, config: ScenarioConfig) -> np.ndarray:
    local_positions = np.zeros(len(sample_info_df), dtype=float)
    for batch in config.batches:
        batch_indices = np.where(sample_info_df["Batch"].to_numpy() == batch)[0]
        if len(batch_indices) <= 1:
            continue
        local_positions[batch_indices] = np.linspace(0.0, 1.0, len(batch_indices))
    return local_positions


def simulate_intensity_matrix(
    sample_info_df: pd.DataFrame,
    feature_df: pd.DataFrame,
    rng: np.random.Generator,
    config: ScenarioConfig,
) -> tuple[np.ndarray, np.ndarray]:
    n_samples = len(sample_info_df)
    n_features = len(feature_df)

    injection_orders = sample_info_df["Injection_Order"].to_numpy(dtype=float)
    batch_labels = sample_info_df["Batch"].tolist()
    sample_types = sample_info_df["Sample_Type"].tolist()
    batch_factor_lookup = dict(zip(config.batches, config.batch_factors))
    residual_batch_lookup = dict(zip(config.batches, config.residual_batch_factors))

    batch_factor_vector = np.array([batch_factor_lookup[batch] for batch in batch_labels], dtype=float)
    scaled_orders = (injection_orders - injection_orders.min()) / (injection_orders.max() - injection_orders.min())
    batch_local_positions = build_batch_local_positions(sample_info_df, config)
    drift_vector = build_drift_vector(scaled_orders, config)
    ionization_vector = rng.lognormal(mean=0.0, sigma=config.ionization_sigma, size=n_samples)
    shared_sample_factor = batch_factor_vector * drift_vector * ionization_vector

    group_factor_vector = np.ones(n_samples, dtype=float)
    for index, sample_type in enumerate(sample_types):
        if sample_type == "Exposure":
            group_factor_vector[index] = 1.10
        elif sample_type == "Control":
            group_factor_vector[index] = 0.96

    base_feature_levels = rng.uniform(1.8e5, 8.5e5, size=n_features)
    intensities = np.zeros((n_samples, n_features), dtype=float)

    analyte_indices = feature_df.index[~feature_df["is_istd"]].to_numpy()
    differential_indices = rng.choice(analyte_indices, size=min(config.differential_feature_count, len(analyte_indices)), replace=False)
    residual_batch_pool = np.setdiff1d(analyte_indices, differential_indices)
    residual_batch_count = min(config.residual_batch_feature_count, len(residual_batch_pool))
    residual_batch_sensitive_indices = rng.choice(residual_batch_pool, size=residual_batch_count, replace=False)
    matrix_effect_count = min(
        len(analyte_indices),
        max(0, int(round(len(analyte_indices) * config.matrix_effect_feature_fraction))),
    )
    matrix_effect_indices = (
        rng.choice(analyte_indices, size=matrix_effect_count, replace=False)
        if matrix_effect_count > 0
        else np.array([], dtype=int)
    )
    istd_indices = feature_df.index[feature_df["is_istd"]].to_numpy()
    istd_interference_count = min(
        len(istd_indices),
        max(0, int(round(len(istd_indices) * config.istd_interference_feature_fraction))),
    )
    istd_interference_indices = (
        rng.choice(istd_indices, size=istd_interference_count, replace=False)
        if istd_interference_count > 0
        else np.array([], dtype=int)
    )
    saturation_feature_count = min(
        len(analyte_indices),
        max(0, int(round(len(analyte_indices) * config.saturation_feature_fraction))),
    )
    saturation_feature_indices = (
        np.argsort(base_feature_levels[analyte_indices])[-saturation_feature_count:]
        if saturation_feature_count > 0
        else np.array([], dtype=int)
    )
    saturation_feature_indices = analyte_indices[saturation_feature_indices] if saturation_feature_count > 0 else np.array([], dtype=int)
    batch_drift_lookup = dict(zip(config.batches, config.batch_drift_multipliers))

    for feature_index, feature_row in feature_df.iterrows():
        base_level = base_feature_levels[feature_index]
        if feature_row["is_istd"]:
            istd_factor = 1.0 + config.istd_tracking_strength * (shared_sample_factor - 1.0)
            if config.istd_order_decay_strength > 0:
                istd_factor *= np.clip(1.0 - config.istd_order_decay_strength * scaled_orders, 0.15, None)
            stable_noise = rng.lognormal(mean=0.0, sigma=config.istd_noise_sigma, size=n_samples)
            values = base_level * istd_factor * stable_noise
            if feature_index in istd_interference_indices and config.istd_interference_sample_type:
                target_mask = np.array([sample_type == config.istd_interference_sample_type for sample_type in sample_types], dtype=bool)
                values[target_mask] *= config.istd_interference_strength
            intensities[:, feature_index] = values
            continue

        feature_noise = rng.lognormal(mean=0.0, sigma=config.analyte_noise_sigma, size=n_samples)
        values = base_level * shared_sample_factor * feature_noise

        if feature_index in differential_indices:
            exposure_boost = rng.uniform(1.35, 1.95)
            control_shift = rng.uniform(0.88, 1.05)
            for sample_index, sample_type in enumerate(sample_types):
                if sample_type == "Exposure":
                    values[sample_index] *= exposure_boost
                elif sample_type == "Control":
                    values[sample_index] *= control_shift

        if feature_index in residual_batch_sensitive_indices:
            values *= np.array([residual_batch_lookup[batch] for batch in batch_labels], dtype=float)

        residual_drift_strength = rng.uniform(0.02, max(0.03, config.drift_strength * 0.35))
        if any(multiplier != 1.0 for multiplier in config.batch_drift_multipliers):
            batch_specific_drift = np.array(
                [
                    1.0 + residual_drift_strength * batch_drift_lookup.get(batch, 1.0) * (2.0 * batch_local_positions[idx] - 1.0)
                    for idx, batch in enumerate(batch_labels)
                ],
                dtype=float,
            )
        else:
            batch_specific_drift = np.array(
                [1.0 + residual_drift_strength * batch_drift_lookup.get(batch, 1.0) * scaled_orders[idx] for idx, batch in enumerate(batch_labels)],
                dtype=float,
            )
        residual_drift = batch_specific_drift
        values *= residual_drift
        values *= group_factor_vector
        if feature_index in matrix_effect_indices and config.matrix_effect_sample_type:
            target_mask = np.array([sample_type == config.matrix_effect_sample_type for sample_type in sample_types], dtype=bool)
            values[target_mask] *= config.matrix_effect_strength

        qc_mask = np.array([sample_type == "QC" for sample_type in sample_types], dtype=bool)
        if qc_mask.any():
            values[qc_mask] = (
                base_level
                * shared_sample_factor[qc_mask]
                * rng.lognormal(mean=0.0, sigma=config.qc_noise_sigma, size=qc_mask.sum())
                * residual_drift[qc_mask]
            )

        if feature_index in saturation_feature_indices and config.saturation_power < 1.0:
            finite_values = values[np.isfinite(values)]
            if finite_values.size:
                threshold = float(np.nanquantile(finite_values, config.saturation_threshold_quantile))
                if np.isfinite(threshold) and threshold > 0:
                    above_mask = np.isfinite(values) & (values > threshold)
                    values[above_mask] = threshold + np.power(values[above_mask] - threshold, config.saturation_power)

        intensities[:, feature_index] = values

    non_istd_indices = feature_df.index[~feature_df["is_istd"]].to_numpy()
    sample_indices = np.arange(n_samples)
    non_qc_sample_indices = sample_indices[np.array([sample_type != "QC" for sample_type in sample_types], dtype=bool)]

    for _ in range(config.outlier_count):
        sample_index = int(rng.choice(non_qc_sample_indices))
        feature_index = int(rng.choice(non_istd_indices))
        if rng.random() > 0.5:
            intensities[sample_index, feature_index] *= rng.uniform(*config.high_jump_range)
        else:
            intensities[sample_index, feature_index] *= rng.uniform(*config.low_jump_range)

    if config.carryover_event_count > 0 and config.carryover_feature_fraction > 0 and config.carryover_decay_factors:
        eligible_sources = non_qc_sample_indices[non_qc_sample_indices < (n_samples - len(config.carryover_decay_factors))]
        if len(eligible_sources) > 0:
            carryover_feature_count = max(1, int(round(len(non_istd_indices) * config.carryover_feature_fraction)))
            carryover_features = rng.choice(
                non_istd_indices,
                size=min(carryover_feature_count, len(non_istd_indices)),
                replace=False,
            )
            source_count = min(config.carryover_event_count, len(eligible_sources))
            for source_index in rng.choice(eligible_sources, size=source_count, replace=False):
                for lag, decay in enumerate(config.carryover_decay_factors, start=1):
                    target_index = int(source_index + lag)
                    intensities[target_index, carryover_features] += intensities[source_index, carryover_features] * decay

    for feature_index, feature_row in feature_df.iterrows():
        if feature_row["is_istd"]:
            missing_rate = rng.uniform(*config.istd_missing_range)
            candidate_indices = non_qc_sample_indices
        else:
            missing_rate = rng.uniform(*config.analyte_missing_range)
            candidate_indices = sample_indices

        missing_count = int(round(n_samples * missing_rate))
        if missing_count <= 0:
            continue

        chosen = rng.choice(candidate_indices, size=min(missing_count, len(candidate_indices)), replace=False)
        intensities[chosen, feature_index] = np.nan

    if config.structured_missing_rate > 0:
        target_sample_indices = [
            index
            for index, row in enumerate(sample_info_df.itertuples())
            if (not config.structured_missing_batch or row.Batch == config.structured_missing_batch)
            and (not config.structured_missing_sample_type or row.Sample_Type == config.structured_missing_sample_type)
        ]
        if target_sample_indices:
            target_feature_count = max(1, int(round(len(non_istd_indices) * 0.35)))
            target_features = rng.choice(
                non_istd_indices,
                size=min(target_feature_count, len(non_istd_indices)),
                replace=False,
            )
            missing_count = max(1, int(round(len(target_sample_indices) * config.structured_missing_rate)))
            for feature_index in target_features:
                chosen = rng.choice(
                    np.array(target_sample_indices, dtype=int),
                    size=min(missing_count, len(target_sample_indices)),
                    replace=False,
                )
                intensities[chosen, feature_index] = np.nan

    intensities = np.where(np.isnan(intensities), np.nan, np.maximum(intensities, 1_000.0))
    return intensities, shared_sample_factor


def build_raw_intensity_df(
    sample_info_df: pd.DataFrame,
    feature_df: pd.DataFrame,
    intensity_matrix: np.ndarray,
) -> pd.DataFrame:
    sample_columns = sample_info_df["Sample_Name"].tolist()
    feature_rows = pd.DataFrame(intensity_matrix.T, columns=sample_columns)
    feature_rows.insert(0, "FeatureID", feature_df["FeatureID"].tolist())

    sample_type_row = {"FeatureID": "Sample_Type"}
    for row in sample_info_df.itertuples():
        sample_type_row[row.Sample_Name] = row.Sample_Type

    return pd.concat([pd.DataFrame([sample_type_row]), feature_rows], ignore_index=True)


def write_workbook(
    raw_intensity_df: pd.DataFrame,
    sample_info_df: pd.DataFrame,
    feature_df: pd.DataFrame,
    output_path: Path,
) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        raw_intensity_df.to_excel(writer, sheet_name="RawIntensity", index=False)
        sample_info_df.to_excel(writer, sheet_name="SampleInfo", index=False)

    workbook = load_workbook(output_path)
    raw_sheet = workbook["RawIntensity"]
    red_font = Font(color="FFFF0000")
    istd_feature_ids = set(feature_df.loc[feature_df["is_istd"], "FeatureID"].tolist())

    for row_index in range(3, raw_sheet.max_row + 1):
        feature_id = raw_sheet.cell(row=row_index, column=1).value
        if feature_id in istd_feature_ids:
            raw_sheet.cell(row=row_index, column=1).font = red_font

    workbook.save(output_path)
    workbook.close()


def summarize_dataset(
    sample_info_df: pd.DataFrame,
    feature_df: pd.DataFrame,
    intensity_matrix: np.ndarray,
    config: ScenarioConfig,
) -> None:
    feature_ids = feature_df["FeatureID"].tolist()
    intensity_df = pd.DataFrame(
        intensity_matrix.T,
        index=feature_ids,
        columns=sample_info_df["Sample_Name"].tolist(),
    )
    qc_samples = sample_info_df.loc[sample_info_df["Sample_Type"] == "QC", "Sample_Name"].tolist()
    istd_feature_ids = feature_df.loc[feature_df["is_istd"], "FeatureID"].tolist()

    qc_cv_values = []
    for feature_id in istd_feature_ids:
        values = pd.to_numeric(intensity_df.loc[feature_id, qc_samples], errors="coerce").dropna()
        if len(values) < 2 or values.mean() <= 0:
            continue
        qc_cv_values.append((feature_id, float(values.std(ddof=1) / values.mean() * 100)))

    all_values = intensity_df.apply(pd.to_numeric, errors="coerce")
    missing_count = int(all_values.isna().sum().sum())

    print("\n" + "=" * 70)
    print(f"Scenario: {config.name}")
    print(config.description)
    print("=" * 70)
    print(f"Samples: {len(sample_info_df)}")
    print(f"Features: {len(feature_df)}")
    print(f"ISTDs: {feature_df['is_istd'].sum()}")
    print(f"Sample type counts: {sample_info_df['Sample_Type'].value_counts().to_dict()}")
    print(f"Batch counts: {sample_info_df['Batch'].value_counts().to_dict()}")
    print(f"Missing cells: {missing_count}")
    if qc_cv_values:
        sorted_cv = sorted(qc_cv_values, key=lambda item: item[1])
        print("Best ISTD QC CV%:")
        for feature_id, cv_value in sorted_cv[: min(8, len(sorted_cv))]:
            print(f"  {feature_id}: {cv_value:.2f}%")
    if config.recommended_steps:
        print(f"Recommended validation focus: {', '.join(config.recommended_steps)}")


def build_scenario_manifest_rows() -> list[dict[str, str]]:
    rows = []
    for scenario_name, config in SCENARIO_LIBRARY.items():
        filename = (
            "feature_matrix_with_qc_non_group_AfterVBA.xlsx"
            if scenario_name == "balanced_pipeline"
            else f"{scenario_name}.xlsx"
        )
        rows.append(
            {
                "scenario": config.name,
                "filename": filename,
                "description": config.description,
                "recommended_steps": " / ".join(config.recommended_steps),
                "primary_checks": "; ".join(config.primary_checks),
            }
        )
    return rows


def write_scenario_manifest(output_root: Path) -> None:
    rows = build_scenario_manifest_rows()
    manifest_df = pd.DataFrame(rows)
    manifest_df.to_csv(output_root / "scenario_manifest.csv", index=False, encoding="utf-8-sig")

    lines = [
        "# Scenario Matrix Guide",
        "",
        "Use these generated workbooks as targeted validation inputs for different pipeline risks.",
        "",
        "| Scenario | File | Recommended Step(s) | What to Look For |",
        "| --- | --- | --- | --- |",
    ]
    for row in rows:
        lines.append(
            f"| `{row['scenario']}` | `{row['filename']}` | {row['recommended_steps']} | {row['primary_checks']} |"
        )
    (output_root / "SCENARIO_MATRIX_GUIDE.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


def generate_scenario_workbook(config: ScenarioConfig, output_path: Path, seed: int = SEED) -> Path:
    rng = np.random.default_rng(seed)
    sample_info_df = build_sample_info(rng, config)
    feature_df = build_feature_metadata(rng, config)
    intensity_matrix, shared_sample_factor = simulate_intensity_matrix(sample_info_df, feature_df, rng, config)
    sample_info_df = sample_info_df.copy()
    sample_info_df["Creatinine_mg_dL"] = build_creatinine_values(sample_info_df, shared_sample_factor, rng, config)
    raw_intensity_df = build_raw_intensity_df(sample_info_df, feature_df, intensity_matrix)

    write_workbook(raw_intensity_df, sample_info_df, feature_df, output_path)
    summarize_dataset(sample_info_df, feature_df, intensity_matrix, config)
    print(f"\nSaved workbook: {output_path}")
    return output_path


def build_argument_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Generate DNP-compatible validation workbooks for different mass-spec stress scenarios.",
    )
    parser.add_argument(
        "--scenario",
        default="balanced_pipeline",
        help=(
            "Scenario profile to generate. Supports a single profile such as "
            "'strong_batch' or a composition like 'strong_drift+random_jump+order_confounding'."
        ),
    )
    parser.add_argument(
        "--all",
        action="store_true",
        help="Generate all scenario profiles into data/scenario_matrices/.",
    )
    parser.add_argument(
        "--seed",
        type=int,
        default=SEED,
        help="Random seed for reproducible generation.",
    )
    return parser


def main() -> None:
    args = build_argument_parser().parse_args()
    repo_root = Path(__file__).resolve().parents[1]

    if args.all:
        output_root = repo_root / "data" / "scenario_matrices"
        output_root.mkdir(parents=True, exist_ok=True)
        for scenario_name, config in SCENARIO_LIBRARY.items():
            output_path = output_root / f"{scenario_name}.xlsx"
            generate_scenario_workbook(config, output_path, seed=args.seed)
        write_scenario_manifest(output_root)
        return

    config = compose_scenario_config(args.scenario)
    if config.name == "balanced_pipeline":
        output_path = repo_root / "data" / "feature_matrix_with_qc_non_group_AfterVBA.xlsx"
    else:
        output_path = repo_root / "data" / "scenario_matrices" / f"{config.name}.xlsx"
    generate_scenario_workbook(config, output_path, seed=args.seed)


if __name__ == "__main__":
    main()
